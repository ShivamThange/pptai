"""Optional rendered-slide export with explicit availability status."""

from __future__ import annotations

import shutil
import subprocess
import hashlib
import json
import sys
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from pptx import Presentation
from domain.evaluation.quality_metadata import PREFIX


@dataclass(frozen=True)
class RenderResult:
    status: str
    image_paths: list[Path]
    backend: str | None = None
    reason: str | None = None
    artifact_hash: str | None = None
    image_hashes: dict[str, str] = field(default_factory=dict)
    findings: list[dict] = field(default_factory=list)
    text_bounds_status: str = "NOT_EVALUATED"


def _text_bounds(slide, slide_no, *, page_size=None, expected_boxes=None):
    findings, errors = [], []

    def inspect(shape, table_metadata=None, table_id=None):
        try:
            # Office table-cell Shape.Type/AlternativeText are not implemented.
            # Read their text geometry and inherit the native table's metadata.
            if table_id is None and shape.Type == 6:  # msoGroup
                for index in range(1, shape.GroupItems.Count + 1):
                    inspect(shape.GroupItems.Item(index))
                return
            if table_id is None and page_size is not None:
                description = shape.AlternativeText or ""
                metadata = json.loads(description[len(PREFIX):]) if description.startswith(PREFIX) else {}
                box = (shape.Left, shape.Top, shape.Width, shape.Height)
                if box[0] < -1 or box[1] < -1 or box[0] + box[2] > page_size[0] + 1 or box[1] + box[3] > page_size[1] + 1:
                    findings.append({"slide_no": slide_no, "shape_id": shape.Id, "required": metadata.get("required", True),
                                     "check": "POWERPOINT_MEASURED_OUT_OF_BOUNDS", "actual_box_pt": box})
                expected = (expected_boxes or {}).get(shape.Name)
                if expected is not None and any(abs(actual - planned) > 2 for actual, planned in zip(box, expected)):
                    findings.append({"slide_no": slide_no, "shape_id": shape.Id, "required": metadata.get("required", True),
                                     "check": "POWERPOINT_MEASURED_GEOMETRY_CHANGE", "actual_box_pt": box, "expected_box_pt": expected})
            if shape.HasTable:
                description = shape.AlternativeText or ""
                metadata = json.loads(description[len(PREFIX):]) if description.startswith(PREFIX) else {}
                for row in range(1, shape.Table.Rows.Count + 1):
                    for col in range(1, shape.Table.Columns.Count + 1):
                        inspect(shape.Table.Cell(row, col).Shape, metadata, shape.Id)
                return
            if not shape.HasTextFrame or not shape.TextFrame2.HasText:
                return
            frame = shape.TextFrame2
            text = frame.TextRange
            available_w = shape.Width - frame.MarginLeft - frame.MarginRight
            available_h = shape.Height - frame.MarginTop - frame.MarginBottom
            if text.BoundHeight > available_h + 1 or text.BoundWidth > available_w + 1:
                metadata = table_metadata or {}
                if table_id is None:
                    description = shape.AlternativeText or ""
                    if description.startswith(PREFIX):
                        metadata = json.loads(description[len(PREFIX):])
                findings.append({"slide_no": slide_no, "shape_id": table_id if table_id is not None else shape.Id, "required": metadata.get("required", True),
                                 "check": "POWERPOINT_MEASURED_TEXT_OVERFLOW", "bound_height_pt": text.BoundHeight,
                                 "bound_width_pt": text.BoundWidth, "available_height_pt": available_h,
                                 "available_width_pt": available_w})
        except Exception as error:
            errors.append(str(error))

    for index in range(1, slide.Shapes.Count + 1):
        inspect(slide.Shapes.Item(index))
    return findings, errors


def _clear_pngs(directory: Path) -> None:
    for image in directory.glob("slide_*.png"):
        image.unlink()


def render_pptx_slides(pptx_path: str | Path, output_dir: str | Path) -> RenderResult:
    """Serialize native rendering across worker and independent CLI processes."""
    from config.service_settings import get_service_settings
    from infrastructure.utils.artifacts import artifact_lock
    with artifact_lock(get_service_settings().root / ".locks" / "native-render.lock"):
        return _render_pptx_slides_isolated(pptx_path, output_dir)


def _render_pptx_slides_isolated(pptx_path: str | Path, output_dir: str | Path) -> RenderResult:
    """Isolate native COM crashes/hangs; failed exports remain explicitly unavailable."""
    source, target = Path(pptx_path).resolve(), Path(output_dir).resolve()
    digest = hashlib.sha256(source.read_bytes()).hexdigest() if source.is_file() else None
    try:
        # Office's export API rejects long Windows paths, so copy into a short
        # work directory. Keep that directory inside the service root: system
        # temp (mkdtemp mode 0700) can be unwritable under the Windows sandbox.
        from config.service_settings import get_service_settings
        work_root = get_service_settings().root / ".render-work"
        work_root.mkdir(parents=True, exist_ok=True)
        temporary = work_root / uuid.uuid4().hex[:12]
        temporary.mkdir()
        try:
            worker_source = temporary / "candidate.pptx"
            worker_target = temporary / "render"
            shutil.copyfile(source, worker_source)
            process = subprocess.run([sys.executable, "-m", "pipeline.slide_renderer", str(worker_source), str(worker_target)],
                                     cwd=Path(__file__).resolve().parents[1], capture_output=True, text=True,
                                     timeout=60, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
            if process.returncode:
                raise RuntimeError(f"Render worker exited with code {process.returncode}: {process.stderr[-1500:]}")
            data = json.loads(process.stdout)
            target.mkdir(parents=True, exist_ok=True)
            copied, hashes = [], {}
            for original in data['image_paths']:
                path = Path(original)
                if path.parent.resolve() != worker_target.resolve():
                    raise ValueError("Render worker returned an image outside its directory")
                if hashlib.sha256(path.read_bytes()).hexdigest() != data['image_hashes'].get(str(path)):
                    raise ValueError("Rendered image changed after export")
                destination = target / path.name
                shutil.copyfile(path, destination)
                copied.append(destination)
                hashes[str(destination)] = data['image_hashes'][str(path)]
            data['image_paths'], data['image_hashes'] = copied, hashes
            result = RenderResult(**data)
        finally:
            shutil.rmtree(temporary, ignore_errors=True)
        if result.artifact_hash != digest or (source.is_file() and hashlib.sha256(source.read_bytes()).hexdigest() != digest):
            raise ValueError("Render worker hash does not match the current artifact")
        if result.status == "EVALUATED_PASS":
            if len(result.image_paths) != len(Presentation(source).slides):
                raise ValueError("Render worker returned an incomplete slide set")
            for path in result.image_paths:
                if hashlib.sha256(path.read_bytes()).hexdigest() != result.image_hashes.get(str(path)):
                    raise ValueError("Rendered image changed after export")
        (target / 'render_manifest.json').write_text(json.dumps({
            'artifact_hash': digest, 'backend': result.backend, 'image_hashes': result.image_hashes,
            'findings': result.findings, 'text_bounds_status': result.text_bounds_status}, indent=2), encoding='utf-8')
        return result
    except Exception as error:
        return RenderResult("FAILED_TO_RUN", [], reason=str(error), artifact_hash=digest)


def _render_pptx_slides_local(pptx_path: str | Path, output_dir: str | Path) -> RenderResult:
    source = Path(pptx_path).resolve()
    target = Path(output_dir).resolve()
    target.mkdir(parents=True, exist_ok=True)
    if not source.is_file():
        return RenderResult("FAILED_TO_RUN", [], reason=f"Presentation not found: {source}")
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    try:
        expected_count = len(Presentation(source).slides)
        if not expected_count:
            raise ValueError("Presentation is empty")
        original = Presentation(source)
        expected_boxes = [{shape.name: tuple(value / 12700 for value in (shape.left, shape.top, shape.width, shape.height))
                           for shape in slide.shapes} for slide in original.slides]
    except Exception as error:
        return RenderResult("FAILED_TO_RUN", [], reason=str(error), artifact_hash=digest)

    def completed(backend, findings=None, bounds_status="NOT_EVALUATED"):
        from PIL import Image
        images = sorted(target.glob("slide_*.png"))
        if len(images) != expected_count or hashlib.sha256(source.read_bytes()).hexdigest() != digest:
            raise ValueError("Incomplete render or presentation changed during rendering")
        for image in images:
            with Image.open(image) as decoded:
                decoded.verify()
        hashes = {str(image): hashlib.sha256(image.read_bytes()).hexdigest() for image in images}
        result = RenderResult("EVALUATED_PASS", images, backend=backend, artifact_hash=digest,
                              image_hashes=hashes, findings=findings or [], text_bounds_status=bounds_status)
        (target / "render_manifest.json").write_text(json.dumps({"artifact_hash": digest, "backend": backend,
            "image_hashes": hashes, "findings": result.findings, "text_bounds_status": bounds_status}, indent=2), encoding="utf-8")
        return result
    _clear_pngs(target)

    presentation = None
    app = None
    initialized = False
    com_error = "PowerPoint COM did not produce slide images"
    try:
        import pythoncom
        import win32com.client

        pythoncom.CoInitialize()
        initialized = True
        app = win32com.client.DispatchEx("PowerPoint.Application")
        # WithWindow=False makes PowerPoint refuse to open presentations that
        # contain raster images (observed with COM on Windows); keep a window.
        presentation = app.Presentations.Open(str(source), ReadOnly=True, Untitled=False, WithWindow=True)
        findings, bounds_errors = [], []
        for index in range(1, presentation.Slides.Count + 1):
            image = target / f"slide_{index:03d}.png"
            presentation.Slides(index).Export(str(image), "PNG", 1920, 1080)
            rows, errors = _text_bounds(presentation.Slides(index), index,
                                       page_size=(presentation.PageSetup.SlideWidth, presentation.PageSetup.SlideHeight),
                                       expected_boxes=expected_boxes[index - 1])
            findings.extend(rows)
            bounds_errors.extend(errors)
        return completed("powerpoint-com", findings, "FAILED_TO_RUN" if bounds_errors else
                         ("EVALUATED_FAIL" if findings else "EVALUATED_PASS"))
    except Exception as error:
        com_error = str(error)
    finally:
        if presentation is not None:
            try:
                presentation.Close()
            except Exception:
                pass
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass
        presentation = None
        app = None
        if initialized:
            try:
                pythoncom.CoUninitialize()
            except Exception:
                pass

    soffice = shutil.which("soffice")
    if not soffice:
        _clear_pngs(target)
        return RenderResult("FAILED_TO_RUN", [], reason=f"PowerPoint COM unavailable ({com_error}); LibreOffice not found", artifact_hash=digest)
    try:
        _clear_pngs(target)
        pdf_path = target / f"{source.stem}.pdf"
        if pdf_path.exists():
            pdf_path.unlink()
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", str(source), "--outdir", str(target)],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=120,
        )
        import fitz

        document = fitz.open(pdf_path)
        try:
            for index, page in enumerate(document, start=1):
                image = target / f"slide_{index:03d}.png"
                page.get_pixmap(dpi=150).save(str(image))
        finally:
            document.close()
        return completed("libreoffice-pymupdf")
    except Exception as error:
        _clear_pngs(target)
        return RenderResult("FAILED_TO_RUN", [], reason=f"LibreOffice/PyMuPDF rendering failed: {error}", artifact_hash=digest)


if __name__ == "__main__":
    from dataclasses import asdict
    result = _render_pptx_slides_local(sys.argv[1], sys.argv[2])
    print(json.dumps(asdict(result), default=str))
