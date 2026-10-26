"""Read-only workbook extraction. Formulas and stored results stay distinct."""
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import posixpath
from schemas.ingestion_schema import ParseFinding
from core.job_context import check
from .common import base, block, table, finish

NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
REL = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'


def tagged(value):
    if value is None:
        return None, 'blank'
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(), type(value).__name__
    if isinstance(value, (timedelta, Decimal)):
        return str(value), type(value).__name__
    if isinstance(value, float):
        import math
        if not math.isfinite(value):
            return str(value), 'nonfinite'
        return repr(value), 'float'
    return value, type(value).__name__


def extract(source, output, options, *, context=None, limits=None):
    from openpyxl import load_workbook
    from openpyxl.utils.cell import coordinate_to_tuple
    ir, raw_sheets = base(source), []
    max_cells = getattr(limits, 'max_workbook_cells', None)
    processed = 0
    formulas = load_workbook(source, read_only=True, data_only=False)
    cached = load_workbook(source, read_only=True, data_only=True)
    try:
        with ZipFile(source) as package:
            workbook = ET.fromstring(package.read('xl/workbook.xml'))
            rels = {r.get('Id'): r.get('Target') for r in ET.fromstring(package.read('xl/_rels/workbook.xml.rels'))}
            for ordinal, sheet in enumerate(workbook.find('m:sheets', NS), 1):
                check(context)
                name = sheet.get('name')
                target = rels[sheet.get(REL)]
                part = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
                xml = ET.fromstring(package.read(part))
                sheet_raw = {'name': name, 'index': ordinal, 'visibility': sheet.get('state', 'visible'),
                             'part': part, 'merged_ranges': [r.get('ref') for r in xml.findall('m:mergeCells/m:mergeCell', NS)],
                             'cells': [], 'unprocessed_cells': []}
                raw_sheets.append(sheet_raw)
                if name not in formulas.sheetnames or not hasattr(formulas[name], 'reset_dimensions'):
                    ir['findings'].append(ParseFinding(code='UNINTERPRETED_WORKBOOK_OBJECT', message='Chart sheet retained without interpretation.',
                        location={'sheet_name': name, 'sheet_index': ordinal}).model_dump())
                    continue
                ws, values = formulas[name], cached[name]
                ws.reset_dimensions(); values.reset_dimensions()
                xml_cells = {c.get('r'): c for c in xml.findall('m:sheetData/m:row/m:c', NS)}
                bounds = [coordinate_to_tuple(c) for c in xml_cells]
                if not bounds:
                    continue
                max_row = max(r for r, c in bounds)
                max_col = max(c for r, c in bounds)
                cell_rows = {}
                for row, value_row in zip(ws.iter_rows(max_row=max_row, max_col=max_col), values.iter_rows(max_row=max_row, max_col=max_col)):
                    check(context)
                    for cell, value_cell in zip(row, value_row):
                        coordinate = getattr(cell, 'coordinate', None)
                        if coordinate not in xml_cells:
                            continue
                        meta = xml_cells[coordinate]
                        raw_xml = ET.tostring(meta, encoding='unicode')
                        if max_cells is not None and processed >= max_cells:
                            sheet_raw['unprocessed_cells'].append(coordinate)
                            continue
                        processed += 1
                        formula = cell.value if cell.data_type == 'f' else None
                        value, value_type = tagged(value_cell.value if formula else cell.value)
                        cache_present = formula and meta.find('m:v', NS) is not None and meta.find('m:v', NS).text is not None
                        state = 'cached_unverified' if cache_present else 'unknown' if formula else None
                        # Formula constants never appear in the evidence text or table rows.
                        text = str(value) if value is not None else ''
                        lexical = meta.findtext('m:v', default=None, namespaces=NS)
                        if not formula and cell.data_type == 'n' and lexical is not None:
                            # Use package decimals as evidence rather than a rounded float.
                            text = lexical
                            if value_type == 'float':
                                value = lexical
                        if formula and not cache_present:
                            text = ''
                            ir['findings'].append(ParseFinding(code='FORMULA_VALUE_UNKNOWN', message='No stored formula result; formula is not numeric evidence.',
                                location={'sheet_name': name, 'cell_range': coordinate}).model_dump())
                        elif formula:
                            text = f'{text} (cached, unverified)'
                        ref = f"{ir['document_id']}:xlsx:sheet:{ordinal}:cell:{coordinate}"
                        ir['source_item_ids'].append(ref)
                        loc = {'kind': 'xlsx_cell', 'document_id': ir['document_id'], 'source_id': ref,
                               'sheet_name': name, 'sheet_index': ordinal, 'cell_range': coordinate}
                        cell_data = {'row': cell.row - 1, 'column': cell.column - 1, 'text': text,
                                     'lexical_value': lexical,
                                     'raw_value': value if not formula else None, 'value_type': cell.data_type + ':' + value_type,
                                     'formula': formula, 'cached_value': value if cache_present else None,
                                     'cached_state': state, 'number_format': cell.number_format,
                                     'source_refs': [ref], 'locations': [loc]}
                        cell_rows.setdefault(cell.row, []).append(cell_data)
                        sheet_raw['cells'].append({'coordinate': coordinate, **cell_data, 'original_xml': raw_xml})
                        block(ir, ref, text, loc, section=name, kind='xlsx_cell')
                if sheet_raw['unprocessed_cells']:
                    ir['findings'].append(ParseFinding(code='WORKBOOK_CELL_LIMIT', severity='error', blocks_generation=True,
                        message='Configured workbook cell limit reached; cells remain unprocessed.',
                        location={'sheet_name': name, 'cells': sheet_raw['unprocessed_cells']}).model_dump())
                for start in range(min(r for r,c in bounds), max_row + 1, 200):
                    check(context)
                    rows, cells = [], []
                    for row_no in range(start, min(start + 200, max_row + 1)):
                        row_values = [''] * max_col
                        for cell in cell_rows.get(row_no, []):
                            row_values[cell['column']] = cell['text']
                            cells.append({**cell, 'row': row_no - start})
                        rows.append(row_values)
                    table(ir, output, f"{ir['document_id']}:xlsx:sheet:{ordinal}:table:{start}",
                          {'headers': [], 'rows': rows, 'cells': cells}, section=name)
            objects = [p for p in package.namelist() if p.startswith(('xl/charts/', 'xl/drawings/', 'xl/pivotTables/', 'xl/chartsheets/', 'xl/externalLinks/', 'xl/comments', 'xl/embeddings/'))]
            if objects:
                ir['findings'].append(ParseFinding(code='UNINTERPRETED_WORKBOOK_OBJECT',
                    message='Workbook drawings, charts or pivots retained without interpretation.', location={'parts': objects}).model_dump())
    finally:
        formulas.close(); cached.close()
    return finish(ir, output, {'sheets': raw_sheets, 'cell_count': processed, 'objects': objects})
