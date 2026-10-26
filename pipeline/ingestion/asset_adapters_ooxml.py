"""Read-only internal OOXML relationships, occurrences, cells and chart literals."""
import json
import posixpath
import re
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as ET
from core.artifacts import atomic_json, file_hash, sha256_value
from core.job_context import check
from schemas.asset_schema import AssetSourceItem, ChartAssetContent, ChartPointBinding, FormulaAssetContent
from schemas.raw_extraction_schema import SourceAsset, SourceLocation
from .asset_adapters_tables import rich_cell, table_content

NS = {'c':'http://schemas.openxmlformats.org/drawingml/2006/chart',
      'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
      'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
      'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm':'http://schemas.openxmlformats.org/officeDocument/2006/math'}


def xml(package, part):
    return ET.fromstring(package.read(part), parser=ET.XMLParser(resolve_entities=False, no_network=True))


def relationships(package, part):
    relpart = posixpath.join(posixpath.dirname(part), '_rels', posixpath.basename(part) + '.rels')
    if relpart not in package.namelist():
        return {}
    result = {}
    for node in xml(package, relpart):
        target = node.get('Target', '')
        if node.get('TargetMode') == 'External' or '\\' in target or ':' in target:
            continue
        target = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(part), target))
        if target.startswith('../') or target not in package.namelist():
            continue
        result[node.get('Id')] = target
    return result


def register(raw, document_id, package_path, part, node, chain=(), embedded=None):
    path = node.getroottree().getpath(node)
    source_id = document_id + ':ooxml:' + part + ':' + sha256_value(path)[:20]
    if source_id not in raw.source_item_ids:
        raw.source_item_ids.append(source_id)
        raw.asset_source_items.append(AssetSourceItem(source_id=source_id, document_id=document_id,
            package_sha256=file_hash(package_path), package_path=str(package_path), part=part, xml_path=path,
            snippet_sha256=__import__('hashlib').sha256(ET.tostring(node)).hexdigest(),
            relationship_chain=list(chain), embedded_sha256=embedded))
    location = SourceLocation(document_id=document_id, source_id=source_id, kind='ooxml_item', package_part=part, xml_path=path,
                              relationship_id=chain[-1] if chain else None)
    return source_id, location


def omml_expression(node):
    kind = ET.QName(node).localname
    children = list(node)
    if kind == 't':
        return node.text or ''
    if kind.endswith('Pr'):
        return ''
    if kind == 'f':
        num, den = node.find('m:num', NS), node.find('m:den', NS)
        if num is None or den is None:
            raise ValueError('Unsupported OMML fraction')
        return '\\frac{' + omml_expression(num) + '}{' + omml_expression(den) + '}'
    if kind in {'sSup','sSub'}:
        base = node.find('m:e', NS)
        script = node.find('m:sup' if kind == 'sSup' else 'm:sub', NS)
        return '{' + omml_expression(base) + '}' + ('^' if kind == 'sSup' else '_') + '{' + omml_expression(script) + '}'
    if kind in {'oMath','oMathPara','num','den','r','e','sup','sub'}:
        return ''.join(omml_expression(c) for c in children)
    raise ValueError('FORMULA_UNSUPPORTED: OMML structure')


def _docx_tables(raw,package,package_path,tree,root,document_id):
    from .asset_adapters import representation
    w='{'+NS['w']+'}'
    footnotes=xml(package,'word/footnotes.xml') if 'word/footnotes.xml' in package.namelist() else None
    for table_node in tree.xpath('.//w:tbl',namespaces=NS):
        table_ref,table_loc=register(raw,document_id,package_path,'word/document.xml',table_node)
        rows=table_node.findall('w:tr',NS)
        width=len(table_node.findall('w:tblGrid/w:gridCol',NS))
        cells=[];active={};refs=[table_ref];locations=[table_loc];annotations=[]
        for ri,row in enumerate(rows):
            before=row.find('w:trPr/w:gridBefore',NS)
            col=int(before.get(w+'val',0)) if before is not None else 0
            header=row.find('w:trPr/w:tblHeader',NS) is not None
            for node in row.findall('w:tc',NS):
                span=node.find('w:tcPr/w:gridSpan',NS);span=int(span.get(w+'val')) if span is not None else 1
                ref,loc=register(raw,document_id,package_path,'word/document.xml',node)
                refs.append(ref);locations.append(loc)
                text='\n'.join(''.join(p.xpath('.//w:t/text()',namespaces=NS)) for p in node.findall('w:p',NS))
                merge=node.find('w:tcPr/w:vMerge',NS)
                continuation=merge is not None and merge.get(w+'val')!='restart'
                if continuation:
                    anchor=active.get(col)
                    if anchor is None or anchor['column_span']!=span:
                        raise ValueError('TABLE_SPAN_INVALID: DOCX continuation without matching anchor')
                    anchor['row_span']+=1;anchor['source_refs'].append(ref);anchor['locations'].append(loc.model_dump(mode='json'))
                    if text:anchor['text']+='\n'+text
                    cell=anchor
                else:
                    cell={'row':ri,'column':col,'row_span':1,'column_span':span,'text':text,'column_header':header,
                          'source_refs':[ref],'locations':[loc.model_dump(mode='json')]}
                    cells.append(cell)
                    for ci in range(col,col+span):
                        if merge is not None:active[ci]=cell
                        else:active.pop(ci,None)
                for reference in node.xpath('.//w:footnoteReference',namespaces=NS):
                    identifier=reference.get(w+'id')
                    found=footnotes.xpath('./w:footnote[@w:id=$identifier]',namespaces=NS,identifier=identifier) if footnotes is not None else []
                    if len(found)==1:
                        fref,floc=register(raw,document_id,package_path,'word/footnotes.xml',found[0])
                        refs.append(fref);locations.append(floc)
                        aid=fref+':annotation'
                        cell.setdefault('footnote_refs',[]).append(aid)
                        annotations.append({'annotation_id':aid,'kind':'footnote','text':' '.join(found[0].xpath('.//w:t/text()',namespaces=NS)),
                            'source_refs':[fref],'scope':'cell','row':ri,'column':col})
                col+=span
            width=max(width,col)
        if not cells or not width:
            continue
        grid=[['']*width for _ in rows]
        for cell in cells:grid[cell['row']][cell['column']]=cell['text']
        aid=table_ref+':table'
        path=root/'assets'/'ooxml'/(sha256_value(aid)+'.json')
        data={'headers':[],'rows':grid,'cells':cells};atomic_json(path,data)
        asset=SourceAsset(asset_id=aid,kind='table',section_path=['Root'],document_id=document_id,
            source_refs=list(dict.fromkeys(refs)),locations=list({l.source_id:l for l in locations}.values()),
            data=data,file_path=str(path),content_sha256=file_hash(path))
        asset.content=table_content(asset,[representation(path,'original','application/json','native_table')])
        asset.content.annotations.extend(annotations)
        raw.assets.append(asset)


def _embedded_chart_sheets(raw,package,package_path,chart_part,chart_node,document_id):
    """Resolve actual embedded XLSX literals; cached/formula results stay unusable."""
    import io
    import hashlib
    from openpyxl import load_workbook
    external=chart_node.find('c:externalData',NS)
    if external is None:
        return {}
    rid=external.get('{'+NS['r']+'}id')
    target=relationships(package,chart_part).get(rid)
    if not target or not target.endswith('.xlsx'):
        return {}
    embedded=package.read(target)
    embedded_hash=hashlib.sha256(embedded).hexdigest()
    register(raw,document_id,package_path,chart_part,external,[rid],embedded_hash)
    result={};main={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    book=load_workbook(io.BytesIO(embedded),read_only=False,data_only=False,keep_links=False)
    try:
        with ZipFile(io.BytesIO(embedded)) as inner:
            workbook=xml(inner,'xl/workbook.xml');rels=relationships(inner,'xl/workbook.xml')
            for sheet in workbook.findall('s:sheets/s:sheet',main):
                part=rels.get(sheet.get('{'+NS['r']+'}id'))
                name=sheet.get('name')
                if not part or name not in book.sheetnames:
                    continue
                tree=xml(inner,part);result[name]={}
                nodes=tree.findall('s:sheetData/s:row/s:c',main)
                if len(nodes)>100000:
                    raise ValueError('CHART_REFERENCE_UNRESOLVED: embedded workbook cell limit')
                for node in nodes:
                    coordinate=node.get('r');cell=book[name][coordinate]
                    path=node.getroottree().getpath(node)
                    ref=document_id+':ooxml:'+target+'!'+part+':'+sha256_value(path)[:20]
                    if ref not in raw.source_item_ids:
                        raw.source_item_ids.append(ref)
                        raw.asset_source_items.append(AssetSourceItem(source_id=ref,document_id=document_id,
                            package_sha256=file_hash(package_path),package_path=str(package_path),part=part,xml_path=path,
                            snippet_sha256=hashlib.sha256(ET.tostring(node)).hexdigest(),embedded_sha256=embedded_hash,
                            embedded_package_part=target,relationship_chain=[rid]))
                    formula=node.find('s:f',main)
                    lexical=node.findtext('s:v',namespaces=main)
                    text=lexical if cell.data_type=='n' and lexical is not None else str(cell.value) if cell.value is not None else ''
                    result[name][coordinate]={'text':text,'formula':(formula.text or 'UNCOMPUTED_FORMULA') if formula is not None else None,
                        'source_refs':[ref],'locations':[SourceLocation(document_id=document_id,source_id=ref,kind='ooxml_item',
                            package_part=target+'!'+part,xml_path=path,cell_range=coordinate,sheet_name=name).model_dump(mode='json')]}
    finally:
        book.close()
    return result


def _native_charts(raw, package, package_path, part, node, chain, root, document_id, sheets, occurrence):
    from decimal import Decimal
    types = node.xpath('.//c:plotArea/*[contains(local-name(), "Chart")]', namespaces=NS)
    refs = [occurrence[0]]
    locs = [occurrence[1]]
    def literal(source_node):
        ref, loc = register(raw, document_id, package_path, part, source_node, chain)
        refs.append(ref); locs.append(loc)
        return ref
    literal(node)
    chart_type = 'unsupported'
    if len(types) == 1:
        kind = ET.QName(types[0]).localname
        chart_type = {'lineChart':'line','areaChart':'area','pieChart':'pie','barChart':'column'}.get(kind, 'unsupported')
        direction=types[0].find('c:barDir', NS)
        if kind == 'barChart' and direction is not None and direction.get('val') == 'bar':
            chart_type = 'bar'
    aid = refs[0] + ':chart'
    supported = chart_type != 'unsupported' and len(node.xpath('.//c:valAx', namespaces=NS)) <= 1 and not node.xpath('.//c:logBase|.//c:pivotSource', namespaces=NS)
    categories, series, points, labels = [], [], [], []
    origin = 'native_literals'
    def values(container):
        nonlocal origin
        if container is None:
            raise ValueError('Missing chart source')
        form = container.find('.//c:f', NS)
        if form is not None:
            # References resolve original retained cells, never cache values.
            origin = 'native_cells'
            match = re.fullmatch(r"(?:'((?:[^']|'')+)'|([^'!]+))!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?", form.text or '')
            if not match:
                raise ValueError('CHART_REFERENCE_UNRESOLVED')
            name = (match[1] or match[2]).replace("''", "'")
            from openpyxl.utils.cell import range_boundaries, get_column_letter
            bounds = range_boundaries(match[3]+match[4]+':'+(match[5] or match[3])+(match[6] or match[4]))
            if bounds[0] != bounds[2] and bounds[1] != bounds[3]:
                raise ValueError('Chart range is not one dimensional')
            cells = sheets.get(name, {})
            result = []
            for row in range(bounds[1],bounds[3]+1):
                for col in range(bounds[0],bounds[2]+1):
                    cell = cells.get(get_column_letter(col)+str(row))
                    if not cell or cell.get('formula') or cell.get('cached_state'):
                        raise ValueError('CHART_DATA_UNVERIFIED')
                    result.append((cell['text'],cell['source_refs'][0]))
                    refs.extend(cell['source_refs']); locs.extend(SourceLocation.model_validate(v) for v in cell['locations'])
            return result
        literal_source = next((c for c in container if ET.QName(c).localname in {'strLit','numLit','v'}), None)
        if literal_source is None:
            raise ValueError('CHART_DATA_UNVERIFIED: cache only')
        if ET.QName(literal_source).localname == 'v':
            return [(literal_source.text or '', literal(literal_source))]
        pts = literal_source.findall('c:pt', NS)
        count = literal_source.find('c:ptCount', NS)
        if count is None or len(pts) != int(count.get('val')) or sorted(int(p.get('idx')) for p in pts) != list(range(len(pts))):
            raise ValueError('CHART_REFERENCE_UNRESOLVED: missing indices')
        return [(p.findtext('c:v', namespaces=NS), literal(p)) for p in sorted(pts, key=lambda p:int(p.get('idx')))]
    try:
        if not supported:
            raise ValueError('CHART_ENCODING_UNSUPPORTED')
        if not sheets:
            sheets=_embedded_chart_sheets(raw,package,package_path,part,node,document_id)
        for si, ser in enumerate(types[0].findall('c:ser', NS)):
            cats = values(ser.find('c:cat', NS))
            nums = values(ser.find('c:val', NS))
            name_values = values(ser.find('c:tx', NS))
            if len(name_values) != 1 or not cats or len(cats) != len(nums):
                raise ValueError('CHART_REFERENCE_UNRESOLVED')
            names = [v[0] for v in cats]
            if series and names != categories:
                raise ValueError('CHART_REFERENCE_UNRESOLVED: mismatched categories')
            categories = names
            labels.extend([v[1] for v in cats] + [name_values[0][1]])
            data = []
            for ci,(text,ref) in enumerate(nums):
                number = Decimal(text)
                if not number.is_finite():
                    raise ValueError('CHART_DATA_UNVERIFIED')
                data.append(str(number))
                points.append(ChartPointBinding(series_index=si, category_index=ci, lexical_value=text,
                                               decimal_value=str(number), value_refs=[ref]))
            series.append({'name':name_values[0][0], 'values':data})
        if not series or chart_type == 'pie' and (len(series)!=1 or any(Decimal(v)<0 for v in series[0]['values']) or not any(Decimal(v) for v in series[0]['values'])):
            raise ValueError('CHART_ENCODING_UNSUPPORTED')
        state, status = 'VERIFIED_SOURCE_DATA', 'EDITABLE'
        findings = []
    except (ValueError, TypeError, ArithmeticError, OSError, KeyError, ET.LxmlError) as error:
        state, status = 'UNAVAILABLE', 'UNSUPPORTED'
        categories, series, points, labels = [], [], [], []
        findings = [str(error)]
    path = root / 'assets' / 'ooxml' / (sha256_value(aid) + '.json')
    atomic_json(path, {'chart_type': chart_type, 'categories':categories, 'series':series, 'data_state':state})
    from .asset_adapters import representation
    content = ChartAssetContent(source_sha256=file_hash(package_path), chart_type=chart_type,
        data_origin=origin, data_state=state, reconstruction_status=status, categories=categories, series=series,
        points=points, label_refs=labels, representations=[representation(path,'chart_data_json','application/json','native_chart')])
    raw.assets.append(SourceAsset(asset_id=aid, kind='chart', section_path=['Root'], document_id=document_id,
        source_refs=list(dict.fromkeys(refs)), locations=list({l.source_id:l for l in locs}.values()),
        file_path=str(path), content_sha256=file_hash(path), content=content, findings=findings))


def supplement_ooxml(raw, bundles, root, *, context=None):
    for document_id, directory in bundles.items():
        check(context)
        descriptor = next((d for d in raw.documents if d.document_id == document_id), None)
        if not descriptor or descriptor.resolved_type not in {'docx','xlsx'}:
            continue
        package_path = directory / 'source_snapshot.bin'
        sheets = {}
        if descriptor.resolved_type == 'xlsx':
            payload = json.loads((directory / 'adapter_raw.json').read_text(encoding='utf-8'))
            for sheet in payload.get('sheets', []):
                sheets[sheet['name']] = {c['coordinate']:c for c in sheet['cells']}
                if not sheet['cells']:
                    continue
                # Canonical whole-sheet grid retains merges crossing F02 projections.
                cells = [{k:v for k,v in c.items() if k not in {'coordinate','original_xml'}} for c in sheet['cells']]
                from openpyxl.utils.cell import range_boundaries
                for merged in sheet['merged_ranges']:
                    c1,r1,c2,r2 = range_boundaries(merged)
                    anchor = next((c for c in cells if (c['row'],c['column']) == (r1-1,c1-1)), None)
                    if anchor:
                        anchor.update(row_span=r2-r1+1,column_span=c2-c1+1)
                        cells = [c for c in cells if c is anchor or not (r1-1 <= c['row'] < r2 and c1-1 <= c['column'] < c2)]
                aid = document_id + ':xlsx:sheet:' + str(sheet['index']) + ':full-table'
                path = root / 'assets' / 'ooxml' / (sha256_value(aid)+'.json')
                height = max(c['row'] + c.get('row_span',1) for c in cells)
                width = max(c['column'] + c.get('column_span',1) for c in cells)
                grid = [[''] * width for _ in range(min(height,200))]
                for cell in cells:
                    if cell['row'] < len(grid):
                        grid[cell['row']][cell['column']] = cell['text']
                # Legacy data is a bounded projection; rich content carries the full grid.
                data = {'headers':[], 'rows':grid, 'cells':cells}
                atomic_json(path, data)
                asset = SourceAsset(asset_id=aid, kind='table', section_path=[sheet['name']], document_id=document_id,
                    source_refs=list(dict.fromkeys(r for c in cells for r in c['source_refs'])),
                    locations=[SourceLocation.model_validate(l) for c in cells for l in c['locations']], data=data,
                    file_path=str(path), content_sha256=file_hash(path))
                from .asset_adapters import representation
                asset.content = table_content(asset,[representation(path,'original','application/json','native_table')])
                asset.content.sheet_name, asset.content.sheet_index = sheet['name'], sheet['index']
                raw.assets.append(asset)
        with ZipFile(package_path) as package:
            chart_occurrences = []
            source_parts = ['word/document.xml'] if descriptor.resolved_type == 'docx' else sorted(p for p in package.namelist() if p.startswith('xl/drawings/') and p.endswith('.xml'))
            for part in source_parts:
                if part not in package.namelist():
                    continue
                tree = xml(package, part)
                if part=='word/document.xml':
                    _docx_tables(raw,package,package_path,tree,root,document_id)
                rels = relationships(package, part)
                for node in tree.xpath('.//a:blip', namespaces=NS):
                    rid = node.get('{'+NS['r']+'}embed')
                    target = rels.get(rid)
                    if not target:
                        continue
                    payload = package.read(target)
                    import hashlib
                    ref, loc = register(raw,document_id,package_path,part,node,[rid],hashlib.sha256(payload).hexdigest())
                    path = root / 'assets' / 'ooxml' / (sha256_value(ref)+Path(target).suffix)
                    path.parent.mkdir(parents=True,exist_ok=True); path.write_bytes(payload)
                    from PIL import Image
                    try:
                        with Image.open(path) as image:
                            width,height = image.size
                    except OSError:
                        width=height=None
                    raw.assets.append(SourceAsset(asset_id=ref,kind='image',section_path=['Root'],document_id=document_id,
                        source_refs=[ref],locations=[loc],file_path=str(path),content_sha256=file_hash(path),width=width,height=height,
                        availability='AVAILABLE' if width else 'MISSING', findings=[] if width else ['Unsupported vector media retained']))
                for node in tree.xpath('.//c:chart',namespaces=NS):
                    rid = node.get('{'+NS['r']+'}id')
                    if rid in rels:
                        occurrence=register(raw,document_id,package_path,part,node,[rid])
                        chart_occurrences.append((rels[rid],rid,occurrence))
                for node in tree.xpath('.//m:oMath',namespaces=NS):
                    ref,loc = register(raw,document_id,package_path,part,node)
                    path = root / 'assets' / 'ooxml' / (sha256_value(ref)+'.xml')
                    path.parent.mkdir(parents=True,exist_ok=True); path.write_bytes(ET.tostring(node))
                    try:
                        expression=omml_expression(node); status='SOURCE_LITERAL'
                    except ValueError:
                        expression=None; status='UNKNOWN'
                    raw.assets.append(SourceAsset(asset_id=ref,kind='formula',section_path=['Root'],document_id=document_id,
                        source_refs=[ref],locations=[loc],file_path=str(path),content_sha256=file_hash(path),
                        content=FormulaAssetContent(source_sha256=file_hash(package_path),expression=expression,
                            expression_refs=[ref],expression_format='omml',transcription_status=status)))
            for part,rid,occurrence in chart_occurrences:
                _native_charts(raw,package,package_path,part,xml(package,part),[rid],root,document_id,sheets,occurrence)
