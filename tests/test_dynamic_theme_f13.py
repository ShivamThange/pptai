"""F13 dynamic theming: tokens, injection, selection, and rendered output."""
import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Pt

from application.jobs.builder import PresentationBuilder
from domain.rendering.theme_engine import (LENS_THEME_MAP, apply_theme_to_ir2, available_themes,
                               load_theme, resolve_theme_name)

pytestmark = pytest.mark.filterwarnings("ignore::DeprecationWarning")

ROOT = Path(__file__).resolve().parents[1]
LAYOUTS = ROOT / "config" / "ppt_default_layouts.json"


def sample_ir2():
    return {
        "title": "Theme Deck",
        "output_filename": "theme.pptx",
        "slides": [
            {
                "layout": "HEADER_WITH_BODY",
                "content": {
                    "header": {"type": "heading", "text": "Scaling laws for fused kernels"},
                    "body": {"type": "bullets", "items": ["Throughput rises 2.4x at 70B parameters",
                                                          "Memory cost falls by 38%"]},
                },
                "source_unit_ids": ["u1"],
            },
            {
                "layout": "TABLE_FULL_WIDTH",
                "content": {
                    "header": {"type": "heading", "text": "Results"},
                    "table": {"type": "table", "headers": ["Model", "Score"],
                              "rows": [["A", "91.2"], ["B", "88.4"]]},
                    "source": {"type": "text_group", "text": "Source: table 2."},
                },
            },
            {
                "layout": "KPI_ROW_WITH_CHART",
                "content": {
                    "header": {"type": "heading", "text": "Metrics"},
                    "kpi_1": {"type": "stat", "items": [{"value": "91.2", "label": "Accuracy"}]},
                    "kpi_2": {"type": "stat", "items": [{"value": "3.5B", "label": "Parameters"}]},
                    "kpi_3": {"type": "stat", "items": [{"value": "2.4x", "label": "Speedup"}]},
                    "chart": {"type": "chart", "chart_type": "column", "categories": ["S", "M", "L"],
                              "series": [{"name": "Score", "values": [70.1, 82.4, 91.2]}]},
                },
            },
        ],
    }


def test_all_shipped_themes_load_and_validate():
    themes = available_themes()
    assert set(themes) == {"academic_slate", "corporate_blue", "warm_editorial", "dark_tech", "minimal_mono"}
    for theme_id, view in themes.items():
        theme = load_theme(theme_id)
        assert theme.name == view["name"]
        assert theme.background.startswith("#") and len(theme.chart_colors) >= 4


def test_auto_resolution_maps_every_lens():
    assert resolve_theme_name("auto", "academic") == "academic_slate"
    assert resolve_theme_name("auto", "executive") == "corporate_blue"
    assert resolve_theme_name("auto", "pitch") == "warm_editorial"
    assert resolve_theme_name("auto", "technical") == "minimal_mono"
    assert set(LENS_THEME_MAP) == {"academic", "executive", "pitch", "technical"}
    assert resolve_theme_name("dark_tech") == "dark_tech"


def test_unknown_theme_is_rejected_loudly():
    with pytest.raises(ValueError, match="Unknown theme"):
        load_theme("neon_cyberpunk")


def test_apply_theme_injects_tokens_without_touching_content():
    theme = load_theme("dark_tech")
    original = sample_ir2()
    themed = apply_theme_to_ir2(original, theme)
    # Factual content and provenance survive byte-for-byte (acceptance: theme
    # is presentation-layer only).
    assert themed["slides"][0]["source_unit_ids"] == original["slides"][0]["source_unit_ids"]
    assert themed["slides"][0]["content"]["body"]["items"] == original["slides"][0]["content"]["body"]["items"]
    assert themed["slides"][1]["content"]["table"]["rows"] == original["slides"][1]["content"]["table"]["rows"]
    assert json.dumps(original) == json.dumps(sample_ir2()), "input IR2 must not be mutated"
    # Style tokens injected.
    assert themed["slides"][0]["content"]["header"]["style"]["color"] == "#E2E8F0"
    assert themed["slides"][0]["background_color"] == "#0F172A"
    assert themed["slides"][1]["content"]["table"]["style"]["header_bg"] == "#1E293B"
    assert themed["slides"][1]["content"]["table"]["style"]["band_bg"] == "#182235"
    assert themed["slides"][1]["content"]["source"]["style"]["color"] == "#94A3B8"
    assert themed["slides"][2]["content"]["chart"]["style"]["chart_colors"][0] == "#38BDF8"
    assert themed["slides"][2]["content"]["kpi_1"]["style"]["color"] == "#38BDF8"
    assert themed["slides"][2]["content"]["kpi_1"]["style"]["label_color"] == "#94A3B8"


def test_model_authored_styles_beat_theme_tokens():
    content = sample_ir2()
    content["slides"][0]["content"]["body"]["style"] = {"color": "#123456"}
    themed = apply_theme_to_ir2(content, load_theme("dark_tech"))
    assert themed["slides"][0]["content"]["body"]["style"]["color"] == "#123456"


def test_custom_slide_background_is_preserved():
    content = sample_ir2()
    content["slides"][0]["background_color"] = "#101828"
    themed = apply_theme_to_ir2(content, load_theme("dark_tech"))
    assert themed["slides"][0]["background_color"] == "#101828"
    assert themed["slides"][1]["background_color"] == "#0F172A"


@pytest.mark.parametrize("theme_id", ["academic_slate", "dark_tech", "warm_editorial"])
def test_rendered_pptx_carries_theme(tmp_path, theme_id):
    theme = load_theme(theme_id)
    builder = PresentationBuilder(LAYOUTS, backend="python", theme=theme, dynamic_layout=False)
    target = tmp_path / f"themed_{theme_id}.pptx"
    builder.build(sample_ir2(), str(target))
    deck = Presentation(str(target))
    slide = deck.slides[0]
    # Slide background is the theme background.
    fill = slide.background.fill
    assert fill.type is not None
    assert str(fill.fore_color.rgb) == theme.background.lstrip("#")
    # Heading run color is the theme text token.
    heading = next(s for s in slide.shapes if s.name.endswith("header"))
    assert str(heading.text_frame.paragraphs[0].runs[0].font.color.rgb) == theme.text.lstrip("#")
    if theme.heading_font:
        assert heading.text_frame.paragraphs[0].runs[0].font.name == theme.heading_font
    # Table header fill and body colors follow the theme.
    table_slide = deck.slides[1]
    table = next(s for s in table_slide.shapes if s.has_table).table
    header_cell = table.cell(0, 0)
    assert str(header_cell.fill.fore_color.rgb) == theme.table_header_bg.lstrip("#")
    assert str(header_cell.text_frame.paragraphs[0].runs[0].font.color.rgb) == "FFFFFF"
    banded = table.cell(2, 0)
    assert str(banded.fill.fore_color.rgb) == theme.table_band.lstrip("#")
    # Chart series picked the theme palette.
    chart_slide = deck.slides[2]
    chart = next(s for s in chart_slide.shapes if s.has_chart).chart
    series_color = chart.plots[0].series[0].format.fill.fore_color.rgb
    assert str(series_color) == theme.chart_colors[0].lstrip("#")


def test_academic_slate_matches_legacy_neutral_look(tmp_path):
    """The default theme must reproduce the pre-F13 hardcoded palette."""
    builder = PresentationBuilder(LAYOUTS, backend="python", theme=load_theme("academic_slate"),
                                  dynamic_layout=False)
    target = tmp_path / "legacy_look.pptx"
    builder.build(sample_ir2(), str(target))
    deck = Presentation(str(target))
    heading = next(s for s in deck.slides[0].shapes if s.name.endswith("header"))
    assert str(heading.text_frame.paragraphs[0].runs[0].font.color.rgb) == "0F172A"
    table = next(s for s in deck.slides[1].shapes if s.has_table).table
    assert str(table.cell(0, 0).fill.fore_color.rgb) == "0F172A"
    assert str(table.cell(2, 0).fill.fore_color.rgb) == "F8FAFC"


def test_options_reject_unknown_theme(tmp_path):
    import sys
    import types
    # Sandbox guard: orchestrator imports the Gemini SDK at module level; the
    # theme contract itself does not need it.
    if "google" not in sys.modules:
        google = types.ModuleType("google")
        genai = types.ModuleType("google.genai")
        genai.Client = object
        genai.types = types.ModuleType("google.genai.types")
        google.genai = genai
        sys.modules["google"] = google
        sys.modules["google.genai"] = genai
        sys.modules["google.genai.types"] = genai.types
    from application.pipeline.orchestrator import MultiAgentOptions
    with pytest.raises(ValueError, match="Unknown theme"):
        MultiAgentOptions(layouts_path=str(LAYOUTS), output_dir=str(tmp_path), theme="nonsense")


def test_generation_request_validates_theme():
    from contracts.service_schema import GenerationRequest
    request = GenerationRequest(theme="dark_tech")
    assert request.theme == "dark_tech"
    with pytest.raises(ValueError, match="Unknown theme"):
        GenerationRequest(theme="nonsense")
