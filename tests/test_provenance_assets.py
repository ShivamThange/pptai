"""R02 counterexamples plus integration against actual installed Docling APIs."""

import copy
import hashlib
import json
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from docling_core.types.doc import (
    DoclingDocument,
    DocItemLabel,
    ProvenanceItem,
    BoundingBox,
    ImageRef,
    TableData as DoclingTableData,
    TableCell as DoclingTableCell,
    Size,
)

from core.asset_registry import AssetRegistry, AssetBindingError
from core.builder import PresentationBuilder
from pipeline.stage1_docling_parser import export_docling_document
from pipeline.source_adapter import (
    adapt_structural_ir,
    build_raw_extraction,
    normalize_table_data,
)
from pipeline.knowledge_extractor import (
    build_legacy_projection,
    validate_knowledge_units,
)
from pipeline.slide_composer import validate_slide_spec
from schemas.ir1_schema import ValidatedIR1, KnowledgeUnit, EvidenceRef
from schemas.ir2_content_schema import SlideSpec


def _prov(page=1, text=""):
    return ProvenanceItem(
        page_no=page,
        bbox=BoundingBox(l=10, t=10, r=90, b=50, coord_origin="TOPLEFT"),
        charspan=(0, len(text)),
    )


def make_document():
    doc = DoclingDocument(name="Measurements")
    for page in (1, 2):
        doc.add_page(
            page_no=page,
            size=Size(width=100, height=100),
            image=ImageRef.from_pil(Image.new("RGB", (100, 100), "white"), dpi=72),
        )
    cap = doc.add_text(
        label=DocItemLabel.CAPTION,
        text="Figure 1: Source measurement apparatus",
        prov=_prov(text="Figure 1: Source measurement apparatus"),
    )
    picture = doc.add_picture(
        image=ImageRef.from_pil(Image.new("RGB", (80, 40), "navy"), dpi=72),
        caption=cap,
        prov=_prov(),
    )
    picture.prov.append(_prov(2))
    values = [
        ("Material", True),
        ("Mass (kg)", True),
        ("Copper", False),
        ("12.50 kg", False),
    ]
    cells = [
        DoclingTableCell(
            text=text,
            column_header=header,
            start_row_offset_idx=i // 2,
            end_row_offset_idx=i // 2 + 1,
            start_col_offset_idx=i % 2,
            end_col_offset_idx=i % 2 + 1,
            bbox=BoundingBox(
                l=10 + i % 2 * 40,
                t=10 + i // 2 * 20,
                r=50 + i % 2 * 40,
                b=30 + i // 2 * 20,
                coord_origin="TOPLEFT",
            ),
        )
        for i, (text, header) in enumerate(values)
    ]
    table_caption = doc.add_text(
        label=DocItemLabel.CAPTION,
        text="Table 1: Mass measurements",
        prov=_prov(text="Table 1: Mass measurements"),
    )
    doc.add_table(
        data=DoclingTableData(num_rows=2, num_cols=2, table_cells=cells),
        caption=table_caption,
        prov=_prov(),
    )
    doc.add_text(
        label=DocItemLabel.FORMULA, text="m = 12.50 kg", prov=_prov(text="m = 12.50 kg")
    )
    return doc


@pytest.fixture
def source_run(tmp_path):
    doc = make_document()
    source = tmp_path / "source.json"
    source.write_text(
        json.dumps(doc.export_to_dict(), ensure_ascii=False), encoding="utf-8"
    )
    extraction = tmp_path / "nondefault-extraction"
    structural = export_docling_document(doc, str(source), str(extraction))
    run = tmp_path / "another-run"
    raw = build_raw_extraction(extraction / "ir.json", run)
    return structural, raw, run


def _layouts(kind):
    return {
        "slide_dimensions": {"width_inches": 13.333, "height_inches": 7.5},
        "layouts": {
            "TEST": {
                "regions": {
                    "title": {
                        "left": 0.06,
                        "top": 0.05,
                        "width": 0.88,
                        "height": 0.1,
                        "allowed_types": ["heading"],
                    },
                    "body": {
                        "left": 0.08,
                        "top": 0.2,
                        "width": 0.84,
                        "height": 0.65,
                        "allowed_types": [kind],
                    },
                }
            }
        },
    }


def _slide(kind, asset_id):
    return SlideSpec(
        layout="TEST",
        content={
            "title": {"type": "heading", "text": "Source measurements"},
            "body": {"type": kind, "source_asset_id": asset_id},
        },
    )


def test_real_docling_caption_images_locations_and_source_cells(source_run):
    structural, raw, run = source_run
    assert raw.schema_version == "2"
    image = next(asset for asset in raw.assets if asset.kind == "image")
    assert image.caption == "Figure 1: Source measurement apparatus"
    assert [location.page for location in image.locations] == [1, 2]
    assert image.caption_refs and image.caption_locations[0].page == 1
    assert (image.width, image.height) == (80, 40)
    assert (
        image.content_sha256
        == hashlib.sha256(Path(image.file_path).read_bytes()).hexdigest()
    )
    assert all(Path(asset.file_path).is_relative_to(run) for asset in raw.assets)
    assert len(raw.assets) == 3
    table = next(asset for asset in raw.assets if asset.kind == "table")
    assert table.data["headers"] == ["Material", "Mass (kg)"]
    assert table.data["rows"] == [["Copper", "12.50 kg"]]
    assert len(table.data["cells"]) == 4
    assert all(
        cell["source_refs"] == table.source_refs and cell["locations"]
        for cell in table.data["cells"]
    )
    # Installed TextItem.get_image(doc) crops formulas from page pixels.
    assert (
        next(asset for asset in raw.assets if asset.kind == "formula").availability
        == "AVAILABLE"
    )


@pytest.mark.parametrize("kind", ["image", "table", "formula"])
def test_compose_and_build_from_run_registry_and_reopen_native_pptx(source_run, kind):
    structural, raw, run = source_run
    asset = next(asset for asset in raw.assets if asset.kind == kind)
    slide = _slide(kind, asset.asset_id)
    layouts = _layouts(kind)
    validate_slide_spec(
        slide,
        layouts,
        allowed_unit_ids=set(),
        assets=raw.assets,
        source_item_ids=raw.source_item_ids,
    )
    target = run / f"{kind}.pptx"
    builder = PresentationBuilder(
        layouts,
        asset_registry=AssetRegistry(
            raw.assets, run_root=run, source_item_ids=raw.source_item_ids
        ),
    )
    original = slide.model_dump(exclude_none=True)
    assert builder.build({"slides": [original]}, str(target)) == []
    assert original == slide.model_dump(exclude_none=True)
    reopened = Presentation(target)
    assert len(reopened.slides) == 1
    if kind == "table":
        table = next(
            shape.table for shape in reopened.slides[0].shapes if shape.has_table
        )
        assert table.cell(0, 1).text == "Mass (kg)"
        assert table.cell(1, 1).text == "12.50 kg"
    else:
        assert (
            len(
                [shape for shape in reopened.slides[0].shapes if shape.shape_type == 13]
            )
            == 1
        )


@pytest.mark.parametrize(
    "kind,failure",
    [
        (kind, failure)
        for kind in ("image", "table", "formula")
        for failure in (
            "unknown",
            "wrong_kind",
            "missing",
            "wrong_document",
            "hash",
            "override",
            "unresolved_ref",
        )
        if failure != "override" or kind == "image"
    ],
)
def test_required_asset_counterexamples_fail_composition_and_preflight(
    source_run, kind, failure
):
    _, raw, run = source_run
    assets = copy.deepcopy(raw.assets)
    asset = next(value for value in assets if value.kind == kind)
    slide = _slide(kind, asset.asset_id)
    if failure == "unknown":
        slide.content["body"].source_asset_id = "invented"
    elif failure == "wrong_kind":
        asset.kind = "table" if kind != "table" else "image"
    elif failure == "missing":
        Path(asset.file_path).unlink()
    elif failure == "wrong_document":
        asset.document_id = "different-document"
    elif failure == "hash":
        Path(asset.file_path).write_bytes(b"changed")
    elif failure == "override":
        slide.content["body"].path = str(run / "override.png")
    elif failure == "unresolved_ref":
        asset.source_refs.append(f"{asset.document_id}:invented-ref")
    with pytest.raises(ValueError):
        validate_slide_spec(
            slide,
            _layouts(kind),
            allowed_unit_ids=set(),
            assets=assets,
            source_item_ids=raw.source_item_ids,
        )
    target = run / "must-not-exist.pptx"
    builder = PresentationBuilder(
        _layouts(kind),
        asset_registry=AssetRegistry(
            assets, run_root=run, source_item_ids=raw.source_item_ids
        ),
    )
    with pytest.raises(ValueError):
        builder.build({"slides": [slide.model_dump(exclude_none=True)]}, str(target))
    assert not target.exists()


def test_nondefault_run_prevents_global_asset_substitution(source_run, tmp_path):
    _, raw, run = source_run
    image = next(value for value in raw.assets if value.kind == "image")
    registry = AssetRegistry(
        raw.assets, run_root=run, source_item_ids=raw.source_item_ids
    )
    moved = tmp_path / "external.png"
    moved.write_bytes(Path(image.file_path).read_bytes())
    image.file_path = str(moved)
    with pytest.raises(AssetBindingError, match="outside the active run"):
        registry.require(image.asset_id, "image")


def test_cross_document_ids_do_not_collide_for_new_and_legacy_inputs(tmp_path):
    doc = make_document()
    ids = []
    for index in (1, 2):
        source = tmp_path / f"document{index}.txt"
        source.write_text(f"source {index}")
        structural = export_docling_document(
            doc, str(source), str(tmp_path / f"extract{index}")
        )
        raw = adapt_structural_ir(structural)
        ids.append(
            (
                {block.block_id for block in raw.blocks},
                {asset.asset_id for asset in raw.assets},
            )
        )
    assert ids[0][0].isdisjoint(ids[1][0]) and ids[0][1].isdisjoint(ids[1][1])
    legacy = {"blocks": [{"block_id": "block_0001", "text": "Same text"}]}
    a = adapt_structural_ir({**legacy, "document_name": "a.pdf"})
    b = adapt_structural_ir({**legacy, "document_name": "b.pdf"})
    assert a.blocks[0].block_id != b.blocks[0].block_id


def test_legacy_adapter_preserves_unknown_provenance_and_all_usable_locations(tmp_path):
    image = tmp_path / "old.png"
    Image.new("RGB", (5, 5)).save(image)
    legacy = {
        "document_name": "old.pdf",
        "blocks": [
            {
                "block_id": "b",
                "text": "Original text",
                "prov": [{"page_no": 1}, {"page_no": 3}],
            }
        ],
        "assets": [{"asset_id": "image", "type": "image", "file_path": str(image)}],
    }
    raw = adapt_structural_ir(legacy)
    assert [location.page for location in raw.blocks[0].locations] == [1, 3]
    assert raw.source_sha256 is None
    assert raw.assets[0].source_refs == [] and raw.assets[0].findings
    with pytest.raises(AssetBindingError, match="re-extract"):
        AssetRegistry(raw.assets, source_item_ids=raw.source_item_ids).require(
            raw.assets[0].asset_id, "image"
        )


def test_table_record_migration_does_not_guess_cell_evidence():
    normalized = normalize_table_data([{"Mass (kg)": "12.50 kg", "Item": "Copper"}])
    assert normalized["rows"] == [["12.50 kg", "Copper"]]
    assert normalized["cells"] == []
    with pytest.raises(ValueError):
        normalize_table_data({"headers": ["a", "b"], "rows": [["a"]]})


def test_table_values_and_cell_membership_cannot_be_changed(source_run):
    _, raw, run = source_run
    table = next(asset for asset in raw.assets if asset.kind == "table")
    slide = _slide("table", table.asset_id)
    slide.content["body"].headers = list(table.data["headers"])
    slide.content["body"].rows = [["Copper", "999 kg"]]
    with pytest.raises(ValueError, match="differ from source"):
        validate_slide_spec(
            slide,
            _layouts("table"),
            allowed_unit_ids=set(),
            assets=raw.assets,
            source_item_ids=raw.source_item_ids,
        )
    table.data["cells"][0]["source_refs"] = ["other-document:table"]
    with pytest.raises(ValueError):
        AssetRegistry(
            raw.assets, run_root=run, source_item_ids=raw.source_item_ids
        ).require(table.asset_id, "table")


def test_source_metadata_survives_ir1_and_legacy_evaluator_projection(source_run):
    _, raw, run = source_run
    block = raw.blocks[0]
    unit = KnowledgeUnit(
        unit_id="u",
        kind="concept",
        title=block.text,
        summary=block.text,
        section_path=block.section_path,
        evidence=[EvidenceRef(source_id=block.block_id, excerpt=block.text)],
    )
    accepted, _, _ = validate_knowledge_units(raw, [unit])
    assert len(accepted) == 1
    ir1 = ValidatedIR1(
        document_title="Measurements",
        units=accepted,
        assets=[asset.model_dump() for asset in raw.assets],
        source_item_ids=raw.source_item_ids,
    )
    legacy = build_legacy_projection(ir1)
    assert [asset.model_dump() for asset in legacy.evidence_assets] == [
        asset.model_dump() for asset in raw.assets
    ]
    assert legacy.source_evidence["u"][0]["content_sha256"] == block.content_sha256
    assert legacy.source_evidence["u"][0]["locations"] == [
        location.model_dump(mode="json") for location in block.locations
    ]


def test_shared_panel_caption_recovery_is_explicit_and_conservative(tmp_path):
    doc = DoclingDocument(name="Panels")
    doc.add_page(page_no=1, size=Size(width=300, height=300))

    def panel(left, right, caption=None):
        return doc.add_picture(
            image=ImageRef.from_pil(Image.new("RGB", (40, 40)), dpi=72),
            caption=caption,
            prov=ProvenanceItem(
                page_no=1,
                bbox=BoundingBox(l=left, t=30, r=right, b=100, coord_origin="TOPLEFT"),
                charspan=(0, 0),
            ),
        )

    caption = doc.add_text(
        label=DocItemLabel.CAPTION,
        text="Figure 2: (left) Input. (right) Output.",
        prov=ProvenanceItem(
            page_no=1,
            bbox=BoundingBox(l=10, t=110, r=250, b=130, coord_origin="TOPLEFT"),
            charspan=(0, 38),
        ),
    )
    left = panel(20, 80)
    right = panel(180, 240, caption)
    source = tmp_path / "source.txt"
    source.write_text("Panels source")
    result = export_docling_document(doc, str(source), str(tmp_path / "extract"))
    left_asset = next(
        asset
        for asset in result["assets"]
        if asset["source_refs"][0].endswith(left.self_ref)
    )
    assert left_asset["caption"] == caption.text
    assert left_asset["caption_binding"] == "SHARED_PANEL_GEOMETRY"
    assert left_asset["caption_refs"][0].endswith(caption.self_ref)
    assert left_asset["findings"]
    # A third panel makes the left/right assignment ambiguous, so no evidence
    # relation is guessed for the previously uncaptioned panel.
    panel(100, 160)
    ambiguous = export_docling_document(doc, str(source), str(tmp_path / "ambiguous"))
    asset = next(
        asset
        for asset in ambiguous["assets"]
        if asset["source_refs"][0].endswith(left.self_ref)
    )
    assert not asset["caption"] and asset["caption_binding"] == "UNKNOWN"


def test_catalog_rejects_invented_reference_even_with_self_consistent_locations(
    source_run,
):
    _, raw, run = source_run
    asset = next(asset for asset in raw.assets if asset.kind == "image")
    invented = f"{asset.document_id}:#/pictures/99999"
    asset.source_refs = [invented]
    for location in asset.locations:
        location.source_id = invented
    with pytest.raises(AssetBindingError, match="source document catalog"):
        AssetRegistry(
            raw.assets, run_root=run, source_item_ids=raw.source_item_ids
        ).require(asset.asset_id, "image")


def test_staging_does_not_overwrite_assets_with_the_same_basename(tmp_path):
    entries = []
    for index, color in enumerate(("red", "blue")):
        directory = tmp_path / str(index)
        directory.mkdir()
        image = directory / "figure.png"
        Image.new("RGB", (10, 10), color).save(image)
        entries.append(
            {"asset_id": f"asset{index}", "type": "image", "file_path": str(image)}
        )
    ir = tmp_path / "ir.json"
    ir.write_text(json.dumps({"document_name": "legacy.pdf", "assets": entries}))
    raw = build_raw_extraction(ir, tmp_path / "run")
    assert raw.assets[0].file_path != raw.assets[1].file_path
    assert (
        len(
            {
                hashlib.sha256(Path(asset.file_path).read_bytes()).hexdigest()
                for asset in raw.assets
            }
        )
        == 2
    )


def test_retrieved_asset_metadata_survives_composition_repair_and_edit(source_run):
    import asyncio
    from core.artifacts import write_validated_json
    from pipeline.chunker import build_chunks, write_chunks
    from pipeline.slide_composer import (
        retrieve_slide_context,
        compose_single_slide,
        SlideCompositionError,
        write_ir2,
    )
    from pipeline.slide_editor import edit_multi_agent_slide, import_legacy_revision
    from core.revisions import active_artifacts
    from schemas.ir2_content_schema import TargetIR2Presentation
    from schemas.storyboard_schema import SlidePlan, Storyboard

    _, raw, run = source_run
    image = next(asset for asset in raw.assets if asset.kind == "image")
    block = raw.blocks[0]
    unit = KnowledgeUnit(
        unit_id="u",
        kind="concept",
        title=block.text,
        summary=block.text,
        section_path=block.section_path,
        evidence=[EvidenceRef(source_id=block.block_id, excerpt=block.text)],
        validation_status="VALIDATED",
    )
    ir1 = ValidatedIR1(
        document_title="Measurements",
        units=[unit],
        assets=[asset.model_dump() for asset in raw.assets],
        source_item_ids=raw.source_item_ids,
    )
    chunks = build_chunks(ir1)
    plan = SlidePlan(
        slide_id=1,
        title="Source measurements",
        objective="Explain measurements",
        narrative_role="opening",
        chunk_refs=[chunks[0].chunk_id],
        source_unit_ids=["u"],
        candidate_asset_ids=[image.asset_id],
        preferred_layouts=["TEST"],
    )
    storyboard = Storyboard(document_title="Measurements", slides=[plan])
    context = retrieve_slide_context(plan, ir1, chunks, storyboard)
    supplied_image = next(
        asset for asset in context["assets"] if asset["asset_id"] == image.asset_id
    )
    assert supplied_image == image.model_dump(mode="json")
    assert set(context["source_item_ids"]).issubset(raw.source_item_ids)

    class Caller:
        def __init__(self):
            self.attempts = 0

        async def generate_structured_async(self, **kwargs):
            self.attempts += 1
            return _slide("image", "invented" if self.attempts == 1 else image.asset_id)

    caller = Caller()
    slide = asyncio.run(
        compose_single_slide(plan, context, _layouts("image"), caller, max_repairs=1)
    )
    assert caller.attempts == 2
    assert slide.content["body"].source_asset_id == image.asset_id
    write_validated_json(run / "ir1_validated.json", ir1)
    write_validated_json(run / "storyboard.json", storyboard)
    write_chunks(chunks, run)
    write_ir2(TargetIR2Presentation(slides=[slide]), run)
    (run / "run_manifest.json").write_text(json.dumps({"pipeline": "multi-agent-v1"}))
    layouts_file = run / "layouts.json"
    layouts_file.write_text(json.dumps(_layouts("image")))
    builder = PresentationBuilder(
        layouts_file,
        asset_registry=AssetRegistry(
            ir1.assets, run_root=run, source_item_ids=ir1.source_item_ids
        ),
    )
    builder.build(
        {"slides": [slide.model_dump(exclude_none=True)]}, str(run / "op_ppt.pptx")
    )
    original_ir1 = (run / "ir1_validated.json").read_bytes()
    imported = import_legacy_revision(output_dir=run, layouts_path=layouts_file, caller=caller)
    result = edit_multi_agent_slide(
        output_dir=run,
        layouts_path=layouts_file,
        slide_index=1,
        instruction="Preserve the source figure",
        base_revision=imported.parent.name,
        caller=caller,
    )
    assert result.exists() and (run / "ir1_validated.json").read_bytes() == original_ir1
    assert (
        len(
            [
                shape
                for shape in Presentation(result).slides[0].shapes
                if shape.shape_type == 13
            ]
        )
        == 1
    )
    original_pptx = result.read_bytes()
    active = active_artifacts(run)
    stored = json.loads((active / "ir1_validated.json").read_text())
    active_image = next(asset for asset in stored["assets"] if asset["asset_id"] == image.asset_id)
    Path(active_image["file_path"]).unlink()
    with pytest.raises(ValueError, match="integrity failure"):
        edit_multi_agent_slide(
            output_dir=run,
            layouts_path=layouts_file,
            slide_index=1,
            instruction="Preserve the source figure",
            base_revision=active.name,
            caller=caller,
        )
    assert result.read_bytes() == original_pptx


def test_headerless_table_does_not_gain_dataframe_index_labels(tmp_path):
    doc = make_document()
    for cell in doc.tables[0].data.table_cells:
        cell.column_header = False
    source = tmp_path / "source.txt"
    source.write_text("Headerless table")
    structural = export_docling_document(doc, str(source), str(tmp_path / "extraction"))
    raw = adapt_structural_ir(structural)
    table = next(asset for asset in raw.assets if asset.kind == "table")
    assert table.data["headers"] == []
    assert table.data["rows"] == [["Material", "Mass (kg)"], ["Copper", "12.50 kg"]]
    assert all(
        block.content_sha256 == hashlib.sha256(block.text.encode()).hexdigest()
        for block in raw.blocks
    )


def test_document_identity_cannot_escape_the_active_asset_directory(tmp_path):
    source = tmp_path / "image.png"
    Image.new("RGB", (10, 10)).save(source)
    ir = tmp_path / "ir.json"
    ir.write_text(
        json.dumps(
            {
                "schema_version": "2",
                "document_id": "../../outside",
                "assets": [
                    {"asset_id": "image", "type": "image", "file_path": str(source)}
                ],
            }
        )
    )
    run = tmp_path / "run"
    raw = build_raw_extraction(ir, run)
    assert Path(raw.assets[0].file_path).resolve().is_relative_to(run.resolve())
