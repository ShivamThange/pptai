"""F13 theme chrome: decorated decks stay valid, quiet and gate-compatible.

Covers the contract the pipeline relies on: decorations never raise, never
collide with content, never pollute role-based counts the quality gate reads,
and a theme without chrome renders byte-identical geometry to the pre-chrome
behavior.
"""
import json
from pathlib import Path

import pytest
from pptx import Presentation

from application.jobs.builder import PresentationBuilder
from domain.rendering.theme_decorations import apply_theme_decorations
from domain.rendering.theme_engine import available_themes, load_theme
from renderers.base import legacy_ir2_to_scene

pytestmark = pytest.mark.filterwarnings("ignore::DeprecationWarning")

ROOT = Path(__file__).resolve().parents[1]
LAYOUTS = ROOT / "config" / "ppt_default_layouts.json"
ALL_THEMES = sorted(available_themes())


def chrome_deck():
    """One slide per chrome archetype: cover, body, section, quote, table, closing."""
    return {"title": "Chrome Deck", "slides": [
        {"layout": "TITLE_CENTERED", "content": {
            "title": {"type": "heading", "text": "Chrome Deck"},
            "subtitle": {"type": "text_group", "text": "A themed tour"}}},
        {"layout": "HEADER_WITH_BODY", "content": {
            "header": {"type": "heading", "text": "Ideas"},
            "body": {"type": "bullets", "items": ["One", "Two", "Three"]}}},
        {"layout": "SECTION_DIVIDER", "content": {
            "section_title": {"type": "heading", "text": "Methods"},
            "section_subtitle": {"type": "text_group", "text": "How it works"}}},
        {"layout": "QUOTE_CALLOUT", "content": {
            "quote": {"type": "quote", "text": "Attention is all you need."},
            "attribution": {"type": "text_group", "text": "Vaswani et al., 2017"}}},
        {"layout": "TABLE_FULL_WIDTH", "content": {
            "header": {"type": "heading", "text": "Results"},
            "table": {"type": "table", "headers": ["Model", "Score"],
                      "rows": [["A", "91.2"], ["B", "88.4"], ["C", "85.1"]]},
            "source": {"type": "text_group", "text": "Source: table 2."}}},
        {"layout": "METRIC_CALLOUT_3_GRID", "content": {
            "header": {"type": "heading", "text": "Numbers"},
            "card_1": {"type": "metric_card", "value": "28.4", "label": "BLEU"},
            "card_2": {"type": "metric_card", "value": "41.0", "label": "FR"},
            "card_3": {"type": "card", "title": "SOTA", "text": "On six benchmarks."}}},
        {"layout": "CLOSING_SLIDE", "content": {
            "title": {"type": "heading", "text": "Thank you"},
            "message": {"type": "text_group", "text": "Questions welcome."},
            "contact": {"type": "text_group", "text": "team@example.com"}}},
    ]}


def test_shipped_themes_carry_validated_chrome_and_effects():
    for theme_id in ALL_THEMES:
        theme = load_theme(theme_id)
        assert theme.chrome.title_bar and theme.chrome.footer, theme_id
        assert 8 <= theme.chrome.footer_font_pt <= 14, theme_id
        assert 0.0 <= theme.effects.radius_in <= 0.3, theme_id
        assert theme.surface and theme.surface_raised and theme.border, theme_id
        assert theme.table_header_text.startswith("#"), theme_id


@pytest.mark.parametrize("theme_id", ALL_THEMES)
def test_decorated_decks_render_without_warnings_or_failures(tmp_path, theme_id):
    builder = PresentationBuilder(LAYOUTS, backend="python", dynamic_layout=True,
                                  theme=load_theme(theme_id))
    target = tmp_path / f"chrome_{theme_id}.pptx"
    warnings = builder.build(chrome_deck(), str(target))
    assert warnings == []
    assert builder.chrome_report["failures"] == []
    # Every decoration is a real native object the coverage verifier accepted;
    # the build itself raises on any coverage mismatch, so reaching here means
    # scene objects and native shapes agree one-to-one.
    decor = [obj for slide in builder.last_scene.slides for obj in slide.objects
             if "//decor-" in obj.object_id]
    assert decor, f"{theme_id} produced no chrome"
    assert all(not obj.required for obj in decor)
    assert all(obj.kind in {"shape", "text_group"} for obj in decor)


@pytest.mark.parametrize("theme_id", ALL_THEMES)
def test_decorations_never_overlap_content_or_leave_the_slide(tmp_path, theme_id):
    builder = PresentationBuilder(LAYOUTS, backend="python", theme=load_theme(theme_id))
    builder.build(chrome_deck(), str(tmp_path / f"bounds_{theme_id}.pptx"))
    scene = builder.last_scene
    width, height = scene.width_inches, scene.height_inches
    for slide in scene.slides:
        decor = [obj for obj in slide.objects if "//decor-" in obj.object_id]
        content = [obj for obj in slide.objects if "//decor-" not in obj.object_id]
        for obj in decor:
            box = obj.box
            assert box.x >= 0 and box.y >= 0, (slide.slide_id, obj.object_id)
            assert box.x + box.w <= width + 1e-6, (slide.slide_id, obj.object_id)
            assert box.y + box.h <= height + 1e-6, (slide.slide_id, obj.object_id)
            for other in content:
                overlap_x = min(box.x + box.w, other.box.x + other.box.w) - max(box.x, other.box.x)
                overlap_y = min(box.y + box.h, other.box.y + other.box.h) - max(box.y, other.box.y)
                assert overlap_x <= 0.02 or overlap_y <= 0.02, \
                    (slide.slide_id, obj.object_id, other.object_id)


def test_chrome_report_records_applied_and_skipped(tmp_path):
    builder = PresentationBuilder(LAYOUTS, backend="python", theme=load_theme("academic_slate"))
    builder.build(chrome_deck(), str(tmp_path / "report.pptx"))
    report = builder.chrome_report
    applied = [name for names in report["applied"].values() for name in names]
    assert "title-bar" in applied and "page-number" in applied
    assert "numeral" in applied  # SECTION_DIVIDER slide
    # TABLE_FULL_WIDTH carries its own source footer: chrome must yield.
    assert any("footer: slide already carries footer content" in reason
               for reasons in report["skipped"].values() for reason in reasons)


def test_table_slide_keeps_chrome_off_and_body_slides_get_footers(tmp_path):
    builder = PresentationBuilder(LAYOUTS, backend="python", theme=load_theme("dark_tech"))
    builder.build(chrome_deck(), str(tmp_path / "dark.pptx"))
    scene = builder.last_scene
    table_slide = scene.slides[4]  # TABLE_FULL_WIDTH
    assert not any("//decor-footer" in obj.object_id for obj in table_slide.objects)
    body_slide = scene.slides[1]  # HEADER_WITH_BODY
    ids = {obj.object_id for obj in body_slide.objects}
    assert "slide-2//decor-title-bar" in ids
    assert "slide-2//decor-page-number" in ids
    assert "slide-2//decor-footer-title" in ids
    # The metric cards carry the theme's raised surface and elevation keys.
    card = next(obj for obj in scene.slides[5].objects if obj.object_id.endswith("card_1/metric-bg"))
    assert card.style["fill"] == "#24324A"
    assert card.style["shadow"] is True


def test_cover_and_closing_get_the_accent_rail(tmp_path):
    builder = PresentationBuilder(LAYOUTS, backend="python", theme=load_theme("corporate_blue"))
    builder.build(chrome_deck(), str(tmp_path / "rail.pptx"))
    scene = builder.last_scene
    for index in (0, 6):  # TITLE_CENTERED cover and CLOSING_SLIDE
        rail = next((obj for obj in scene.slides[index].objects
                     if obj.object_id.endswith("//decor-rail")), None)
        assert rail is not None, f"slide {index + 1} lacks the cover rail"
        assert rail.box.x == 0.0 and rail.box.h == scene.height_inches
    # Content slides never receive a rail.
    assert not any(obj.object_id.endswith("//decor-rail") for obj in scene.slides[1].objects)


def test_theme_without_chrome_keeps_legacy_scene_untouched(tmp_path):
    from domain.rendering.theme_engine import ThemeSpec
    legacy = ThemeSpec(id="legacy", name="Legacy", background="#FFFFFF", text="#111111",
                       text_muted="#666666", accent="#2563EB", accent_soft="#EFF6FF",
                       connector="#64748B", table_header_bg="#0F172A",
                       table_row_color="#334155", table_band="#F8FAFC",
                       chart_colors=["#2563EB"])
    layouts = json.loads(LAYOUTS.read_text(encoding="utf-8"))
    plain = legacy_ir2_to_scene(chrome_deck(), layouts)
    themed, report = apply_theme_decorations(plain, legacy)
    assert themed is plain and report.applied == {} and report.failures == []


def test_page_numbers_and_footer_survive_the_fitting_pass(tmp_path):
    builder = PresentationBuilder(LAYOUTS, backend="python", theme=load_theme("warm_editorial"))
    builder.build(chrome_deck(), str(tmp_path / "editorial.pptx"))
    scene = builder.last_scene
    page = next(obj for obj in scene.slides[1].objects
                if obj.object_id == "slide-2//decor-page-number")
    assert page.content["text"] == "2 / 7"
    assert page.role == "caption" and page.style["font_size"] == 10
    mark = next((obj for obj in scene.slides[3].objects
                 if obj.object_id.endswith("//decor-quote-mark")), None)
    assert mark is not None and mark.content["text"] == "\u201c"


def test_decorations_render_into_the_native_artifact(tmp_path):
    builder = PresentationBuilder(LAYOUTS, backend="python", theme=load_theme("minimal_mono"))
    builder.build(chrome_deck(), str(tmp_path / "mono.pptx"))
    deck = Presentation(str(tmp_path / "mono.pptx"))
    names = {shape.name for shape in deck.slides[1].shapes}
    assert "slide-2//decor-title-bar" in names
    assert "slide-2//decor-page-number" in names
    bar = next(shape for shape in deck.slides[1].shapes
               if shape.name == "slide-2//decor-title-bar")
    assert str(bar.fill.fore_color.rgb) == "111827"  # minimal_mono accent
