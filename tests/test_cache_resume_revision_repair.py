"""R05 failure/cancellation, integrity, resume and concurrency counterexamples."""
import asyncio
import json
import os
import subprocess
import sys
import threading
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from core.artifacts import (CheckpointCache, atomic_json, file_hash, parse_document_cached,
                            portable_value, sha256_value, artifact_lock)
from core.llm_client import LLMCancelledError
from core.revisions import RevisionStore, active_artifacts
from pipeline.knowledge_extractor import KnowledgeUnitBatch, KnowledgeUnitDraft, extract_knowledge_units
from pipeline.orchestrator import MultiAgentOptions, run_multi_agent_pipeline
from pipeline.slide_composer import SlideCompositionError, slide_cache_key
from pipeline.slide_editor import edit_multi_agent_slide
from pipeline.slide_renderer import RenderResult
from pipeline.source_adapter import adapt_structural_ir
from schemas.ir1_schema import EvidenceRef
from schemas.ir2_content_schema import SlideSpec
from schemas.storyboard_schema import SlidePlan, Storyboard


class FixtureCaller:
    model = "fixture-v1"
    provider = "fixture"
    def __init__(self, count=4):
        self.count, self.calls, self.fail_slide, self.cancel_slide, self.barrier = count, [], None, None, None
    def generate_structured(self, *, stage, prompt, **kwargs):
        self.calls.append(stage)
        source = json.loads(prompt)["SOURCE_DATA"]
        if stage == "knowledge_extraction":
            block = source["blocks"][0]
            return KnowledgeUnitBatch(units=[KnowledgeUnitDraft(kind="concept", title=block["text"], summary=block["text"],
                                       evidence=[EvidenceRef(source_id=block["id"], excerpt=block["text"])], priority="HIGH")])
        chunks = source["chunks"]
        return Storyboard(document_title="Attention", slides=[SlidePlan(slide_id=index, title="Attention combines values",
                   objective="Explain attention", narrative_role="opening" if index == 1 else "evidence",
                   chunk_refs=[chunks[0]["chunk_id"]], source_unit_ids=chunks[0]["unit_ids"],
                   preferred_layouts=["TITLE_CENTERED"]) for index in range(1, self.count + 1)])
    async def generate_structured_async(self, *, prompt, **kwargs):
        context = json.loads(prompt)["SOURCE_DATA"]["context"]
        index = context["plan"]["slide_id"]
        self.calls.append(index)
        if self.barrier:
            await asyncio.to_thread(self.barrier.wait, 10)
        if index == self.cancel_slide:
            raise LLMCancelledError()
        if index == self.fail_slide:
            raise RuntimeError("Injected slide failure")
        return SlideSpec(layout="TITLE_CENTERED", content={"title": {"type": "heading", "text":
                         "Attention combines values" if "user_edit_instruction" not in json.loads(prompt).get("requested_changes", {}) else "Attention uses learned weights",
                         "style": {"font_size": 32}}})


def fixture_source(root):
    source = root / "ir.json"
    atomic_json(source, {"document_name": "source.txt", "blocks": [{"block_id": "b1", "type": "text",
                 "text": "Attention combines values using learned weights.", "section": "Intro"}], "assets": []})
    return source


def render_fixture(path, output):
    root = Path(output); root.mkdir(parents=True, exist_ok=True)
    images = []
    for index in range(len(Presentation(path).slides)):
        image = root / f"slide_{index + 1:03}.png"
        Image.new("RGB", (12, 12), "white").save(image)
        images.append(image)
    return RenderResult("EVALUATED_PASS", images, artifact_hash=file_hash(path),
                        image_hashes={str(p): file_hash(p) for p in images}, text_bounds_status="EVALUATED_PASS")


@pytest.fixture
def setup_run(tmp_path, monkeypatch):
    monkeypatch.setattr("pipeline.quality_controller.render_pptx_slides", render_fixture)
    source = fixture_source(tmp_path)
    options = MultiAgentOptions(layouts_path=str(Path(__file__).parents[1] / "config/ppt_default_layouts.json"),
                                output_dir=str(tmp_path / "run"), min_slides=4, max_slides=4,
                                max_repairs_per_slide=0, max_concurrency=1)
    return source, options, FixtureCaller()


def test_failure_on_slide_four_preserves_first_three_and_resume_only_calls_four(setup_run):
    source, options, caller = setup_run
    caller.fail_slide = 4
    with pytest.raises(SlideCompositionError):
        run_multi_agent_pipeline(source, options, caller=caller)
    root = Path(options.output_dir)
    assert not (root / "active.json").exists()
    assert len(list((root / ".cache/slides").glob("*.json"))) == 3
    assert json.loads(next((root / "jobs").glob("*.json")).read_text())["status"] == "FAILED"
    caller.fail_slide = None
    before = len(caller.calls)
    final = run_multi_agent_pipeline(source, replace(options, resume=True), caller=caller)
    assert caller.calls[before:] == [4]
    assert len(Presentation(final).slides) == 4


def test_edit_survives_resume_without_generation_build_or_render(setup_run, monkeypatch):
    source, options, caller = setup_run
    original = run_multi_agent_pipeline(source, options, caller=caller)
    original_hashes = RevisionStore(options.output_dir).verify(original.parent)["artifacts"]
    updated = edit_multi_agent_slide(output_dir=options.output_dir, layouts_path=options.layouts_path, slide_index=2,
                instruction="Use the learned-weights takeaway", base_revision=original.parent.name, caller=caller)
    count = len(caller.calls)
    def forbidden(*args, **kwargs):
        raise AssertionError("Unchanged resume tried to regenerate/render")
    monkeypatch.setattr("pipeline.orchestrator.PresentationBuilder.build", forbidden)
    monkeypatch.setattr("pipeline.quality_controller.render_pptx_slides", forbidden)
    resumed = run_multi_agent_pipeline(source, replace(options, resume=True), caller=caller)
    assert resumed == updated and len(caller.calls) == count
    assert "learned weights" in Presentation(resumed).slides[1].shapes[0].text
    record = RevisionStore(options.output_dir).verify(updated.parent)
    assert record["parent_id"] == original.parent.name
    assert RevisionStore(options.output_dir).verify(original.parent)["artifacts"] == original_hashes
    old = json.loads((original.parent / "ppt_content.json").read_text())["slides"]
    new = json.loads((updated.parent / "ppt_content.json").read_text())["slides"]
    assert old[0] == new[0] and old[2:] == new[2:] and old[1] != new[1]


def test_base_revision_is_required_and_stale_edit_makes_no_calls(setup_run):
    source, options, caller = setup_run
    initial = run_multi_agent_pipeline(source, options, caller=caller)
    kwargs = dict(output_dir=options.output_dir, layouts_path=options.layouts_path, slide_index=1, instruction="Concise", caller=caller)
    with pytest.raises(ValueError, match="base_revision is required"):
        edit_multi_agent_slide(**kwargs)
    current = edit_multi_agent_slide(**kwargs, base_revision=initial.parent.name)
    before = len(caller.calls)
    with pytest.raises(ValueError, match="Stale"):
        edit_multi_agent_slide(**kwargs, base_revision=initial.parent.name)
    assert len(caller.calls) == before and active_artifacts(options.output_dir) == current.parent


def test_two_concurrent_edits_have_exactly_one_winner(setup_run):
    source, options, caller = setup_run
    initial = run_multi_agent_pipeline(source, options, caller=caller)
    caller.barrier = threading.Barrier(2)
    def edit(index):
        try:
            return edit_multi_agent_slide(output_dir=options.output_dir, layouts_path=options.layouts_path,
                   slide_index=index, instruction="Concise", base_revision=initial.parent.name, caller=caller)
        except ValueError as error:
            return error
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(edit, [1, 2]))
    assert sum(isinstance(row, Path) for row in results) == 1
    assert sum(isinstance(row, ValueError) and "changed" in str(row) for row in results) == 1
    RevisionStore(options.output_dir).resolve()


def test_process_kill_after_staging_keeps_active_pair_and_releases_lock(setup_run):
    source, options, caller = setup_run
    initial = run_multi_agent_pipeline(source, options, caller=caller)
    pointer = (Path(options.output_dir) / "active.json").read_bytes()
    script = """import os, sys
from pathlib import Path
from core.revisions import RevisionStore
store = RevisionStore(sys.argv[1])
store.commit(Path(sys.argv[2]), base_revision=sys.argv[3], operation='crash-test', failpoint=lambda *args: os._exit(73))
"""
    process = subprocess.run([sys.executable, "-c", script, options.output_dir, str(initial.parent), initial.parent.name],
                             cwd=Path(__file__).parents[1], capture_output=True, timeout=30)
    assert process.returncode == 73, process.stderr.decode()
    assert (Path(options.output_dir) / "active.json").read_bytes() == pointer
    assert active_artifacts(options.output_dir) == initial.parent
    assert len(list((Path(options.output_dir) / "revisions").iterdir())) == 2
    current = RevisionStore(options.output_dir).commit(initial.parent, base_revision=initial.parent.name, operation="retry")
    assert active_artifacts(options.output_dir) == current


@pytest.mark.parametrize("file", ["ppt_content.json", "op_ppt.pptx", "revision.json"])
def test_corrupt_committed_artifacts_cannot_be_resumed(setup_run, file):
    source, options, caller = setup_run
    initial = run_multi_agent_pipeline(source, options, caller=caller)
    (initial.parent / file).write_bytes(b"corrupted")
    with pytest.raises(ValueError, match="integrity failure"):
        run_multi_agent_pipeline(source, replace(options, resume=True), caller=caller)


def test_restore_moves_pointer_and_keeps_both_immutable_bundles(setup_run):
    source, options, caller = setup_run
    initial = run_multi_agent_pipeline(source, options, caller=caller)
    updated = edit_multi_agent_slide(output_dir=options.output_dir, layouts_path=options.layouts_path,
               slide_index=1, instruction="Concise", base_revision=initial.parent.name, caller=caller)
    store = RevisionStore(options.output_dir)
    store.restore(initial.parent.name, base_revision=updated.parent.name)
    assert store.resolve() == initial.parent and updated.exists()
    with pytest.raises(ValueError, match="changed"):
        store.restore(updated.parent.name, base_revision=updated.parent.name)


def test_cancelled_slide_preserves_checkpoints_and_records_terminal_state(setup_run):
    source, options, caller = setup_run
    caller.cancel_slide = 4
    with pytest.raises(LLMCancelledError):
        run_multi_agent_pipeline(source, options, caller=caller)
    root = Path(options.output_dir)
    assert len(list((root / ".cache/slides").glob("*.json"))) == 3
    assert json.loads(next((root / "jobs").glob("*.json")).read_text())["status"] == "CANCELLED"
    assert json.loads((root / "run_manifest.json").read_text())["status"] == "CANCELLED"


def test_recovery_records_interrupted_jobs_and_recognizes_published_job(setup_run):
    source, options, caller = setup_run
    initial = run_multi_agent_pipeline(source, options, caller=caller)
    store = RevisionStore(options.output_dir)
    committed_job = store.verify(initial.parent)["job_id"]
    atomic_json(store.job_path(committed_job), {"job_id": committed_job, "status": "RUNNING", "operation": "generation"})
    atomic_json(store.job_path("abandoned"), {"job_id": "abandoned", "status": "RUNNING", "operation": "slide_edit"})
    atomic_json(store.job_path("live"), {"job_id": "live", "status": "RUNNING", "operation": "slide_edit"})
    with artifact_lock(store.root / ".locks/edit-live.lock"):
        store.recover_jobs()
    assert json.loads(store.job_path(committed_job).read_text())["status"] == "COMPLETED"
    assert json.loads(store.job_path("abandoned").read_text())["status"] == "INTERRUPTED"
    assert json.loads(store.job_path("live").read_text())["status"] == "RUNNING"


def test_extraction_failure_retains_completed_batch(tmp_path):
    source = fixture_source(tmp_path)
    data = json.loads(source.read_text())
    data["blocks"].append({"block_id": "b2", "type": "text", "text": "Attention combines values.", "section": "Methods"})
    raw = adapt_structural_ir(data)
    caller = FixtureCaller()
    original = caller.generate_structured
    def fail(**kwargs):
        if json.loads(kwargs["prompt"])["SOURCE_DATA"]["section"] == "Methods":
            raise RuntimeError("Interrupted second batch")
        return original(**kwargs)
    caller.generate_structured = fail
    cache = CheckpointCache(tmp_path / "cache", reuse=False)
    with pytest.raises(RuntimeError):
        extract_knowledge_units(raw, caller, checkpoint_cache=cache)
    assert len(list(cache.root.glob("*.json"))) == 1
    caller.generate_structured = original
    before = len(caller.calls)
    units = extract_knowledge_units(raw, caller, checkpoint_cache=CheckpointCache(cache.root, reuse=True))
    assert len(units) == 2 and caller.calls[before:] == ["knowledge_extraction"]


def test_parser_cache_uses_bytes_versions_and_asset_hashes_without_presentation_preferences(tmp_path):
    source = tmp_path / "source.txt"; source.write_text("original")
    calls = []
    def parser(source, output, **kwargs):
        calls.append(source)
        root = Path(output); root.mkdir(parents=True, exist_ok=True)
        (root / "asset.png").write_bytes(b"pixels")
        atomic_json(root / "ir.json", {"document_name": "source.txt", "assets": []})
    first = parse_document_cached(source, tmp_path / "run", parser=parser)
    resumed = parse_document_cached(source, tmp_path / "run", resume=True, parser=parser)
    assert resumed == first and len(calls) == 1
    assert (first.parent / "source_snapshot.bin").read_bytes() == b"original"
    (first.parent / "asset.png").write_bytes(b"corrupt")
    parse_document_cached(source, tmp_path / "run", resume=True, parser=parser)
    assert len(calls) == 2
    source.write_text("changed")
    assert parse_document_cached(source, tmp_path / "run", resume=True, parser=parser) != first
    assert len(calls) == 3


@pytest.mark.parametrize("change", ["model", "provider", "layout", "prompt", "schema", "asset"])
def test_slide_cache_keys_include_actual_dependencies(setup_run, monkeypatch, change):
    _, _, caller = setup_run
    plan = SlidePlan(slide_id=1, title="Attention", objective="Explain", narrative_role="opening", chunk_refs=["chunk_001"], source_unit_ids=["ku_1"])
    context = {"assets": [{"asset_id": "a", "content_sha256": "aaa", "file_path": "C:/first"}]}
    layouts = {"theme": "blue"}
    original = slide_cache_key(plan, context, layouts, caller, 1)
    if change in {"model", "provider"}:
        setattr(caller, change, "changed")
    elif change == "layout":
        layouts["theme"] = "red"
    elif change == "asset":
        context["assets"][0]["content_sha256"] = "bbb"
    elif change == "prompt":
        monkeypatch.setattr("pipeline.slide_composer.dependency_fingerprint", lambda *args: {"actual_prompt_source": "changed"})
    else:
        monkeypatch.setattr("pipeline.slide_composer.SlideDraft.model_json_schema", lambda: {"schema": "changed"})
    assert original != slide_cache_key(plan, context, layouts, caller, 1)


def test_corrupt_checkpoint_is_a_cache_miss(tmp_path):
    cache = CheckpointCache(tmp_path, reuse=True)
    key = sha256_value("input")
    cache.put(key, {"result": "original"})
    row = json.loads((tmp_path / (key + ".json")).read_text()); row["payload"] = {"result": "tampered"}
    atomic_json(tmp_path / (key + ".json"), row)
    assert cache.get(key) is None


@pytest.mark.parametrize("pipeline_args", [[], ["--pipeline", "multi-agent"]])
def test_cli_unchanged_resume_makes_no_parser_generation_build_or_render_calls(setup_run, monkeypatch, pipeline_args):
    import main
    source, options, caller = setup_run
    def parser(source_path, output, **kwargs):
        Path(output).mkdir(parents=True, exist_ok=True)
        atomic_json(Path(output) / "ir.json", json.loads(source.read_text()))
    monkeypatch.setattr("pipeline.stage1_docling_parser.run_docling_extraction", parser)
    real_run = run_multi_agent_pipeline
    monkeypatch.setattr("pipeline.service.run_multi_agent_pipeline", lambda path, opts, **kwargs: real_run(path, opts, caller=caller,
                      **{k: v for k, v in kwargs.items() if k != "caller"}))
    cli_source = source.with_suffix(".md")
    cli_source.write_text("Attention combines values using learned weights.", encoding="utf-8")
    args = ["main.py", *pipeline_args, "--input", str(cli_source), "--layouts", options.layouts_path,
            "--output-dir", options.output_dir, "--slides-min", "4", "--slides-max", "4", "--max-repairs-per-slide", "0"]
    monkeypatch.setattr(sys, "argv", args)
    main.main()
    before = len(caller.calls)
    def forbidden(*args, **kwargs):
        raise AssertionError("Unchanged CLI resume did work")
    monkeypatch.setattr("pipeline.stage1_docling_parser.run_docling_extraction", forbidden)
    monkeypatch.setattr("pipeline.orchestrator.PresentationBuilder.build", forbidden)
    monkeypatch.setattr("pipeline.quality_controller.render_pptx_slides", forbidden)
    monkeypatch.setattr(sys, "argv", [*args, "--resume"])
    main.main()
    assert len(caller.calls) == before


def test_changed_json_after_render_cannot_be_published(setup_run):
    from core.revisions import copy_bundle
    source, options, caller = setup_run
    original = run_multi_agent_pipeline(source, options, caller=caller)
    store = RevisionStore(options.output_dir)
    work = store.root / ".work/tampered"
    copy_bundle(original.parent, work)
    data = json.loads((work / "ppt_content.json").read_text())
    data["slides"][0]["content"]["title"]["text"] = "JSON no longer matches the checked deck"
    atomic_json(work / "ppt_content.json", data)
    with pytest.raises(ValueError, match="changed after artifact assessment"):
        store.commit(work, base_revision=original.parent.name, operation="tampered")
    assert store.resolve() == original.parent


def test_edit_cache_binding_survives_changed_quality_options(setup_run):
    source, options, caller = setup_run
    initial = run_multi_agent_pipeline(source, options, caller=caller)
    updated = edit_multi_agent_slide(output_dir=options.output_dir, layouts_path=options.layouts_path, slide_index=2,
                 instruction="Use learned weights", base_revision=initial.parent.name, caller=caller)
    before = len(caller.calls)
    checked = run_multi_agent_pipeline(source, replace(options, resume=True, polish="all"), caller=caller)
    assert not any(isinstance(stage, int) for stage in caller.calls[before:])
    assert "learned weights" in Presentation(checked).slides[1].shapes[0].text
    assert updated.exists()


def test_explicit_legacy_import_rejects_forged_source_claims(setup_run, tmp_path):
    from core.revisions import copy_bundle
    from pipeline.slide_editor import import_legacy_revision
    source, options, caller = setup_run
    original = run_multi_agent_pipeline(source, options, caller=caller)
    legacy = tmp_path / "legacy"
    copy_bundle(original.parent, legacy)
    data = json.loads((legacy / "ir1_validated.json").read_text())
    data["units"][0]["summary"] = "Attention improves profit by 88%."
    atomic_json(legacy / "ir1_validated.json", data)
    before = (legacy / "ir1_validated.json").read_bytes()
    with pytest.raises(ValueError, match="source grounding"):
        import_legacy_revision(output_dir=legacy, layouts_path=options.layouts_path, caller=caller)
    assert (legacy / "ir1_validated.json").read_bytes() == before
    assert not (legacy / "active.json").exists()
