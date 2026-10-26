"""F02 source retention, routing, cache, collection and API counterexamples."""
import copy
import json
import sqlite3
import sys
from dataclasses import replace
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from infrastructure.utils.artifacts import atomic_json, file_hash, parse_document_cached
from application.jobs.job_context import JobContext
from infrastructure.persistence.jobs import JobStore, SCHEMA
from infrastructure.llm.llm_client import LLMCancelledError
from config.service_settings import ServiceSettings, get_service_settings
from api.app import create_app
from infrastructure.ingestion.detection import detect, decode_text
from infrastructure.ingestion.collection import assemble
from infrastructure.ingestion.validation import validate_ingestion
from infrastructure.ingestion.router import resolve_options
from infrastructure.ingestion.knowledge_extractor import source_passage_units, _section_batches, _assets_for_section
from infrastructure.ingestion.source_adapter import build_raw_extraction
from application.services.deck_service import parse_sources
from contracts.ingestion_schema import ParseOptions, ParseResult, IngestionError
from contracts.raw_extraction_schema import RawExtraction
from contracts.ir1_schema import EvidenceRef
from domain.evaluation.claim_verification import evidence_text
from workers.local import LocalWorker

FIXTURES = Path(__file__).parent / 'fixtures/ingestion'


def load(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def parse(name, tmp_path, **kwargs):
    return parse_document_cached(FIXTURES / name, tmp_path / name, resume=True, **kwargs)


def test_fixture_manifest_binds_actual_source_bytes():
    for name, item in load(FIXTURES / 'manifest.json')['fixtures'].items():
        assert file_hash(FIXTURES / name) == item['sha256']


@pytest.mark.parametrize('name,kind', [('source.pdf','pdf'),('source.docx','docx'),('unicode.txt','txt'),
    ('bom.txt','txt'),('source.md','markdown'),('records.csv','csv'),('source.xlsx','xlsx'),('ocr.png','png'),('ocr.jpg','jpeg')])
def test_signature_and_text_detection(name, kind):
    assert detect(FIXTURES / name)['resolved_type'] == kind


@pytest.mark.parametrize('name,selected,code', [('source.pdf','docx','TYPE_MISMATCH'),('ocr.png','pdf','TYPE_MISMATCH'),
    ('corrupt.png','auto','CORRUPT_SOURCE'),('invalid.txt','auto','CORRUPT_SOURCE'),('latin.txt','auto','TEXT_ENCODING_REQUIRED')])
def test_detection_rejects_mismatches_and_invalid_text(name, selected, code):
    with pytest.raises(IngestionError) as error:
        detect(FIXTURES / name, selected_type=selected)
    assert error.value.code == code


def test_detection_extension_authority_and_explicit_text_override(tmp_path):
    assert detect(FIXTURES / 'latin.txt', encoding='cp1252')['text_encoding'] == 'cp1252'
    value = detect(FIXTURES / 'unicode.txt', selected_type='markdown')
    assert value['resolved_type'] == 'markdown' and value['detected_type'] == 'txt'
    assert value['findings'][0].code == 'TEXT_TYPE_OVERRIDE'
    with pytest.raises(IngestionError) as error:
        detect(FIXTURES / 'source.pdf', filename='data.exe')
    assert error.value.code == 'UNSUPPORTED_TYPE'
    with pytest.raises(IngestionError) as error:
        detect(FIXTURES / 'ocr.png', filename='data.pdf')
    assert error.value.details['detected_type'] == 'png'


@pytest.mark.parametrize('name', ['unicode.txt','bom.txt','records.csv','source.xlsx'])
def test_deterministic_real_adapters_roundtrip_retention_and_member_identity(tmp_path, name):
    path = parse(name,tmp_path)
    raw = validate_ingestion(path)
    assert raw.schema_version == '3' and len(raw.documents) == 1
    assert RawExtraction.model_validate(raw.model_dump()) == raw
    assert (path.parent / 'source_snapshot.bin').read_bytes() == (FIXTURES / name).read_bytes()
    result = ParseResult.model_validate(load(path.parent / 'parse_result.json'))
    assert result.status == 'complete'
    assert all(b.document_id == 'doc_' + file_hash(FIXTURES / name) for b in raw.blocks)
    assert all(l.source_id in raw.source_item_ids for b in raw.blocks for l in b.locations)


def test_text_spans_recover_exact_unicode_and_newlines(tmp_path):
    path = parse('unicode.txt',tmp_path)
    text = (path.parent / 'decoded_source.txt').read_bytes().decode('utf-8')
    raw = validate_ingestion(path)
    for block in raw.blocks:
        start, end = block.locations[0].charspan
        assert block.text == text[start:end]
        assert block.locations[0].line_start == text.count('\n',0,start)+1
    assert '\r\n\r\n' in text and text.endswith('\n\n')


def test_csv_lexical_values_duplicate_headers_and_multiline_locations(tmp_path):
    path = parse('records.csv',tmp_path)
    raw = validate_ingestion(path)
    asset = raw.assets[0]
    assert asset.data['headers'][:2] == ['Subject','Subject']
    assert asset.data['rows'][0] == ['Alpha\nEast','001','1.2300','']
    cell = asset.data['cells'][4]
    assert cell['locations'][0]['record_index'] == 2
    assert (cell['locations'][0]['line_start'],cell['locations'][0]['line_end']) == (2,3)
    assert len(load(path.parent / 'adapter_raw.json')['records']) == 3


def test_csv_partition_header_refs_and_no_header_policy(tmp_path):
    raw = validate_ingestion(parse('many.csv',tmp_path))
    assert [len(a.data['rows']) for a in raw.assets] == [200,5]
    assert raw.assets[0].data['cells'][0]['source_refs'] == raw.assets[1].data['cells'][0]['source_refs']
    raw = validate_ingestion(parse('noheader.csv',tmp_path,parse_options={'csv_has_header':False}))
    assert not raw.assets[0].data['headers'] and raw.assets[0].data['rows'][0] == ['Alpha','001']


@pytest.mark.parametrize('name,code,options', [('ragged.csv','CSV_ROW_WIDTH',{'csv_delimiter':','}),
    ('malformed.csv','CORRUPT_SOURCE',{'csv_delimiter':','}),('headeronly.csv','NO_USABLE_CONTENT',{'csv_delimiter':','})])
def test_partial_csv_preserves_raw_and_never_publishes_cache(name, code, options, tmp_path):
    with pytest.raises(IngestionError) as error:
        parse(name,tmp_path,parse_options=options)
    assert any(f['code'] == code for f in error.value.result['findings'])
    root = tmp_path / name
    assert not list(root.rglob('parse_record.json'))
    raw = next(root.glob('incomplete/*/adapter_raw.json'))
    assert load(raw)['records']
    assert (raw.parent / 'source_snapshot.bin').read_bytes() == (FIXTURES / name).read_bytes()


def test_ambiguous_csv_requires_choice(tmp_path):
    with pytest.raises(IngestionError) as error:
        parse('ambiguous.csv',tmp_path)
    assert error.value.code == 'CSV_DELIMITER_REQUIRED'
    assert validate_ingestion(parse('ambiguous.csv',tmp_path,parse_options={'csv_delimiter':','})).assets


def test_xlsx_hidden_sparse_merged_stale_dimensions_and_formula_quarantine(tmp_path):
    path = parse('source.xlsx',tmp_path)
    raw = validate_ingestion(path)
    workbook = load(path.parent / 'adapter_raw.json')
    assert workbook['sheets'][1]['visibility'] == 'hidden'
    cells = {c['coordinate']: c for c in workbook['sheets'][0]['cells']}
    assert cells['D205']['text'] == 'Sparse tail'
    assert cells['A3']['raw_value'] == 0 and cells['C3']['number_format'] == '0%'
    assert workbook['sheets'][0]['merged_ranges'] == ['A7:C7']
    assert cells['B5']['formula'] == '=SUM(8888,9999)' and cells['B5']['text'] == ''
    assert cells['B5']['cached_state'] == 'unknown'
    assert cells['C5']['cached_state'] == 'cached_unverified' and cells['C5']['cached_value'] == 3
    units = source_passage_units(raw)
    assert all('8888' not in u.summary and '9999' not in u.summary and '=1+2' not in u.summary for u in units)
    assert len({s for s in raw.source_item_ids if s.endswith(':cell:A1')}) == 2
    assert any(f['code'] == 'UNINTERPRETED_WORKBOOK_OBJECT' for f in load(path)['findings'])


def test_xlsx_limit_lists_unprocessed_cells_and_keeps_prior_complete_cache(tmp_path):
    complete = parse('source.xlsx',tmp_path)
    before = file_hash(complete)
    with pytest.raises(IngestionError) as error:
        parse('source.xlsx',tmp_path,limits=ServiceSettings(root=tmp_path,max_workbook_cells=3))
    findings = error.value.result['findings']
    assert any(f['code'] == 'WORKBOOK_CELL_LIMIT' and f['location']['cells'] for f in findings)
    assert file_hash(complete) == before


def test_cache_hits_no_adapter_calls_and_unused_gpu_is_ignored(tmp_path, monkeypatch):
    first = parse('records.csv',tmp_path)
    def forbidden(*args, **kwargs):
        raise AssertionError('Adapter ran on a complete cache hit')
    monkeypatch.setattr('pipeline.ingestion.router.extract', forbidden)
    assert parse('records.csv',tmp_path,use_gpu=True) == first


@pytest.mark.parametrize('artifact',['source_snapshot.bin','adapter_raw.json','assets','ir.json','parse_result.json'])
def test_missing_or_tampered_retained_inputs_force_reparse(tmp_path, monkeypatch, artifact):
    path = parse('records.csv',tmp_path)
    target = next((path.parent / 'assets').rglob('*.csv')) if artifact == 'assets' else path.parent / artifact
    target.write_bytes(b'tampered')
    from pipeline.ingestion import router
    original, calls = router.extract, []
    monkeypatch.setattr(router,'extract',lambda *a,**kw: calls.append(a[0]) or original(*a,**kw))
    assert parse('records.csv',tmp_path) == path and len(calls) == 1
    validate_ingestion(path)


def test_option_and_byte_changes_get_new_keys(tmp_path):
    path = parse('records.csv',tmp_path)
    changed = parse('records.csv',tmp_path,parse_options={'csv_has_header':False})
    assert path != changed
    source = tmp_path / 'own.csv'; source.write_bytes((FIXTURES / 'records.csv').read_bytes())
    original = parse_document_cached(source,tmp_path/'root',resume=True)
    source.write_bytes(source.read_bytes().replace(b'1.2300',b'2.3400'))
    assert parse_document_cached(source,tmp_path/'root',resume=True) != original


def test_cancelled_candidate_is_retained_without_cache_record(tmp_path,monkeypatch):
    from pipeline.ingestion import router
    original = router.extract; context = JobContext()
    def cancel(*a,**kw):
        value = original(*a,**kw); context.cancel_event.set(); return value
    monkeypatch.setattr(router,'extract',cancel)
    with pytest.raises(LLMCancelledError):
        parse('unicode.txt',tmp_path,context=context)
    assert not list(tmp_path.rglob('parse_record.json'))
    assert list(tmp_path.glob('unicode.txt/incomplete/*/source_snapshot.bin'))


def test_collection_keeps_member_ids_order_and_isolates_shared_sections(tmp_path):
    members, outcomes = parse_sources([{'source_path':str(FIXTURES/'unicode.txt')},
                                      {'source_path':str(FIXTURES/'records.csv')}],tmp_path/'parse')
    aggregate = assemble(members,tmp_path/'deck')
    raw = validate_ingestion(aggregate)
    assert raw.document_id is None and len(raw.documents) == 2
    assert raw.documents[0].source_sha256 == file_hash(FIXTURES/'unicode.txt')
    assert {b.document_id for b in raw.blocks} == {d.document_id for d in raw.documents}
    batches = _section_batches(raw,100000)
    assert all(len({b.document_id for b in blocks}) == 1 for _,blocks in batches if blocks)
    assert all(all(a.document_id in section for a in _assets_for_section(raw,section)) for section,_ in batches)
    other = assemble(list(reversed(members)),tmp_path/'deck')
    assert load(other)['collection_id'] != load(aggregate)['collection_id']
    copied = build_raw_extraction(aggregate,tmp_path/'raw')
    assert copied.documents == raw.documents
    from infrastructure.ingestion.chunker import build_chunks
    from infrastructure.ingestion.knowledge_extractor import generate_validated_ir1
    ir1 = generate_validated_ir1(raw, None, tmp_path / 'ir1', extraction_mode='source-passages')
    chunks = build_chunks(ir1)
    units = {u.unit_id: u for u in ir1.units}
    for chunk in chunks:
        assert len({ref.document_id for u in chunk.unit_ids for ref in units[u].evidence}) == 1


def test_incomplete_source_set_evaluates_remaining_files_and_preserves_outcomes(tmp_path):
    outcomes = []
    with pytest.raises(IngestionError) as error:
        parse_sources([{'source_path':str(FIXTURES/'ragged.csv'),'parse_options':{'csv_delimiter':','}},
                       {'source_path':str(FIXTURES/'unicode.txt')}],tmp_path,on_outcome=lambda _,v:outcomes.append(v))
    assert error.value.code == 'SOURCE_SET_INCOMPLETE'
    assert [o['status'] for o in error.value.result['sources']] == ['partial','complete']
    assert len(outcomes) == 2 and list(tmp_path.rglob('parse_record.json'))


def test_foreign_source_and_cell_identity_are_rejected(tmp_path):
    path = parse('records.csv',tmp_path)
    data = load(path); data['assets'][0]['data']['cells'][4]['locations'][0]['document_id'] = 'foreign'
    atomic_json(path,data)
    with pytest.raises(IngestionError):
        validate_ingestion(path)
    path = parse('records.csv',tmp_path)
    raw = validate_ingestion(path)
    _,errors = evidence_text(raw,[EvidenceRef(source_id=raw.assets[0].asset_id,kind='cell',row=1,column=1,excerpt='001',document_id='foreign')])
    assert any('identity mismatch' in e for e in errors)


def test_database_migration_preserves_legacy_rows_and_backfills_members(tmp_path):
    database = tmp_path/'legacy.db'
    with sqlite3.connect(database) as db:
        db.executescript(SCHEMA)
        db.execute("INSERT INTO documents VALUES ('doc','source.md','source.md','hash',1,NULL,'time')")
        db.execute("INSERT INTO decks VALUES ('deck','doc','output','time')")
    store = JobStore(database)
    assert store.document('doc')['source_sha256'] == 'hash'
    assert store.deck_sources('deck')[0]['document_id'] == 'doc'
    assert database.with_name('legacy.db.pre-f02.bak').is_file()
    with store.connection() as db:
        assert db.execute('PRAGMA user_version').fetchone()[0] == 1
        assert len(list(db.execute("SELECT name FROM sqlite_master WHERE type='table'"))) == 4
    assert JobStore(database).deck_sources('deck') == store.deck_sources('deck')


def test_api_intake_formats_metadata_corrections_and_atomic_sets(tmp_path,monkeypatch):
    settings = ServiceSettings(root=tmp_path/'service', max_source_files=3)
    def forbidden(*args,**kwargs):
        raise AssertionError('Submission called converter/provider')
    monkeypatch.setattr('pipeline.ingestion.router.extract',forbidden)
    with TestClient(create_app(settings)) as client:
        formats = client.get('/api/input-formats').json()
        assert len(formats['formats']) == 8 and formats['limits']['upload_limit_bytes'] == 104857600
        assert client.get('/upload/').status_code == 200 and client.get('/upload/upload.js').status_code == 200
        bad = client.post('/api/documents',files={'file':('wrong.pdf',(FIXTURES/'ocr.png').read_bytes())})
        assert bad.status_code == 422 and bad.json()['detail']['detected_type'] == 'png'
        empty = client.post('/api/documents',files={'file':('empty.txt',b'')})
        assert empty.status_code == 422 and empty.json()['detail']['code'] == 'EMPTY_UPLOAD'
        undecoded = client.post('/api/documents',files={'file':('latin.txt',(FIXTURES/'latin.txt').read_bytes())})
        assert undecoded.json()['detail']['code'] == 'TEXT_ENCODING_REQUIRED'
        fixed = client.post('/api/documents',files={'file':('latin.txt',(FIXTURES/'latin.txt').read_bytes())},data={'encoding':'cp1252'})
        assert fixed.status_code == 201 and fixed.json()['upload']['text_encoding'] == 'cp1252'
        ids = [client.post('/api/documents',files={'file':(name,(FIXTURES/name).read_bytes())}).json()['id'] for name in ['unicode.txt','records.csv']]
        assert not any(client.app.state.store.document(i)['structural_ir_path'] for i in ids)
        body = {'sources':[{'document_id':i} for i in ids]}
        submitted = client.post('/api/decks',json=body)
        assert submitted.status_code == 202 and len(submitted.json()['sources']) == 2
        job = client.app.state.store.get(submitted.json()['job_id'])
        assert [s['document_id'] for s in job.request['sources']] == ids
        assert 'source_path' not in json.dumps(submitted.json())
        assert client.post('/api/decks',json={'document_id':ids[0],**body}).status_code == 422
        assert client.post('/api/decks',json={'sources':[{'document_id':ids[0]}]*2}).status_code == 422
        assert client.post('/api/decks',json={'sources':[{'document_id':ids[0]},{'document_id':'missing'}]}).status_code == 404
        with client.app.state.store.connection() as db:
            assert db.execute('SELECT count(*) FROM decks').fetchone()[0] == 1
        invalid = client.post(f'/api/documents/{ids[0]}/parse',json={'parse_options':{'ocr':'force'}})
        assert invalid.status_code == 422 and invalid.json()['detail']['code'] == 'INVALID_PARSE_OPTIONS'


def test_duplicate_bytes_collapse_with_aliases_and_conflicting_options_fail(tmp_path):
    settings=ServiceSettings(root=tmp_path)
    with TestClient(create_app(settings)) as client:
        ids=[client.post('/api/documents',files={'file':(name,(FIXTURES/'records.csv').read_bytes())}).json()['id'] for name in ['one.csv','two.csv']]
        request={'sources':[{'document_id':i} for i in ids]}
        response=client.post('/api/decks',json=request)
        assert response.status_code == 202
        assert response.json()['sources'][0]['upload_ids'] == ids and len(response.json()['sources']) == 1
        request['sources'][1]['parse_options']={'csv_has_header':False}
        response=client.post('/api/decks',json=request)
        assert response.status_code == 422 and response.json()['detail']['code'] == 'DUPLICATE_SOURCE_OPTIONS'


def test_worker_blocked_set_keeps_successful_caches_and_public_outcomes(tmp_path):
    settings=ServiceSettings(root=tmp_path,poll_interval=.01)
    with TestClient(create_app(settings)) as client:
        ids=[client.post('/api/documents',files={'file':(name,(FIXTURES/name).read_bytes())}).json()['id'] for name in ['ragged.csv','unicode.txt']]
        request={'sources':[{'document_id':ids[0],'parse_options':{'csv_delimiter':','}},{'document_id':ids[1]}]}
        job=client.post('/api/decks',json=request).json()
        worker=LocalWorker(settings,store=client.app.state.store)
        with worker.lifetime():
            assert worker.process_next()
        result=client.get(job['status_url']).json()
        assert result['status']=='failed' and result['error']['code']=='SOURCE_SET_INCOMPLETE'
        assert [s['status'] for s in result['result']['sources']] == ['partial','complete']
        assert client.get('/api/documents/'+ids[0]).json()['completeness']=='partial'
        assert client.get('/api/documents/'+ids[1]).json()['parsed']
        assert not (settings.root/'decks'/job['deck_id']/'active.json').exists()


@pytest.mark.parametrize('key',['PPTAI_MAX_SOURCE_FILES','PPTAI_MAX_SOURCE_BYTES','PPTAI_MAX_SOURCE_PAGES','PPTAI_MAX_WORKBOOK_CELLS'])
def test_configured_limits_must_be_positive(key):
    with pytest.raises(ValueError):
        get_service_settings({key:'0'})


@pytest.mark.parametrize('name,options', [('source.pdf',{'ocr':'off'}),('source.docx',{}),('source.md',{}),
    ('scanned.pdf',{'ocr':'force'}),('ocr.png',{'ocr':'force'}),('ocr.jpg',{'ocr':'force'}),('textless.png',{'ocr':'off'})])
def test_real_docling_format_adapters(tmp_path,name,options):
    path=parse(name,tmp_path,parse_options=options)
    raw=validate_ingestion(path)
    assert (path.parent/'docling_document.json').is_file()
    if name!='textless.png':
        assert any('Attention' in b.text for b in raw.blocks)
    if name.endswith(('.md','.docx')):
        assert all(l.page is None and l.bbox is None for b in raw.blocks for l in b.locations)
    if name.endswith(('.png','.jpg')):
        original=next(a for a in raw.assets if a.asset_id.endswith(':image:original'))
        assert original.content_sha256==file_hash(FIXTURES/name)
    if name.endswith('.md'):
        assert any(f['code']=='UNRESOLVED_LINKED_ASSET' for f in load(path)['findings'])


@pytest.mark.parametrize('name,code',[('corrupt.pdf','CORRUPT_SOURCE'),('encrypted.pdf','ENCRYPTED_SOURCE')])
def test_real_unreadable_pdf_failure_retains_original(tmp_path,name,code):
    with pytest.raises(IngestionError) as error:
        parse(name,tmp_path,parse_options={'ocr':'off'})
    assert error.value.code==code and error.value.result['status']=='failed'
    assert list(tmp_path.rglob('source_snapshot.bin')) and not list(tmp_path.rglob('parse_record.json'))


def test_cli_multiple_inputs_pairs_types_and_evaluates_aggregate(tmp_path,monkeypatch):
    import main
    import evaluation.reference.build_reference as reference
    calls=[]
    aggregate=tmp_path/'aggregate.json'
    def generation(sources,root,options,**kwargs):
        calls.extend(sources); kwargs['on_parsed'](aggregate); return tmp_path/'revision'/'op_ppt.pptx'
    monkeypatch.setattr(main,'generate_deck_from_sources',generation)
    monkeypatch.setattr(sys,'argv',['main.py','--input',str(FIXTURES/'unicode.txt'),'--input',str(FIXTURES/'records.csv'),
                        '--input-type','txt','--input-type','csv','--output-dir',str(tmp_path)])
    main.main()
    assert [s['selected_type'] for s in calls]==['txt','csv']
    monkeypatch.setattr(sys,'argv',['main.py','--input',str(FIXTURES/'unicode.txt'),'--input',str(FIXTURES/'records.csv'),'--input-type','txt'])
    with pytest.raises(SystemExit) as error:
        main.main()
    assert error.value.code==2


def test_partial_docling_and_required_runtime_failure_block_and_preserve_diagnostics(tmp_path, monkeypatch):
    from pipeline.ingestion import adapters
    import pipeline.ingestion.adapters.docling as adapter
    real = adapter.extract
    def partial(*args, **kwargs):
        ir = real(*args, **kwargs)
        ir['findings'].append({'code':'PARTIAL_CONVERSION','severity':'error','message':'Injected mixed-page partial conversion.',
                               'location':{'missing_pages':[2]},'blocks_generation':True})
        return ir
    monkeypatch.setattr(adapter, 'extract', partial)
    with pytest.raises(IngestionError) as error:
        parse('source.md',tmp_path)
    assert error.value.result['status']=='partial'
    incomplete = next((tmp_path/'source.md').glob('incomplete/*/ir.json'))
    validate_ingestion(incomplete,require_complete=False)
    assert error.value.result['structural_sha256'] == file_hash(incomplete)
    assert not list((tmp_path/'source.md').rglob('parse_record.json'))
    def unavailable(*args,**kwargs):
        raise IngestionError('OCR_UNAVAILABLE','Required OCR model unavailable.')
    monkeypatch.setattr('pipeline.ingestion.router.extract',unavailable)
    with pytest.raises(IngestionError) as error:
        parse('ocr.png',tmp_path)
    assert error.value.code=='OCR_UNAVAILABLE' and error.value.result['status']=='failed'


def test_parse_job_freezes_selection_independently_of_later_parse(tmp_path):
    settings=ServiceSettings(root=tmp_path,poll_interval=.01)
    with TestClient(create_app(settings)) as client:
        uploaded=client.post('/api/documents',files={'file':('source.txt',(FIXTURES/'unicode.txt').read_bytes())}).json()
        first=client.post(f"/api/documents/{uploaded['id']}/parse",json={'selected_type':'txt'}).json()
        second=client.post(f"/api/documents/{uploaded['id']}/parse",json={'selected_type':'markdown'}).json()
        store=client.app.state.store
        assert store.get(first['job_id']).request['sources'][0]['resolved_type']=='txt'
        assert store.get(second['job_id']).request['sources'][0]['resolved_type']=='markdown'
        worker=LocalWorker(settings,store=store)
        with worker.lifetime():
            worker.process_next(); worker.process_next()
        assert store.get(first['job_id']).result['parse_result']['resolved_options']['resolved_type']=='txt'
        assert store.get(second['job_id']).result['parse_result']['resolved_options']['resolved_type']=='markdown'


def test_chartsheet_is_inventory_not_a_worksheet_failure(tmp_path):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    book=Workbook(); ws=book.active; ws['A1']='Measured'; ws['A2']=10
    chart=BarChart(); chart.add_data(Reference(ws,min_col=1,min_row=2,max_row=2))
    sheet=book.create_chartsheet('Chart'); sheet.add_chart(chart)
    source=tmp_path/'chart.xlsx'; book.save(source)
    path=parse_document_cached(source,tmp_path/'parse')
    raw=load(path.parent/'adapter_raw.json')
    assert any(s['name']=='Chart' for s in raw['sheets'])
    assert any(f['code']=='UNINTERPRETED_WORKBOOK_OBJECT' for f in load(path)['findings'])


def test_progress_retains_file_outcomes_across_generation_stages():
    snapshots=[]; context=JobContext(progress_callback=lambda stage,data:snapshots.append(data))
    context.report('parsing',2,2,{'files_completed':2,'files_total':2,'outcomes':[{'status':'complete'}]})
    context.report('composition',1,4)
    assert snapshots[-1]['files_completed']==2 and snapshots[-1]['completed']==1 and snapshots[-1]['outcomes']


def test_multi_source_cli_evaluation_uses_collection_path(tmp_path,monkeypatch):
    import main
    from evalute import config, runner
    from evaluation.reference import build_reference
    aggregate=tmp_path/'collection.json'; revision=tmp_path/'revision'
    seen=[]
    def generate(sources,root,options,**kwargs):
        kwargs['on_parsed'](aggregate);return revision/'op_ppt.pptx'
    monkeypatch.setattr(main,'generate_deck_from_sources',generate)
    monkeypatch.setattr(build_reference,'build',lambda structural,ir1,out:seen.append((structural,ir1)))
    monkeypatch.setattr(config,'load',lambda path:{'faithfulness':{}})
    monkeypatch.setattr(runner,'evaluate_presentation',lambda *a,**kw:{'metrics':{'automated_overall':None,'faithfulness':{}}})
    monkeypatch.setattr(sys,'argv',['main.py','--input',str(FIXTURES/'unicode.txt'),'--input',str(FIXTURES/'records.csv'),
                                  '--evaluate','--output-dir',str(tmp_path)])
    main.main()
    assert seen==[(str(aggregate),str(revision/'content_analysis.json'))]
