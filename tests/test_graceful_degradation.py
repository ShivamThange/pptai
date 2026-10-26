"""Graceful-degradation counterexamples for the robustness overhaul (Steps 1-5).

Each test pins one failure mode the pipeline used to answer with a hard
ValueError / PIPELINE_REJECTED: an unattached section asset, overflowing text,
a throwing chart primitive, a deterministic quality failure and a slide whose
visual repair cannot be validated.
"""
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

from application.pipeline import orchestrator as orchestrator_module
from application.pipeline.orchestrator import MultiAgentOptions, run_multi_agent_pipeline
from contracts.asset_schema import AssetInterpretation, FigureAssetContent
from contracts.ir1_schema import EvidenceRef, KnowledgeUnit, ValidatedIR1, VisualAsset
from contracts.ir2_content_schema import SlideSpec
from contracts.storyboard_schema import SlidePlan, Storyboard
from infrastructure.ingestion import slide_composer
from infrastructure.ingestion.chunker import build_chunks
from infrastructure.ingestion.knowledge_extractor import KnowledgeUnitBatch, KnowledgeUnitDraft
from infrastructure.ingestion.quality_controller import (artifact_hash, assess_final_artifact,
                                                          run_deterministic_quality_checks)
from infrastructure.ingestion import quality_controller as quality_gate
from infrastructure.ingestion.slide_renderer import RenderResult
from infrastructure.ingestion.storyboard_planner import attach_asset_candidates, plan_storyboard
from infrastructure.llm.context_budget import ContextBudgetError
from infrastructure.utils.artifacts import atomic_json, file_hash
from renderers.base import Box, PresentationScene, SceneObject, SlideScene
from renderers.fitting import fit_scene
from renderers.python_adapter import PythonRenderer
from renderers.python_primitives import PythonPrimitives

ROOT = Path(__file__).resolve().parents[1]
LAYOUTS = ROOT / "config" / "ppt_default_layouts.json"
SPEC = json.loads(LAYOUTS.read_text(encoding="utf-8"))


def unit(unit_id, section, *, origin_asset_ids=(), status="VALIDATED", priority="HIGH"):
    return KnowledgeUnit(unit_id=unit_id, kind="concept", title=section[-1], summary=f"{section[-1]} overview.",
                         section_path=[section], evidence=[EvidenceRef(source_id="b1", excerpt="Attention combines values.")],
                         origin_asset_ids=list(origin_asset_ids), validation_status=status, priority=priority)


def figure(asset_id, section="Results"):
    return VisualAsset(asset_id=asset_id, kind="image", section_path=[section], caption=f"{asset_id} caption",
                       availability="AVAILABLE",
                       content=FigureAssetContent(source_sha256="a" * 64,
                                                  interpretation=AssetInterpretation(asset_id=asset_id, context_sha256="b" * 64,
                                                                                    summary=f"{asset_id} shows the comparison")))


def ir1_with(units, assets):
    return ValidatedIR1(document_title="Attention", units=units, assets=assets,
                        completeness="BASIC_CHECKED")


def test_section_asset_reaches_a_later_slide_after_the_break_fix():
    """Verification 3: the loop must continue past a slide that already has a visual."""
    ir1 = ir1_with([unit("u1", "Results"), unit("u2", "Results")], [figure("img-taken"), figure("img-free")])
    storyboard = Storyboard(document_title="Attention", lens="academic", slides=[
        SlidePlan(slide_id=1, title="First", objective="Explain", narrative_role="evidence", chunk_refs=["c1"],
                  source_unit_ids=["u1"], candidate_asset_ids=["img-taken"], preferred_layouts=["HEADER_WITH_BODY"]),
        SlidePlan(slide_id=2, title="Second", objective="Explain", narrative_role="evidence", chunk_refs=["c1"],
                  source_unit_ids=["u2"], candidate_asset_ids=[], preferred_layouts=["HEADER_WITH_BODY"])])

    attach_asset_candidates(storyboard, ir1, SPEC)

    assert storyboard.slides[0].candidate_asset_ids == ["img-taken"], "an existing visual must not be displaced"
    assert "img-free" in storyboard.slides[1].candidate_asset_ids, (
        "the fixed loop must keep scanning same-section slides instead of breaking early")
    assert storyboard.slides[1].preferred_layouts == ["HEADER_WITH_BODY"], "a layout that already takes visuals is left alone"


def test_context_budget_digest_offers_only_validated_units():
    """Step 1.2: the recovered payload must not hand the model an unvalidated unit."""
    ir1 = ir1_with([unit("u-good", "Results", priority="CRITICAL"),
                    unit("u-bad", "Results", status="REJECTED", priority="CRITICAL")], [])
    chunks = build_chunks(ir1)
    home = next(chunk for chunk in chunks if "u-good" in chunk.unit_ids)

    class BudgetCaller:
        model, provider = "fixture-v1", "fixture"

        def __init__(self):
            self.payloads = []
            self.deny_first = True

        def prepare_structured_request(self, *, prompt, **kwargs):
            self.payloads.append(json.loads(prompt))
            if self.deny_first:
                self.deny_first = False
                raise ContextBudgetError("REQUIRED_CONTEXT_TOO_LARGE", stage="storyboard_planning",
                                         required=9_000, available=100)
            return SimpleNamespace(prompt=prompt, context=SimpleNamespace(omitted_optional_refs=[]))

        def generate_prepared(self, prepared):
            return Storyboard(document_title="Attention", slides=[SlidePlan(
                slide_id=1, title="Results", objective="Explain", narrative_role="evidence",
                chunk_refs=[home.chunk_id], source_unit_ids=["u-good"], preferred_layouts=["TITLE_CENTERED"])])

    caller = BudgetCaller()
    storyboard = plan_storyboard(ir1, chunks, SPEC, caller, min_slides=1, max_slides=4)

    assert caller.deny_first is False, "the budget rejection must trigger the digest fallback"
    offered = [row["unit_id"] for row in caller.payloads[-1]["SOURCE_DATA"]["ir1"]["units"]]
    assert offered == ["u-good"], offered
    assert storyboard.slides[0].source_unit_ids == ["u-good"]


def test_overflowing_text_is_stepped_down_and_ellipsized_instead_of_crashing():
    long_line = ("Attention combines learned values across every token in the sequence and "
                 "recombines them into a single contextual representation.")
    body = SceneObject(object_id="s1/body", kind="text_group", role="body", box=Box(x=1, y=1, w=2.5, h=0.6),
                       content={"lines": [long_line]})
    scene = PresentationScene(width_inches=13.33, height_inches=7.5,
                              slides=[SlideScene(slide_id="s1", objects=[body])])

    fitted, report = fit_scene(scene)

    actions = [row.action for row in report.adjustments]
    assert "font_scaled" in actions and "text_truncated" in actions, report.as_dict()
    lines = fitted.slides[0].objects[0].content["lines"]
    assert any(line.endswith("...") for line in lines), lines
    assert all(len(line) < 200 for line in lines)


def test_residual_overflow_mangles_no_text_and_adds_no_new_build_warning():
    """The F13 counterexample: degradation must not rewrite what it cannot fix.

    A '+1.0' trend label in a shallow box is vertically short, not horizontally long,
    so ellipsizing cannot reduce its measured height; the text has to survive intact.
    Its residual overflow is recorded as information, because a heading's box height
    belongs to the layout solver and the committed F13 tests pin an empty build-warning
    list for merely tight regions. A diagram node keeps the warning it always had.
    """
    trend = SceneObject(object_id="s1/trend", kind="heading", role="label",
                        box=Box(x=1, y=1, w=3.67, h=0.05), content={"text": "+1.0"})
    node = SceneObject(object_id="s1/node", kind="shape", role="label",
                       box=Box(x=1, y=2, w=0.4, h=0.05), content={"text": "Attention layers"})
    scene = PresentationScene(width_inches=13.33, height_inches=7.5,
                              slides=[SlideScene(slide_id="s1", objects=[trend, node])])

    fitted, report = fit_scene(scene)

    assert fitted.slides[0].objects[0].content["text"] == "+1.0"
    rows = {row.object_id: row for row in report.adjustments if row.action == "kept_overflow"}
    assert set(rows) == {"s1/trend", "s1/node"}, report.as_dict()
    assert rows["s1/trend"].severity == "info" and rows["s1/node"].severity == "warning"
    assert [warning for warning in report.warnings if "s1/trend" in warning] == []


def test_failing_chart_primitive_degrades_to_a_labeled_placeholder(tmp_path, capsys):
    scene = PresentationScene(width_inches=13.33, height_inches=7.5, slides=[SlideScene(slide_id="s1", objects=[
        SceneObject(object_id="s1/plot", kind="chart", role="chart_label", box=Box(x=1, y=1, w=5, h=3),
                    content={"chart_type": "column", "categories": ["a", "b"],
                             "series": [{"name": "Revenue", "values": [1, 2]}]}),
        SceneObject(object_id="s1/heading", kind="heading", role="title", box=Box(x=1, y=0.3, w=5, h=0.6),
                    content={"text": "Attention combines values"})])])

    original = PythonPrimitives._render_region_content

    def explode(self, slide, region_def, content_obj, region_name):
        if content_obj.get("type") == "chart":
            raise ValueError("chart primitive cannot be drawn")
        return original(self, slide, region_def, content_obj, region_name)

    renderer = PythonRenderer()
    PythonPrimitives._render_region_content = explode
    try:
        artifact = renderer.render(scene, None, {}, tmp_path / "degraded.pptx")
    finally:
        PythonPrimitives._render_region_content = original

    assert any("rendered as placeholder" in line for line in artifact.warnings), artifact.warnings
    assert "[WARN] Visual render failed for s1/plot" in capsys.readouterr().out
    shapes = Presentation(tmp_path / "degraded.pptx").slides[0].shapes
    assert "[Visual Render Error: ValueError]" in {shape.text_frame.text for shape in shapes if shape.has_text_frame}
    assert [shape.name for shape in shapes] == ["s1/plot", "s1/heading"], "coverage must survive the fallback"


def test_force_accept_exports_a_deck_the_quality_gate_would_reject(tmp_path, capsys):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for _ in range(2):  # identical geometry is a hard overlap failure
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        box.text = "Readable content"
        for paragraph in box.text_frame.paragraphs:
            paragraph.font.size = Pt(18)
    path = tmp_path / "overlapping.pptx"
    prs.save(path)

    with pytest.raises(ValueError, match="quality gate failed"):
        run_deterministic_quality_checks(path, tmp_path / "strict")

    report = run_deterministic_quality_checks(path, tmp_path / "lenient", force_accept=True)

    assert report.status == "EVALUATED_PASS"
    assert report.readiness == "BASIC_CHECKED"
    assert report.deterministic.details["force_accept"] == "BYPASSED"
    assert report.warnings, "bypassed failures stay recorded for the audit"
    assert "[WARN] Quality gate bypassed:" in capsys.readouterr().out


def test_force_accept_also_bypasses_the_rendered_bounds_gate(tmp_path, monkeypatch, capsys):
    """force_accept only covered the deterministic stage, so a measured text overflow
    recorded afterwards re-blocked readiness and the deck was rejected with an empty reason."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text = "Attention weighs value vectors"
    for paragraph in box.text_frame.paragraphs:
        paragraph.font.size = Pt(18)
    path = tmp_path / "overflowing.pptx"
    prs.save(path)

    shot = tmp_path / "slide-1.png"
    Image.new("RGB", (4, 3)).save(shot)
    render = SimpleNamespace(status="EVALUATED_PASS", reason=None, artifact_hash=artifact_hash(path),
                             backend="fixture", image_paths=[shot], image_hashes={str(shot): artifact_hash(shot)},
                             text_bounds_status="EVALUATED_FAIL",
                             findings=[{"slide_no": 1, "shape_id": 4, "required": True,
                                        "check": "POWERPOINT_MEASURED_TEXT_OVERFLOW"}])
    monkeypatch.setattr(quality_gate, "render_pptx_slides", lambda *args, **kwargs: render)

    strict = assess_final_artifact(path, tmp_path / "strict")
    assert strict.readiness == "BLOCKED" and strict.rendered.status == "EVALUATED_FAIL"

    lenient = assess_final_artifact(path, tmp_path / "lenient", force_accept=True)
    assert lenient.readiness != "BLOCKED" and lenient.hard_failures == []
    assert lenient.rendered.status == "EVALUATED_PASS"
    assert lenient.rendered.details["force_accept"] == "BYPASSED"
    assert any("overflows its allocated box" in warning for warning in lenient.warnings), lenient.warnings
    assert "[WARN] Final quality gate bypassed:" in capsys.readouterr().out


class StoryboardCaller:
    """Emits a hallucinated unit id once, then a clean storyboard."""

    model, provider = "fixture-v1", "fixture"

    def __init__(self):
        self.catalog, self.repair_prompt = None, None

    def generate_structured(self, *, stage, prompt, **kwargs):
        payload = json.loads(prompt)
        chunk = payload["SOURCE_DATA"]["chunks"][0]
        self.catalog = [unit["unit_id"] for unit in chunk["units"]] if "units" in chunk else None
        real = chunk["unit_ids"]
        slides = [SlidePlan(slide_id=index, title="Attention combines values", objective="Explain attention",
                            narrative_role="opening" if index == 1 else "evidence", chunk_refs=[chunk["chunk_id"]],
                            source_unit_ids=list(real), preferred_layouts=["TITLE_CENTERED"])
                  for index in range(1, 4)]
        if "validation_errors" not in payload:
            slides[0].source_unit_ids = [*real, "ku-hallucinated"]
        else:
            self.repair_prompt = payload["validation_errors"]
        return Storyboard(document_title="Attention", slides=slides)


def test_storyboard_hallucinated_unit_ids_are_auto_repaired_once():
    ir1 = ir1_with([unit(f"u{index}", "Results") for index in range(1, 4)], [])
    chunks = build_chunks(ir1)
    caller = StoryboardCaller()

    storyboard = plan_storyboard(ir1, chunks, SPEC, caller, min_slides=1, max_slides=6)

    assert all("ku-hallucinated" not in slide.source_unit_ids for slide in storyboard.slides)
    assert caller.repair_prompt, "the retry must be told what it invented"
    assert any("hallucinated unit IDs" in row and "MUST ONLY use unit IDs" in row for row in caller.repair_prompt), caller.repair_prompt


def test_slide_two_repair_failure_keeps_the_rest_of_the_deck(capsys, tmp_path, monkeypatch):
    """Verification 2: validate_slide_spec fails for Slide 2 -> deck still exports."""
    def render_stub(pptx_path, output_dir):
        root = Path(output_dir)
        root.mkdir(parents=True, exist_ok=True)
        images = []
        for index in range(len(Presentation(pptx_path).slides)):
            image = root / f"slide_{index + 1:03d}.png"
            Image.new("RGB", (12, 12), "white").save(image)
            images.append(image)
        return RenderResult("EVALUATED_PASS", images, artifact_hash=file_hash(pptx_path),
                            image_hashes={str(p): file_hash(p) for p in images}, text_bounds_status="EVALUATED_PASS")

    monkeypatch.setattr("infrastructure.ingestion.quality_controller.render_pptx_slides", render_stub)
    monkeypatch.setattr(orchestrator_module, "collect_visual_repair_feedback",
                        lambda report, builder: {number: ["text overflows its region"] for number in range(1, 5)})

    real_validate = slide_composer.validate_slide_spec

    def validate_except_slide_two(slide, layouts_spec, **kwargs):
        if slide.asset_usage.get("repair_slide", {}).get("index") == 2:
            raise ValueError("numeric subject/unit/period binding failed")
        return real_validate(slide, layouts_spec, **kwargs)

    monkeypatch.setattr(slide_composer, "validate_slide_spec", validate_except_slide_two)

    source = tmp_path / "ir.json"
    atomic_json(source, {"document_name": "source.txt",
                         "blocks": [{"block_id": "b1", "type": "text",
                                     "text": "Attention combines values using learned weights.", "section": "Intro"}],
                         "assets": []})
    options = MultiAgentOptions(layouts_path=str(LAYOUTS), output_dir=str(tmp_path / "run"), min_slides=4, max_slides=4,
                                max_repairs_per_slide=1, max_concurrency=1)
    caller = RepairCaller()

    pptx_path = run_multi_agent_pipeline(source, options, caller=caller)

    captured = capsys.readouterr().out
    assert "[WARN] Scoped visual repair failed for slide 2" in captured, captured
    assert "Scoped visual repair failed for slide 1" not in captured

    slides = list(Presentation(pptx_path).slides)
    assert len(slides) == 4
    texts = ["\n".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame) for slide in slides]
    assert "learned weights" in texts[0] and "learned weights" in texts[2] and "learned weights" in texts[3], texts
    assert "Attention combines values" in texts[1], "slide 2 must revert to its original IR2"


class RepairCaller:
    model, provider = "fixture-v1", "fixture"

    def generate_structured(self, *, stage, prompt, **kwargs):
        source = json.loads(prompt)["SOURCE_DATA"]
        if stage == "knowledge_extraction":
            block = source["blocks"][0]
            return KnowledgeUnitBatch(units=[KnowledgeUnitDraft(kind="concept", title=block["text"], summary=block["text"],
                                                      evidence=[EvidenceRef(source_id=block["id"], excerpt=block["text"])],
                                                      priority="HIGH")])
        chunks = source["chunks"]
        return Storyboard(document_title="Attention", slides=[
            SlidePlan(slide_id=index, title="Attention combines values", objective="Explain attention",
                      narrative_role="opening" if index == 1 else "evidence", chunk_refs=[chunks[0]["chunk_id"]],
                      source_unit_ids=chunks[0]["unit_ids"], preferred_layouts=["TITLE_CENTERED"])
            for index in range(1, 5)])

    async def generate_structured_async(self, *, prompt, **kwargs):
        payload = json.loads(prompt)
        index = payload["SOURCE_DATA"]["context"]["plan"]["slide_id"]
        repairing = "visual_critic_feedback" in payload.get("requested_changes", {})
        return SlideSpec(layout="TITLE_CENTERED",
                         content={"title": {"type": "heading", "style": {"font_size": 32},
                                            "text": ("Attention uses learned weights" if repairing
                                                     else "Attention combines values")}},
                         asset_usage={"repair_slide": {"index": index}} if repairing else {})
