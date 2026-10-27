import json
from types import SimpleNamespace
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from domain.evaluation.quality_metadata import declare_metadata
from infrastructure.utils.artifacts import sha256_value
from evaluation.extract.pptx_extractor import extract
from infrastructure.ingestion.quality_controller import (run_deterministic_quality_checks, assess_final_artifact,
    artifact_hash, factual_check, resolve_readiness, adapt_legacy_quality_report, QualityCheck)
from infrastructure.ingestion.slide_renderer import RenderResult
from infrastructure.ingestion.visual_reviewer import audit_slide_images
from contracts.ir1_schema import ValidationResult
from contracts.ir2_content_schema import SlideSpec, TargetIR2Presentation


def textbox(slide, text='Readable content', left=1, top=1, width=3, height=1, size=18):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    shape.text = text
    for p in shape.text_frame.paragraphs:
        p.font.size = Pt(size)
    return shape


def deck(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    return prs, slide, tmp_path / 'deck.pptx'


def checked_facts(text='Readable content', semantic='UNKNOWN', completeness='BASIC_CHECKED'):
    slide = SlideSpec(layout='fixture', content={'body': {'type': 'text_group', 'text': text}})
    slide.grounding['body.text'] = ValidationResult(citation_exists='CHECKED', numeric_binding='NOT_APPLICABLE',
        semantic_support=semantic, status='VERIFIED' if semantic == 'CHECKED' else 'BASIC_CHECKED', claim_sha256=sha256_value(text))
    return SimpleNamespace(completeness=completeness), TargetIR2Presentation(slides=[slide])


def test_identical_text_boxes_are_hard_failures(tmp_path):
    prs, slide, path = deck(tmp_path)
    textbox(slide); textbox(slide)
    prs.save(path)
    report = run_deterministic_quality_checks(path, tmp_path, raise_on_failure=False)
    assert report.status == 'EVALUATED_FAIL'
    assert report.slides[0]['suspicious_overlap_count'] == 1
    assert report.artifact_hash == artifact_hash(path)
    with pytest.raises(ValueError, match='quality gate failed'):
        run_deterministic_quality_checks(path, tmp_path)


@pytest.mark.parametrize('declared,expected', [(False, 1), (True, 0)])
def test_card_containment_requires_relationship(tmp_path, declared, expected):
    prs, slide, path = deck(tmp_path)
    card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.8), Inches(.8), Inches(4), Inches(2))
    text = textbox(slide)
    declare_metadata(card, role='container')
    if declared:
        declare_metadata(text, role='body', container_id=card.shape_id)
    prs.save(path)
    report = run_deterministic_quality_checks(path, tmp_path, raise_on_failure=False)
    assert report.slides[0]['suspicious_overlap_count'] == expected


def test_declared_background_does_not_hide_text_pair_overlap(tmp_path):
    prs, slide, path = deck(tmp_path)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    declare_metadata(bg, role='background')
    textbox(slide); textbox(slide)
    prs.save(path)
    assert run_deterministic_quality_checks(path, tmp_path, raise_on_failure=False).slides[0]['suspicious_overlap_count'] == 1


def test_small_table_fonts_and_paragraphs_are_extracted(tmp_path):
    prs, slide, path = deck(tmp_path)
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(5), Inches(2)).table
    table.cell(0, 0).text = 'First line\nSecond line'
    table.cell(1, 0).text = 'tiny cell'
    table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.size = Pt(7)
    prs.save(path)
    shape = extract(path)['slides'][0]['shapes'][0]
    assert shape['table'][0][0] == 'First line\nSecond line'
    assert 7 in shape['font_sizes']
    report = run_deterministic_quality_checks(path, tmp_path, raise_on_failure=False)
    assert report.readiness == 'BLOCKED'
    assert report.slides[0]['font_findings'][0]['role'] == 'table'


def test_group_transform_siblings_chart_data_and_list_boundaries(tmp_path):
    prs, slide, path = deck(tmp_path)
    group = slide.shapes.add_group_shape()
    textbox(group, 'First\nSecond', left=1)
    textbox(group, left=1)
    group.left = Inches(2)
    group.width *= 2
    data = CategoryChartData()
    data.categories = ['Q1', 'Q2']
    data.add_series('Revenue', [10, 20])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(4), Inches(4), Inches(2), data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = 'Quarterly revenue'
    prs.save(path)
    shapes = extract(path)['slides'][0]['shapes']
    assert len(shapes) == 3
    assert shapes[0]['parent_id'] == group.shape_id
    assert shapes[0]['left'] == pytest.approx(2)
    assert shapes[0]['width'] == pytest.approx(6)
    assert shapes[0]['text'] == 'First\nSecond'
    assert shapes[2]['chart_data'] == [{'name': 'Revenue', 'values': [10., 20.]}]
    assert 'Q2' in shapes[2]['chart_labels']
    assert run_deterministic_quality_checks(path, tmp_path, raise_on_failure=False).slides[0]['suspicious_overlap_count'] == 1


def test_only_connected_endpoints_are_allowed_to_cross(tmp_path):
    prs, slide, path = deck(tmp_path)
    a = textbox(slide)
    b = textbox(slide, left=6)
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1), Inches(1.5), Inches(8), Inches(1.5))
    line.begin_connect(a, 3); line.end_connect(b, 1)
    middle = textbox(slide, left=4, width=1)
    prs.save(path)
    report = run_deterministic_quality_checks(path, tmp_path, raise_on_failure=False)
    assert report.slides[0]['overlap_pairs'] == [[line.shape_id, middle.shape_id]]


def test_role_specific_floor_and_unknown_fonts(tmp_path):
    prs, slide, path = deck(tmp_path)
    caption = textbox(slide, size=10)
    declare_metadata(caption, role='caption')
    body = textbox(slide, left=5, size=10)
    prs.save(path)
    findings = run_deterministic_quality_checks(path, tmp_path, raise_on_failure=False).slides[0]['font_findings']
    assert [f['shape_id'] for f in findings] == [body.shape_id]


def test_broken_archive_is_failed_to_run(tmp_path):
    path = tmp_path / 'broken.pptx'
    path.write_bytes(b'not a pptx')
    report = run_deterministic_quality_checks(path, tmp_path, raise_on_failure=False)
    assert report.deterministic.status == 'FAILED_TO_RUN'
    assert report.readiness == 'BLOCKED'


def test_contradiction_cannot_average_into_readiness(tmp_path):
    prs, slide, path = deck(tmp_path); textbox(slide); prs.save(path)
    report = run_deterministic_quality_checks(path, tmp_path, raise_on_failure=False)
    ir1, ir2 = checked_facts(semantic='FAILED')
    report.factual = factual_check(ir1, ir2, report.artifact_hash)
    report.rendered = QualityCheck(status='EVALUATED_PASS')
    report.judge = QualityCheck(status='EVALUATED_PASS')
    assert resolve_readiness(report).readiness == 'BLOCKED'


@pytest.mark.parametrize('change', ['missing', 'stale', 'deep_unknown'])
def test_incomplete_or_stale_grounding_is_not_full_assessment(change):
    ir1, ir2 = checked_facts()
    if change == 'missing':
        ir2.slides[0].grounding = {}
    elif change == 'stale':
        ir2.slides[0].content['body'].text = 'Changed factual claim'
    result = factual_check(ir1, ir2, 'hash', deep_verification=change == 'deep_unknown')
    assert result.status == 'NOT_EVALUATED'


@pytest.mark.parametrize('mode,expected', [('off', 'BASIC_CHECKED'), ('unavailable', 'DRAFT_INCOMPLETE'),
                                         ('render_fail', 'DRAFT_INCOMPLETE'), ('clip', 'BLOCKED'), ('pass', 'VERIFIED'), ('stale', 'DRAFT_INCOMPLETE')])
def test_final_assessment_statuses(tmp_path, monkeypatch, mode, expected):
    prs, slide, path = deck(tmp_path); textbox(slide); prs.save(path)
    digest = artifact_hash(path)
    ir1, ir2 = checked_facts(semantic='CHECKED', completeness='VERIFIED')
    from PIL import Image
    png = tmp_path / 'slide.png'; Image.new('RGB', (10, 10), 'white').save(png)
    hashes = {str(png): artifact_hash(png)}
    render = RenderResult('FAILED_TO_RUN' if mode == 'render_fail' else 'EVALUATED_PASS', [png], artifact_hash='old' if mode == 'stale' else digest,
                          image_hashes=hashes, findings=[{'slide_no': 1, 'shape_id': 2}] if mode == 'clip' else [], text_bounds_status='EVALUATED_PASS')
    monkeypatch.setattr('pipeline.quality_controller.render_pptx_slides', lambda *args: render)
    class Judge:
        def generate_multimodal_structured(self, **kwargs):
            if mode == 'unavailable':
                raise RuntimeError('Judge offline')
            return {'audits': [{'slide_id': 1, 'whitespace_balance_score': 8, 'text_density_score': 8,
                    'visual_hierarchy_score': 8, 'contrast_readability_score': 8, 'passed': True, 'critique_summary': 'Readable'}]}
    report = assess_final_artifact(path, tmp_path, ir1=ir1, ir2=ir2, caller=Judge(), judge_enabled=mode != 'off', deep_verification=mode == 'pass')
    assert report.readiness == expected
    assert report.artifact_hash == digest
    if mode == 'unavailable':
        assert report.judge.status == 'FAILED_TO_RUN'
        assert report.factual.status == 'EVALUATED_PASS'


def test_vision_missing_and_duplicate_ids_fail_to_run():
    class Judge:
        def generate_multimodal_structured(self, **kwargs):
            row = {'slide_id': 1, 'whitespace_balance_score': 8, 'text_density_score': 8,
                   'visual_hierarchy_score': 8, 'contrast_readability_score': 8, 'passed': True, 'critique_summary': 'ok'}
            return {'audits': [row, row]}
    assert audit_slide_images([], Judge()).status == 'FAILED_TO_RUN'
    assert audit_slide_images(['a'], Judge()).status == 'FAILED_TO_RUN'
    assert audit_slide_images(['a'], Judge(), expected_slide_count=2).status == 'FAILED_TO_RUN'


def test_legacy_success_is_rechecked_and_original_preserved(tmp_path):
    prs, slide, path = deck(tmp_path); textbox(slide); textbox(slide); prs.save(path)
    legacy = {'status': 'PASSED', 'visual': {'score': 1}}
    result = adapt_legacy_quality_report(legacy, path, tmp_path)
    assert result.readiness == 'BLOCKED'
    assert legacy['status'] == 'PASSED'
    prs.slides[0].shapes[0].text = 'Changed'; prs.save(path)
    with pytest.raises(ValueError, match='different artifact'):
        adapt_legacy_quality_report(result.model_dump(), path, tmp_path)


def pipeline_fixture(tmp_path):
    import hashlib
    from pathlib import Path
    from infrastructure.ingestion.knowledge_extractor import KnowledgeUnitBatch, KnowledgeUnitDraft
    from infrastructure.ingestion.source_adapter import adapt_structural_ir
    from contracts.ir1_schema import EvidenceRef
    from contracts.storyboard_schema import SlidePlan, Storyboard
    structural = {'document_name': 'source.txt', 'blocks': [{'block_id': 'b1', 'type': 'text',
                  'text': 'Attention combines values using learned weights.', 'section': 'Intro'}], 'assets': []}
    path = tmp_path / 'ir.json'; path.write_text(json.dumps(structural))
    source_id = adapt_structural_ir(structural).blocks[0].block_id
    unit_id = 'ku_' + hashlib.sha256(f'concept|attention combines values using learned weights.|{source_id}'.encode()).hexdigest()[:16]
    class Caller:
        compositions = 0
        audits = 0
        fail_judge = False
        def generate_structured(self, *, stage, **kwargs):
            if stage == 'knowledge_extraction':
                return KnowledgeUnitBatch(units=[KnowledgeUnitDraft(kind='concept', title='Attention',
                    summary=structural['blocks'][0]['text'], evidence=[EvidenceRef(source_id=source_id, excerpt=structural['blocks'][0]['text'])], priority='HIGH')])
            return Storyboard(document_title='Attention', slides=[SlidePlan(slide_id=1, title='Attention combines values',
                objective='Explain attention', narrative_role='opening', chunk_refs=[json.loads(kwargs['prompt'])['SOURCE_DATA']['chunks'][0]['chunk_id']],
                source_unit_ids=[unit_id], preferred_layouts=['TITLE_CENTERED'])])
        async def generate_structured_async(self, **kwargs):
            self.compositions += 1
            return SlideSpec(layout='TITLE_CENTERED', content={'title': {'type': 'heading', 'text': 'Attention combines values',
                             'style': {'font_size': 32 if self.compositions == 1 else 28}}})
        def generate_multimodal_structured(self, **kwargs):
            self.audits += 1
            return {'audits': [{'slide_id': 1, 'whitespace_balance_score': 8, 'text_density_score': 8,
                'visual_hierarchy_score': 8, 'contrast_readability_score': 8, 'passed': self.audits > 1 and not self.fail_judge,
                'critique_summary': 'Readable' if self.audits > 1 else 'Needs repair', 'actionable_feedback': 'Use a smaller heading font.'}]}
    return path, Path(__file__).parents[1] / 'config/ppt_default_layouts.json', Caller()


def rendered_fixture(path, out):
    from pathlib import Path
    from PIL import Image
    root = Path(out); root.mkdir(parents=True, exist_ok=True)
    image = root / 'slide_001.png'; Image.new('RGB', (10, 10), 'white').save(image)
    return RenderResult('EVALUATED_PASS', [image], artifact_hash=artifact_hash(path),
                        image_hashes={str(image): artifact_hash(image)}, text_bounds_status='EVALUATED_PASS')


@pytest.mark.parametrize('final_fail', [False, True])
def test_scoped_repair_rerenders_and_uses_final_audit(tmp_path, monkeypatch, final_fail):
    from application.pipeline.orchestrator import run_multi_agent_pipeline, MultiAgentOptions
    source, layouts, caller = pipeline_fixture(tmp_path)
    caller.fail_judge = final_fail
    rendered = []
    def render(path, out):
        result = rendered_fixture(path, out)
        rendered.append(result.artifact_hash)
        return result
    monkeypatch.setattr('pipeline.quality_controller.render_pptx_slides', render)
    root = tmp_path / 'run'
    options = MultiAgentOptions(layouts_path=str(layouts), output_dir=str(root), min_slides=1, max_slides=1, polish='all')
    if final_fail:
        with pytest.raises(ValueError, match='Final PPTX quality gate failed'):
            run_multi_agent_pipeline(source, options, caller=caller)
        assert not (root / 'op_ppt.pptx').exists()
    else:
        final = run_multi_agent_pipeline(source, options, caller=caller)
        report = json.loads((final.parent / 'quality_report.json').read_text())
        manifest = json.loads((final.parent / 'run_manifest.json').read_text())
        assert report['artifact_hash'] == rendered[1] == artifact_hash(final)
        assert report['judge']['status'] == 'EVALUATED_PASS'
        assert manifest['status'] == 'BASIC_CHECKED'
        assert manifest['artifact_hash'] == rendered[1]
    assert caller.audits == 2
    assert len(rendered) == 2 and rendered[0] != rendered[1]
    work = next((root / '.work').iterdir())
    initial = json.loads((work / 'quality_attempts' / rendered[0] / 'quality_report.json').read_text())
    assert initial['judge']['status'] == 'EVALUATED_FAIL'


@pytest.mark.parametrize('reason', ['render', 'judge', 'geometry'])
def test_failed_edit_preserves_previous_pptx_ir2_and_quality_report(tmp_path, monkeypatch, reason):
    from application.pipeline.orchestrator import run_multi_agent_pipeline, MultiAgentOptions
    from application.services.slide_editor import edit_multi_agent_slide
    source, layouts, caller = pipeline_fixture(tmp_path)
    root = tmp_path / 'run'
    monkeypatch.setattr('pipeline.quality_controller.render_pptx_slides', rendered_fixture)
    final = run_multi_agent_pipeline(source, MultiAgentOptions(layouts_path=str(layouts), output_dir=str(root), min_slides=1, max_slides=1,
                                      polish='all' if reason == 'judge' else 'off'), caller=caller)
    revision = final.parent
    previous = {name: (revision / name).read_bytes() for name in ['op_ppt.pptx', 'ppt_content.json', 'quality_report.json']}
    if reason == 'render':
        monkeypatch.setattr('pipeline.quality_controller.render_pptx_slides', lambda path, out:
            RenderResult('FAILED_TO_RUN', [], artifact_hash=artifact_hash(path), reason='Unavailable'))
    elif reason == 'judge':
        caller.fail_judge = True
    else:
        async def tiny(**kwargs):
            return SlideSpec(layout='TITLE_CENTERED', content={'title': {'type': 'heading', 'text': 'Attention combines values', 'style': {'font_size': 8}}})
        caller.generate_structured_async = tiny
    with pytest.raises(ValueError):
        edit_multi_agent_slide(output_dir=root, layouts_path=layouts, slide_index=1, instruction='Refine typography',
                               base_revision=revision.name, caller=caller)
    from infrastructure.persistence.revisions import active_artifacts
    assert active_artifacts(root) == revision
    assert all((revision / name).read_bytes() == data for name, data in previous.items())


def test_native_render_worker_failure_is_explicit(tmp_path, monkeypatch):
    from infrastructure.ingestion.slide_renderer import render_pptx_slides
    import subprocess
    prs, slide, path = deck(tmp_path); textbox(slide); prs.save(path)
    def failed(*args, **kwargs):
        raise subprocess.TimeoutExpired('worker', 60)
    monkeypatch.setattr('pipeline.slide_renderer.subprocess.run', failed)
    result = render_pptx_slides(path, tmp_path / 'render')
    assert result.status == 'FAILED_TO_RUN'
    assert result.artifact_hash == artifact_hash(path)


def test_partial_factual_judge_failure_retains_contradiction(monkeypatch):
    import numpy as np
    from evaluation.config import load
    from evaluation.metrics.faithfulness import evaluate, ClaimVerdict
    monkeypatch.setattr('evalute.metrics.faithfulness.similarities', lambda *a: (np.array([[1.0], [1.0]]), 'fixture'))
    class Judge:
        calls = 0
        def generate_structured(self, **kwargs):
            self.calls += 1
            if self.calls > 1:
                raise RuntimeError('Judge unavailable')
            return ClaimVerdict(label='CONTRADICTED', explanation='Wrong figure')
    records, summary = evaluate([{'slide_no': 1, 'text': 'Revenue was 10 USD. Profit was 20 USD.'}],
                               [{'chunk_id': 'source', 'text': 'Revenue was 10 USD. Profit was 2 USD.'}], load(), caller=Judge())
    assert len(records) == 1 and records[0]['label'] == 'CONTRADICTED'
    assert summary['status'] == 'FAILED' and summary['score'] == 0.0
    assert summary['evaluated_claim_count'] == 1 and summary['claim_count'] == 2
    assert summary['assessment_completeness'] == 0.5 and summary['interval'] == [0.0, 0.5]


def test_evaluator_missing_judging_does_not_create_complete_overall(tmp_path, monkeypatch):
    from evaluation.config import load
    from evaluation.runner import evaluate_presentation
    from evaluation.reference.build_reference import build
    prs, slide, path = deck(tmp_path); textbox(slide, text='Attention combines values using learned weights.'); prs.save(path)
    ir = tmp_path / 'source.json'; ir.write_text(json.dumps({'document_name': 'source.txt', 'blocks': [{'block_id': 'b1', 'text': 'Attention combines values using learned weights.'}]}))
    ir1 = tmp_path / 'ir1.json'; ir1.write_text(json.dumps({'document_title': 'Attention', 'concepts': [{'title': 'Attention', 'summary': 'Attention combines values using learned weights.', 'source_block_ids': ['b1']}]}))
    ref = tmp_path / 'ref.json'; build(ir, ir1, ref)
    report = evaluate_presentation(ref, path, 'fixture', 'run', tmp_path / 'result', load(), ai_judging=False)
    assert report['quality_gate']['factual_status'] == 'NOT_EVALUATED'
    assert report['metrics']['automated_overall'] is None
    assert report['metrics']['available_metrics_score'] is not None
    assert report['readiness'] == 'DRAFT_INCOMPLETE'


def test_rotated_group_does_not_claim_completed_native_geometry(tmp_path):
    prs, slide, path = deck(tmp_path)
    group = slide.shapes.add_group_shape(); textbox(group); group.rotation = 30
    prs.save(path)
    report = run_deterministic_quality_checks(path, tmp_path, raise_on_failure=False)
    assert report.deterministic.status == 'NOT_EVALUATED'
    assert report.slides[0]['unknown_geometry_shape_ids']
    assert report.readiness == 'DRAFT_INCOMPLETE'


def test_unavailable_deep_judging_is_failed_to_run():
    ir1, ir2 = checked_facts()
    ir2.slides[0].grounding['body.text'].reason.append('Semantic judge unavailable (auth); support remains unknown.')
    assert factual_check(ir1, ir2, 'hash', deep_verification=True).status == 'FAILED_TO_RUN'


def test_foreground_background_cannot_hide_required_text(tmp_path):
    prs, slide, path = deck(tmp_path)
    textbox(slide)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    declare_metadata(bg, role='background')
    prs.save(path)
    assert run_deterministic_quality_checks(path, tmp_path, raise_on_failure=False).readiness == 'BLOCKED'


def test_missing_presentation_has_explicit_unavailable_hash(tmp_path):
    report = run_deterministic_quality_checks(tmp_path / 'absent.pptx', tmp_path, raise_on_failure=False)
    assert report.status == 'FAILED_TO_RUN'
    assert report.artifact_hash is None
    assert report.revision_id is None
