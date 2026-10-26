"""R06 actual-package parity plus malformed-content/backend counterexamples."""
import copy
import json
import subprocess
from dataclasses import replace
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from infrastructure.utils.artifacts import file_hash
from application.jobs.builder import PresentationBuilder
from infrastructure.persistence.revisions import RevisionStore
from domain.evaluation.quality_metadata import read_metadata
from renderers.base import PresentationScene, legacy_ir2_to_scene
from renderers.factory import get_renderer
from application.pipeline.orchestrator import run_multi_agent_pipeline
from application.services.slide_editor import edit_multi_agent_slide
from test_cache_resume_revision_repair import setup_run
from test_provenance_assets import source_run


def catalog():
    types = ["heading", "text_group", "bullets", "numbered_list", "table", "chart", "image", "formula", "diagram", "stat"]
    return {"slide_dimensions": {"width_inches": 13.333, "height_inches": 7.5}, "layouts": {"fixture": {"regions": {
        "title": {"left": .06, "top": .05, "width": .88, "height": .1, "allowed_types": ["heading"]},
        "body": {"left": .08, "top": .25, "width": .84, "height": .6, "allowed_types": types}}}}}


def deck(content):
    return {"slides": [{"layout": "fixture", "notes": "Retained speaker notes", "content": {"body": content}}]}


@pytest.mark.parametrize("backend", ["python", "pptxgenjs"])
@pytest.mark.parametrize("content", [
    {"type": "table", "headers": ["Subject", "Value"], "rows": [["Profit", "-12"], ["Loss", "0"]]},
    {"type": "chart", "chart_type": "bar", "categories": ["A", "B"], "series": [{"name": "Result", "values": [-3, 5]}]},
    {"type": "chart", "chart_type": "column", "categories": ["A", "B"], "series": [{"name": "Result", "values": [-3, 0]}]},
    {"type": "bullets", "items": ["First", {"text": "Second", "level": 1}]},
    {"type": "numbered_list", "items": ["First", "Second"]},
    {"type": "diagram", "nodes": [{"id": "a", "label": "Source"}, {"id": "b", "label": "Target"}], "edges": [{"from_id": "b", "to_id": "a"}]},
    {"type": "stat", "items": [{"value": "0", "label": "Count"}]},
])
def test_actual_native_objects_values_notes_and_roles(tmp_path, backend, content):
    original = deck(content); snapshot = copy.deepcopy(original)
    builder = PresentationBuilder(catalog(), backend=backend)
    target = tmp_path / "deck.pptx"
    builder.build(original, target)
    assert original == snapshot
    assert builder.last_result.artifact_hash == file_hash(target)
    slide = Presentation(target).slides[0]
    assert "Retained speaker notes" in slide.notes_slide.notes_text_frame.text
    assert {s.name for s in slide.shapes} == {o.object_id for o in builder.last_scene.slides[0].objects}
    assert all(read_metadata(s)["required"] for s in slide.shapes)
    if content["type"] == "chart":
        chart = next(s.chart for s in slide.shapes if s.has_chart)
        direction = chart._chartSpace.xpath(".//c:barDir")[0].get("val")
        assert direction == ("bar" if content["chart_type"] == "bar" else "col")
        assert not chart._chartSpace.xpath(".//c:scaling/c:min[@val='0']")
    if content["type"] == "diagram":
        connector = next(s for s in slide.shapes if s._element.tag.endswith("cxnSp"))
        assert len(connector._element.xpath(".//a:stCxn | .//a:endCxn")) == 2


@pytest.mark.parametrize("backend", ["python", "pptxgenjs"])
@pytest.mark.parametrize("kind", ["bullets", "table"])
def test_dense_native_content_preserves_resolved_font_size(tmp_path, backend, kind):
    content = ({"type": "bullets", "items": ["Source-backed explanation " * 4] * 7} if kind == "bullets" else
               {"type": "table", "headers": [f"H{i}" for i in range(6)], "rows": [["x"] * 6 for _ in range(9)]})
    content["style"] = {"font_size": 16}
    target = tmp_path / "dense.pptx"
    builder = PresentationBuilder(catalog(), backend=backend)
    builder.build(deck(content), target)
    assert builder.last_scene.slides[0].objects[0].style["font_size"] == 16
    shape = Presentation(target).slides[0].shapes[0]
    frames = [cell.text_frame for row in shape.table.rows for cell in row.cells] if shape.has_table else [shape.text_frame]
    assert {run.font.size.pt for frame in frames for paragraph in frame.paragraphs for run in paragraph.runs} == {16}


@pytest.mark.parametrize("height,out_of_bounds", [(80, False), (120, True)])
def test_native_office_expansion_is_checked_against_scene_and_page(height, out_of_bounds):
    from types import SimpleNamespace
    from infrastructure.ingestion.slide_renderer import _text_bounds
    shape = SimpleNamespace(Type=19, Id=3, Name="table", AlternativeText="", Left=0, Top=0,
                            Width=100, Height=height, HasTable=False, HasTextFrame=False)
    slide = SimpleNamespace(Shapes=SimpleNamespace(Count=1, Item=lambda index: shape))
    findings, errors = _text_bounds(slide, 1, page_size=(100, 100), expected_boxes={"table": (0, 0, 100, 50)})
    assert not errors
    checks = {finding["check"] for finding in findings}
    assert "POWERPOINT_MEASURED_GEOMETRY_CHANGE" in checks
    assert ("POWERPOINT_MEASURED_OUT_OF_BOUNDS" in checks) == out_of_bounds


@pytest.mark.parametrize("backend", ["python", "pptxgenjs"])
@pytest.mark.parametrize("bad", [
    {"type": "video", "text": "unsupported"},
    {"type": "table", "headers": ["A", "B"], "rows": [["one"]]},
    {"type": "chart", "categories": ["A", "B"], "series": [{"name": "s", "values": [1]}]},
    {"type": "chart", "categories": ["A"], "series": [{"name": "s", "values": [float("nan")]}]},
    {"type": "diagram", "nodes": [{"id": "a", "label": "A"}], "edges": [{"from_id": "a", "to_id": "missing"}]},
    {"type": "diagram", "caption": "No native structure"},
    {"type": "image", "path": "missing.png"},
])
def test_invalid_required_object_preserves_previous_output(tmp_path, backend, bad):
    target = tmp_path / "prior.pptx"; target.write_bytes(b"prior accepted bytes")
    with pytest.raises(ValueError):
        PresentationBuilder(catalog(), backend=backend).build(deck(bad), target)
    assert target.read_bytes() == b"prior accepted bytes"


@pytest.mark.parametrize("backend", ["python", "pptxgenjs"])
def test_real_image_formula_hash_and_resolved_solver_boxes(tmp_path, backend):
    image = tmp_path / "source.png"; Image.new("RGB", (200, 100), "blue").save(image)
    scene = legacy_ir2_to_scene(deck({"type": "image", "path": str(image), "caption": "Source image"}), catalog())
    obj = scene.slides[0].objects[0]
    obj.box.x, obj.box.y, obj.box.w, obj.box.h = 2, 3, 4, 2
    scene.slides[0].objects = [obj]
    for kind in ["image", "formula"]:
        obj.kind = kind
        path = tmp_path / f"{kind}.pptx"
        get_renderer(backend).render(scene, None, {}, path)
        shape = Presentation(path).slides[0].shapes[0]
        assert (shape.left / 914400, shape.top / 914400, shape.width / 914400, shape.height / 914400) == (2, 3, 4, 2)
        assert file_hash(image) == obj.content["content_sha256"]
    image.write_bytes(b"corrupted after adaptation")
    with pytest.raises(ValueError, match="hash changed"):
        get_renderer(backend).render(scene, None, {}, tmp_path / "rejected.pptx")


def test_direct_scene_rejects_unversioned_and_unsupported_objects():
    scene = legacy_ir2_to_scene(deck({"type": "heading", "text": "Facts"}), catalog()).model_dump()
    scene["schema_version"] = "scene-v0"
    with pytest.raises(ValueError): PresentationScene.model_validate(scene)
    scene["schema_version"] = "scene-v1"; scene["slides"][0]["objects"][0]["kind"] = "video"
    with pytest.raises(ValueError): PresentationScene.model_validate(scene)


def test_node_failure_does_not_fallback_or_replace_artifact(tmp_path, monkeypatch):
    from renderers import pptxgenjs_adapter
    scene = legacy_ir2_to_scene(deck({"type": "heading", "text": "Facts"}), catalog())
    target = tmp_path / "deck.pptx"; target.write_bytes(b"accepted")
    monkeypatch.setattr(pptxgenjs_adapter.shutil, "which", lambda name: None)
    with pytest.raises(RuntimeError, match="requires Node"):
        get_renderer("pptxgenjs").render(scene, None, {}, target)
    assert target.read_bytes() == b"accepted"


def test_node_deadline_and_version_mismatch_preserve_prior(tmp_path, monkeypatch):
    from renderers import pptxgenjs_adapter
    scene = legacy_ir2_to_scene(deck({"type": "heading", "text": "Facts"}), catalog())
    target = tmp_path / "deck.pptx"; target.write_bytes(b"accepted")
    def timed_out(*args, **kwargs): raise subprocess.TimeoutExpired("node", 60)
    monkeypatch.setattr(pptxgenjs_adapter.subprocess, "run", timed_out)
    with pytest.raises(RuntimeError, match="deadline"):
        get_renderer("pptxgenjs").render(scene, None, {}, target)
    monkeypatch.setattr(pptxgenjs_adapter.subprocess, "run", lambda *a, **k: subprocess.CompletedProcess("node", 0, '{"version":"4.0.0"}', ""))
    with pytest.raises(RuntimeError, match="locked acceptance baseline"):
        get_renderer("pptxgenjs").render(scene, None, {}, target)
    assert target.read_bytes() == b"accepted"


def test_node_cli_uses_canonical_validation(tmp_path):
    scene = legacy_ir2_to_scene(deck({"type": "heading", "text": "Portable scene"}), catalog())
    request = tmp_path / "scene.json"; request.write_text(scene.model_dump_json(), encoding="utf-8")
    target = tmp_path / "cli.pptx"
    root = Path(__file__).parents[1]
    process = subprocess.run(["node", str(root / "renderers/pptxgenjs/index.js"), str(request), str(target)],
                             cwd=root, capture_output=True, text=True, timeout=60)
    assert process.returncode == 0, process.stderr
    assert json.loads(process.stdout)["backend"] == "pptxgenjs"
    bad = scene.model_dump(); bad["slides"][0]["objects"][0]["kind"] = "unsupported"
    request.write_text(json.dumps(bad), encoding="utf-8")
    prior = file_hash(target)
    process = subprocess.run(["node", str(root / "renderers/pptxgenjs/index.js"), str(request), str(target)],
                             cwd=root, capture_output=True, text=True, timeout=60)
    assert process.returncode != 0 and file_hash(target) == prior


def test_js_generation_edit_resume_keep_backend_and_scene(setup_run, monkeypatch):
    # Native rendering is separately checked by verify_r06_native.py. This
    # scoped regression isolates revision/cache behavior and uses real JS builds.
    source, options, caller = setup_run
    options = replace(options, renderer="pptxgenjs")
    initial = run_multi_agent_pipeline(source, options, caller=caller)
    store = RevisionStore(options.output_dir)
    record = store.verify(initial.parent)
    assert record["renderer"]["backend"] == "pptxgenjs"
    assert record["renderer"]["version"] == "3.12.0"
    assert "scene_graph.json" in record["artifacts"]
    with pytest.raises(ValueError, match="retains the revision renderer"):
        edit_multi_agent_slide(output_dir=options.output_dir, layouts_path=options.layouts_path, slide_index=2,
            instruction="Use learned weights", base_revision=initial.parent.name, caller=caller, renderer="python")
    updated = edit_multi_agent_slide(output_dir=options.output_dir, layouts_path=options.layouts_path, slide_index=2,
        instruction="Use learned weights", base_revision=initial.parent.name, caller=caller)
    assert store.verify(updated.parent)["renderer"]["backend"] == "pptxgenjs"
    before = len(caller.calls)
    assert run_multi_agent_pipeline(source, replace(options, resume=True), caller=caller) == updated
    assert len(caller.calls) == before


def test_table_cell_com_bounds_do_not_call_unimplemented_properties():
    from types import SimpleNamespace
    from infrastructure.ingestion.slide_renderer import _text_bounds
    class Cell:
        @property
        def Type(self): raise RuntimeError("Office table cell Shape.Type is E_NOTIMPL")
        @property
        def AlternativeText(self): raise RuntimeError("Office table cell AlternativeText is E_NOTIMPL")
        HasTable = False; HasTextFrame = True; Width = 100; Height = 30
        TextFrame2 = SimpleNamespace(HasText=True, MarginLeft=0, MarginRight=0, MarginTop=0, MarginBottom=0,
                                    TextRange=SimpleNamespace(BoundWidth=150, BoundHeight=20))
    table = SimpleNamespace(Type=19, HasTable=True, Id=7, AlternativeText="",
                            Table=SimpleNamespace(Rows=SimpleNamespace(Count=1), Columns=SimpleNamespace(Count=1),
                                                  Cell=lambda r,c: SimpleNamespace(Shape=Cell())))
    slide = SimpleNamespace(Shapes=SimpleNamespace(Count=1, Item=lambda i:table))
    findings, errors = _text_bounds(slide, 1)
    assert not errors and findings[0]["shape_id"] == 7 and findings[0]["required"]


@pytest.mark.parametrize("kind", ["image", "formula", "table"])
def test_js_registry_assets_are_real_and_provenance_is_required(source_run, kind):
    _, raw, root = source_run
    from domain.extraction.asset_registry import AssetRegistry
    asset = next(a for a in raw.assets if a.kind == kind)
    registry = AssetRegistry(raw.assets, run_root=root, source_item_ids=raw.source_item_ids)
    builder = PresentationBuilder(catalog(), backend="pptxgenjs", asset_registry=registry)
    target = root / "registry.pptx"
    builder.build(deck({"type": kind, "source_asset_id": asset.asset_id}), target)
    prior = file_hash(target)
    asset.source_refs = []
    with pytest.raises(ValueError, match="source membership"):
        builder.build(deck({"type": kind, "source_asset_id": asset.asset_id}), target)
    assert file_hash(target) == prior


def test_office_literal_series_name_is_retained_by_extractor(tmp_path):
    from pptx.oxml.xmlchemy import OxmlElement
    from domain.rendering.chart_content import native_series_name
    from evaluation.extract.pptx_extractor import extract
    path = tmp_path / "literal.pptx"
    PresentationBuilder(catalog()).build(deck({"type":"chart", "categories":["A"],
                                      "series":[{"name":"Original", "values":[1]}]}), path)
    presentation = Presentation(path); series = presentation.slides[0].shapes[0].chart.series[0]
    tx = series._element.xpath("./c:tx")[0]
    for child in list(tx): tx.remove(child)
    literal = OxmlElement("c:v"); literal.text = "Edited native series"; tx.append(literal)
    presentation.save(path)
    assert native_series_name(Presentation(path).slides[0].shapes[0].chart.series[0]) == "Edited native series"
    assert "Edited native series" in extract(path)["slides"][0]["shapes"][0]["chart_labels"]


def test_changed_scene_after_assessment_cannot_publish(setup_run):
    source, options, caller = setup_run
    initial = run_multi_agent_pipeline(source, options, caller=caller)
    store = RevisionStore(options.output_dir)
    from infrastructure.persistence.revisions import copy_bundle
    work = store.root / ".work" / "tampered_scene"
    copy_bundle(initial.parent, work)
    scene = json.loads((work / "scene_graph.json").read_text())
    scene["slides"][0]["objects"][0]["content"]["text"] = "Changed after assessment"
    (work / "scene_graph.json").write_text(json.dumps(scene), encoding="utf-8")
    with pytest.raises(ValueError, match="Scene changed after artifact assessment"):
        store.commit(work, base_revision=initial.parent.name, operation="slide_edit")
    assert store.resolve() == initial.parent
