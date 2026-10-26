from pathlib import Path

import pytest
from pydantic import ValidationError
from pptx import Presentation

from application.jobs.builder import PresentationBuilder
from contracts.ir2_content_schema import TargetIR2Presentation


ROOT = Path(__file__).parents[1]
LAYOUTS = ROOT / "config" / "ppt_default_layouts.json"


def test_ir2_schema_rejects_empty_slide_content():
    with pytest.raises(ValidationError, match="at least 1 item"):
        TargetIR2Presentation.model_validate(
            {"slides": [{"layout": "TITLE_CENTERED", "content": {}}]}
        )


def test_builder_refuses_to_save_empty_slide(tmp_path):
    output = tmp_path / "blank.pptx"

    with pytest.raises(ValueError, match="has no content regions"):
        PresentationBuilder(LAYOUTS).build(
            {"slides": [{"layout": "TITLE_CENTERED", "content": {}}]}, output
        )

    assert not output.exists()


def test_builder_creates_visible_shape_for_valid_ir2(tmp_path):
    deck = TargetIR2Presentation.model_validate(
        {
            "slides": [
                {
                    "layout": "TITLE_CENTERED",
                    "content": {"title": {"type": "heading", "text": "Transformer"}},
                }
            ]
        }
    )
    output = tmp_path / "deck.pptx"
    PresentationBuilder(LAYOUTS).build(deck.model_dump(exclude_none=True), output)

    presentation = Presentation(output)
    assert len(presentation.slides) == 1
    assert len(presentation.slides[0].shapes) == 1
    assert presentation.slides[0].shapes[0].text == "Transformer"
