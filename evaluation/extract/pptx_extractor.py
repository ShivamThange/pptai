"""Native extraction retaining paragraphs, hierarchy, cell styles and chart data."""
from pathlib import Path
from math import cos, sin, radians
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from domain.evaluation.quality_metadata import read_metadata
from domain.rendering.chart_content import native_series_name

EMU = 914400


def _font_value(run, paragraph, shape, attribute):
    value = getattr(run.font, attribute)
    if value is not None:
        return value
    value = getattr(paragraph.font, attribute)
    if value is not None:
        return value
    sources = [shape]
    if getattr(shape, 'is_placeholder', False):
        idx = shape.placeholder_format.idx
        sources.extend(p for p in shape.part.slide.slide_layout.placeholders if p.placeholder_format.idx == idx)
    for source in sources:
        nodes = source._element.xpath(f'.//a:lstStyle/a:lvl{paragraph.level + 1}pPr/a:defRPr')
        for node in nodes:
            if attribute == 'size' and node.get('sz'):
                from pptx.util import Pt
                return Pt(int(node.get('sz')) / 100)
            if attribute == 'name':
                face = node.xpath('./a:latin')
                if face and face[0].get('typeface'):
                    return face[0].get('typeface')
    if getattr(shape, 'is_placeholder', False):
        master = shape.part.slide.slide_layout.slide_master
        kind = 'titleStyle' if 'TITLE' in str(shape.placeholder_format.type) else 'bodyStyle'
        nodes = master._element.xpath(f'.//p:txStyles/p:{kind}/a:lvl{paragraph.level + 1}pPr/a:defRPr')
        for node in nodes:
            if attribute == 'size' and node.get('sz'):
                from pptx.util import Pt
                return Pt(int(node.get('sz')) / 100)
            if attribute == 'name':
                face = node.xpath('./a:latin')
                if face:
                    return face[0].get('typeface')
    return None


def _frame(frame, shape, role, paragraph_roles=()):
    paragraphs, runs = [], []
    for index, paragraph in enumerate(frame.paragraphs):
        paragraph_role = paragraph_roles[index] if index < len(paragraph_roles) else role
        paragraphs.append({'text': paragraph.text.replace('\v', '\n'), 'level': paragraph.level, 'role': paragraph_role})
        for run in paragraph.runs:
            size = _font_value(run, paragraph, shape, 'size')
            name = _font_value(run, paragraph, shape, 'name')
            if run.text.strip():
                runs.append({'text': run.text, 'font_size': round(size.pt, 2) if size else None,
                             'font_name': name, 'role': paragraph_role,
                             'bold': run.font.bold, 'italic': run.font.italic})
    return '\n'.join(p['text'] for p in paragraphs), paragraphs, runs


def _shape(shape, transform=(1, 1, 0, 0), parent_id=None):
    sx, sy, ox, oy = transform
    metadata = read_metadata(shape)
    role = metadata.get('role', 'body')
    if not metadata and getattr(shape, 'is_placeholder', False):
        if 'TITLE' in str(shape.placeholder_format.type):
            role = 'title'
    result = {'shape_id': shape.shape_id, 'shape_type': str(shape.shape_type),
              'left': (ox + shape.left * sx) / EMU, 'top': (oy + shape.top * sy) / EMU,
              'width': shape.width * sx / EMU, 'height': shape.height * sy / EMU,
              'text': '', 'paragraphs': [], 'text_runs': [], 'font_sizes': [], 'font_names': [],
              'table': [], 'table_cells': [], 'chart_labels': [], 'chart_categories': [], 'chart_data': [],
              'role': role, 'required': metadata.get('required', True), 'parent_id': parent_id,
              'container_id': metadata.get('container_id'), 'connector_endpoints': [],
              'is_line': shape.shape_type == MSO_SHAPE_TYPE.LINE, 'rotation': shape.rotation}
    if getattr(shape, 'has_text_frame', False):
        result['text'], result['paragraphs'], result['text_runs'] = _frame(
            shape.text_frame, shape, role, metadata.get('paragraph_roles', []))
    if getattr(shape, 'has_table', False):
        result['role'] = 'table'
        for r, row in enumerate(shape.table.rows):
            values = []
            for c, cell in enumerate(row.cells):
                text, paragraphs, runs = _frame(cell.text_frame, shape, 'table')
                values.append(text)
                if not cell.is_spanned:
                    result['table_cells'].append({'row': r, 'column': c, 'text': text,
                        'paragraphs': paragraphs, 'text_runs': runs,
                        'margins_inches': [cell.margin_left/EMU, cell.margin_top/EMU, cell.margin_right/EMU, cell.margin_bottom/EMU],
                        'row_span': cell.span_height, 'column_span': cell.span_width})
                    result['text_runs'].extend(runs)
            result['table'].append(values)
        result['text'] = '\n'.join('\t'.join(row) for row in result['table'])
    if getattr(shape, 'has_chart', False):
        chart = shape.chart
        labels = []
        if chart.has_title:
            labels.append(chart.chart_title.text_frame.text)
        for plot in chart.plots:
            try:
                categories = [str(cat.label) for cat in plot.categories]
                labels.extend(categories)
                result['chart_categories'].append(categories)
            except (AttributeError, ValueError):
                pass  # XY charts have no category axis.
        for series in chart.series:
            data = {'name': native_series_name(series), 'values': list(series.values)}
            for axis in ('x_values', 'y_values'):
                if hasattr(series, axis):
                    data[axis] = list(getattr(series, axis))
            result['chart_data'].append(data)
            labels.append(native_series_name(series))
        for node in chart._chartSpace.xpath('.//a:rPr | .//a:defRPr | .//a:endParaRPr'):
            if node.get('sz'):
                result['text_runs'].append({'text': '[chart style]', 'font_size': int(node.get('sz')) / 100,
                                            'font_name': None, 'role': 'chart_label'})
        result['chart_labels'] = labels
        values = [str(value) for series in result['chart_data'] for value in series['values']]
        result['text'] = '\n'.join(labels + values)
    if result['is_line']:
        result['connector_endpoints'] = [int(n.get('id')) for n in shape._element.xpath('.//a:stCxn | .//a:endCxn')]
        xfrm = shape._element.spPr.xfrm
        flip_h = xfrm is not None and xfrm.get('flipH') in {'1', 'true'}
        flip_v = xfrm is not None and xfrm.get('flipV') in {'1', 'true'}
        x, y, w, h = result['left'], result['top'], result['width'], result['height']
        result['line_segment'] = [x + (w if flip_h else 0), y + (h if flip_v else 0),
                                  x + (0 if flip_h else w), y + (0 if flip_v else h)]
    result['font_sizes'] = [run['font_size'] for run in result['text_runs'] if run['font_size'] is not None]
    result['font_names'] = [run['font_name'] for run in result['text_runs'] if run['font_name']]
    if result['rotation'] and not result['is_line']:
        angle = radians(result['rotation'])
        w, h = result['width'], result['height']
        rw, rh = abs(w*cos(angle)) + abs(h*sin(angle)), abs(w*sin(angle)) + abs(h*cos(angle))
        result.update(left=result['left'] + (w-rw)/2, top=result['top'] + (h-rh)/2, width=rw, height=rh)
    return result


def _walk(shapes, transform=(1, 1, 0, 0), parent_id=None, uncertainties=()):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            xfrm = shape._element.grpSpPr.xfrm
            sx, sy, ox, oy = transform
            gx = shape.width / xfrm.chExt.cx if xfrm.chExt.cx else 1
            gy = shape.height / xfrm.chExt.cy if xfrm.chExt.cy else 1
            child_transform = (sx*gx, sy*gy, ox + sx*(shape.left - gx*xfrm.chOff.x),
                               oy + sy*(shape.top - gy*xfrm.chOff.y))
            unresolved = uncertainties
            if shape.rotation or xfrm.get('flipH') in {'1', 'true'} or xfrm.get('flipV') in {'1', 'true'}:
                unresolved = (*uncertainties, 'Rotated/flipped group transform requires rendered inspection')
            yield from _walk(shape.shapes, child_transform, shape.shape_id, unresolved)
        else:
            result = _shape(shape, transform, parent_id)
            result['geometry_uncertainties'] = list(uncertainties)
            yield result


def extract(path, tool_name='unknown'):
    prs = Presentation(path)
    slides = []
    for number, slide in enumerate(prs.slides, 1):
        shapes = list(_walk(slide.shapes))
        for index, shape in enumerate(shapes):
            shape['z_order'] = index
        text = '\n'.join(s['text'] for s in shapes if s['text'])
        slides.append({'slide_no': number, 'text': text, 'word_count': len(text.split()), 'shapes': shapes})
    return {'candidate_id': Path(path).stem, 'tool_name': tool_name,
            'slide_width': prs.slide_width / EMU, 'slide_height': prs.slide_height / EMU,
            'slides': slides, 'image_flattened_candidate': bool(slides) and not any(s['text'] for s in slides)}
