"""Build retained real-format source fixtures, then record their byte hashes."""
import io
import json
import sys
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from infrastructure.utils.artifacts import atomic_json, file_hash


def main():
    from PIL import Image, ImageDraw, ImageFont
    from docx import Document
    from docx.shared import Inches
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    import fitz
    root = ROOT / 'tests/fixtures/ingestion'
    root.mkdir(parents=True, exist_ok=True)
    text = 'Attention combines values using learned weights.'
    (root / 'unicode.txt').write_bytes(('  ' + text + '\r\n\r\nUnicode café α.\n\n').encode('utf-8'))
    (root / 'bom.txt').write_bytes((text + '\nUnicode café.').encode('utf-16'))
    (root / 'latin.txt').write_bytes('café'.encode('cp1252'))
    (root / 'invalid.txt').write_bytes(b'hello\x00binary')
    (root / 'source.md').write_text(f'# Attention\n\n{text}\n\n| Subject | Value |\n| --- | --- |\n| Alpha | 10 |\n\n![unavailable](https://example.invalid/image.png)\n', encoding='utf-8')
    (root / 'records.csv').write_bytes(b'Subject,Subject,Value,Empty\r\n"Alpha\nEast",001,1.2300,\r\nBeta,002,0,\r\n')
    (root / 'many.csv').write_text('Subject,Value\n' + ''.join(f'Row {i:03},{i}.00\n' for i in range(205)), encoding='utf-8')
    (root / 'ragged.csv').write_bytes(b'Subject,Value\nAlpha,10,extra\nBeta,20\n')
    (root / 'malformed.csv').write_bytes(b'Subject,Value\n"Alpha,10\n')
    (root / 'ambiguous.csv').write_bytes(b'A,B;C\nD,E;F\n')
    (root / 'semicolon.csv').write_bytes(b'Label;Value\nAlpha;001.00\n')
    (root / 'noheader.csv').write_bytes(b'Alpha\t001\nBeta\t0\n')
    (root / 'headeronly.csv').write_bytes(b'Subject,Value\n')
    image = Image.new('RGB', (1200, 300), 'white')
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 42)
    draw.text((40, 75), text, fill='black', font=font)
    image.save(root / 'ocr.png')
    image.save(root / 'ocr.jpg', quality=95)
    oriented = image.transpose(Image.Transpose.ROTATE_90)
    exif = oriented.getexif(); exif[274] = 6
    oriented.save(root / 'oriented.jpg', quality=95, exif=exif)
    Image.new('RGB', (100, 80), 'white').save(root / 'textless.png')
    (root / 'corrupt.png').write_bytes(b'\x89PNG\r\n\x1a\ncorrupt')
    doc = Document(); doc.add_heading('Attention', 1); doc.add_paragraph(text)
    doc.add_paragraph('Unicode café α.', style='List Bullet')
    table = doc.add_table(rows=3, cols=2)
    table.cell(0,0).text = 'Subject'; table.cell(0,1).text = 'Value'
    table.cell(1,0).text = 'Alpha'; table.cell(1,1).text = '10'
    table.cell(2,0).merge(table.cell(2,1)).text = 'Merged note'
    doc.add_picture(str(root / 'ocr.png'), width=Inches(5))
    doc.save(root / 'source.docx')
    workbook = Workbook(); ws = workbook.active; ws.title = 'Visible'
    ws['A1'] = text; ws['A3'] = 0; ws['B3'] = '001'; ws['C3'] = 0.25; ws['C3'].number_format = '0%'
    ws['A5'] = datetime(2026,10,13); ws['B5'] = '=SUM(8888,9999)'; ws['C5'] = '=1+2'
    ws['A7'] = 'Merged'; ws.merge_cells('A7:C7'); ws['D205'] = 'Sparse tail'
    hidden = workbook.create_sheet('Hidden'); hidden.sheet_state = 'hidden'; hidden['A1'] = 'Hidden evidence'; hidden['B2'] = '#DIV/0!'
    chart = BarChart(); chart.add_data(Reference(ws,min_col=1,min_row=3,max_row=3)); ws.add_chart(chart,'F1')
    buffer = io.BytesIO(); workbook.save(buffer)
    with ZipFile(buffer) as source, ZipFile(root / 'source.xlsx','w',ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == 'xl/worksheets/sheet1.xml':
                data = data.replace(b'<dimension ref="A1:D205"/>', b'<dimension ref="A1:A1"/>')
                data = data.replace(b'<f>1+2</f><v></v>', b'<f>1+2</f><v>3</v>')
            target.writestr(info, data)
    pdf = fitz.open(); page = pdf.new_page()
    page.insert_text((72,100), text, fontsize=14)
    page.insert_text((72,145), 'Subject     Value\nAlpha          10',fontsize=12)
    page.insert_image(fitz.Rect(72,200,500,307), filename=str(root / 'ocr.png'))
    pdf.save(root / 'source.pdf', no_new_id=True)
    pdf.save(root / 'encrypted.pdf', encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw='owner', user_pw='locked', no_new_id=True)
    pdf.close()
    scanned = fitz.open(); page = scanned.new_page(width=1200,height=300)
    page.insert_image(page.rect, filename=str(root / 'ocr.png')); scanned.save(root / 'scanned.pdf', no_new_id=True); scanned.close()
    (root / 'corrupt.pdf').write_bytes(b'%PDF-1.7\ncorrupt')
    expectations = {
        'unicode.txt': {'blocks': 2, 'location_kind': 'text_span', 'preserved_newlines': 'CRLF and LF', 'unicode': 'café α'},
        'bom.txt': {'encoding': 'UTF-16 BOM', 'blocks': 1},
        'latin.txt': {'automatic_outcome': 'TEXT_ENCODING_REQUIRED', 'explicit_encoding': 'cp1252'},
        'records.csv': {'records': 3, 'data_rows': 2, 'columns': 4, 'duplicate_headers': True,
                        'record_2_physical_lines': [2, 3], 'lexical_values': ['001', '1.2300', '0', '']},
        'many.csv': {'records': 206, 'data_rows_per_projection': [200, 5]},
        'source.xlsx': {'worksheets': ['Visible', 'Hidden'], 'hidden_sheet': 'Hidden', 'tail_cell': 'D205',
                        'merged_range': 'A7:C7', 'formula_unknown': 'B5', 'formula_cached_unverified': 'C5',
                        'cached_result': 3, 'stale_dimension': 'A1:A1', 'objects': 'drawing and chart'},
        'source.docx': {'heading': 'Attention', 'merged_table': True, 'embedded_image': True, 'physical_page': None},
        'source.md': {'heading': 'Attention', 'table': True, 'linked_asset': 'UNRESOLVED_LINKED_ASSET'},
        'source.pdf': {'pages': 1, 'text': text, 'figure': True},
        'scanned.pdf': {'pages': 1, 'known_ocr_text': text},
        'ocr.png': {'pixels': [1200, 300], 'known_ocr_text': text},
        'ocr.jpg': {'pixels': [1200, 300], 'known_ocr_text': text},
        'oriented.jpg': {'pixels': [300, 1200], 'exif_orientation': 6},
        'textless.png': {'pixels': [100, 80], 'raw_status': 'complete', 'generation': 'SOURCE_CONTENT_UNAVAILABLE'},
        'ragged.csv': {'explicit_delimiter': ',', 'outcome': 'CSV_ROW_WIDTH'},
        'malformed.csv': {'explicit_delimiter': ',', 'outcome': 'CORRUPT_SOURCE'},
        'ambiguous.csv': {'outcome': 'CSV_DELIMITER_REQUIRED'},
        'headeronly.csv': {'outcome': 'NO_USABLE_CONTENT'},
        'corrupt.pdf': {'outcome': 'CORRUPT_SOURCE'},
        'encrypted.pdf': {'outcome': 'ENCRYPTED_SOURCE'},
    }
    manifest = {'schema_version': 1, 'known_text': text, 'baseline': {'docling':'2.121.0','docling-core':'2.92.0','openpyxl':'3.1.5'},
                'limits': 'These fixtures establish retained structure and known OCR samples, not general OCR accuracy.',
                'fixtures': {p.name:{'sha256':file_hash(p),'size_bytes':p.stat().st_size, 'expected': expectations.get(p.name, {})} for p in root.iterdir() if p.is_file() and p.name != 'manifest.json'}}
    atomic_json(root / 'manifest.json', manifest)
    print(json.dumps(manifest, indent=2))


if __name__ == '__main__':
    main()
