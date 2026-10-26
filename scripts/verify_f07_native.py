"""Real F02 mixed parsing and F07 progressive/gateway/revision/native acceptance.

Only model answers are fixtures. Source integrity, budgets, extraction, grounding,
native tables, both renderers and PowerPoint COM checks execute normally.
"""
import argparse
import asyncio
import json
import sys
from dataclasses import replace
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT)); sys.path.insert(0, str(ROOT/'tests')); sys.path.insert(0, str(ROOT/'scripts'))
from config.settings import get_multi_agent_settings
from config.service_settings import ServiceSettings
from pipeline.service import generate_deck_from_sources, edit_slide
from pipeline.orchestrator import MultiAgentOptions
from core.artifacts import atomic_json, file_hash
from core.revisions import RevisionStore
from core.context_budget import ContextBudgetError
from core.llm_client import LLMBudgetError
from test_context_coverage_f07 import SourceClient
from test_context_budget_f07 import Response
from verify_f02_native import SourceFixtureCaller
from pptx import Presentation


class NativeClient(SourceClient):
    cache_version = 'f07-native-v1'
    def __init__(self):
        settings = get_multi_agent_settings({'LLM_UNKNOWN_CONTEXT_TOKENS':'262144', 'LLM_CONTEXT_COUNT_MODE':'estimate',
            'LLM_KNOWLEDGE_EXTRACTION_INPUT_TOKENS':'22000', 'LLM_KNOWLEDGE_EXTRACTION_OUTPUT_TOKENS':'1792'})
        super().__init__(settings=settings)
        self.fixture = SourceFixtureCaller()
    def respond(self, request, timeout):
        body = json.loads(request.data)
        name = body['response_format']['json_schema']['name']
        if name == 'knowledgeunitbatch':
            return super().respond(request, timeout)
        self.requests.append(body)
        prompt = body['messages'][1]['content']
        if name == 'storyboard':
            result = self.fixture.generate_structured(stage='storyboard_planning', prompt=prompt)
            return Response(result.model_dump(mode='json'))
        if name == 'slidedraft':
            result = asyncio.run(self.fixture.generate_structured_async(stage='slide_composition', prompt=prompt))
            return Response({key:value for key,value in result.model_dump(mode='json').items() if key in {'layout','content'}})
        raise AssertionError(name)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--output-dir', default=str(ROOT.parent/'audit-artifacts/f07/native'))
    parser.add_argument('--parse-root', help='Optional existing integrity-checked F02 cache directory')
    args = parser.parse_args()
    root = Path(args.output_dir).resolve(); root.mkdir(parents=True, exist_ok=True)
    parse_root = Path(args.parse_root).resolve() if args.parse_root else root/'parsing'
    fixtures = ROOT/'tests/fixtures/ingestion'
    sources = [{'source_path': fixtures/name, 'parse_root': parse_root/name,
        'parse_options': {'ocr':'off'} if name.endswith('.pdf') else {}}
        for name in ['source.pdf','source.docx','records.csv','source.xlsx']]
    outcomes = []
    for backend in ['python','pptxgenjs']:
        options = MultiAgentOptions(layouts_path=str(ROOT/'config/ppt_default_layouts.json'),
            output_dir=str(root/backend), min_slides=4, max_slides=4, max_concurrency=1,
            max_repairs_per_slide=0, max_section_retries=0, resume=True, renderer=backend)
        first = NativeClient(); first.fail_at = 2
        try:
            generate_deck_from_sources(sources, parse_root, options, caller=first)
            raise AssertionError('Expected injected mid-extraction billing failure')
        except LLMBudgetError:
            pass
        assert not RevisionStore(options.output_dir).active_id()
        cached = list((Path(options.output_dir)/'.cache/extraction').glob('*.json'))
        assert len(cached) == 1
        gateway = NativeClient()
        deck = generate_deck_from_sources(sources, parse_root, options, caller=gateway)
        store = RevisionStore(options.output_dir); initial = store.resolve()
        raw = json.loads((initial/'raw_extraction.json').read_text(encoding='utf-8'))
        assert len(raw['documents']) == 4
        coverage = json.loads((initial/'context_coverage.json').read_text(encoding='utf-8'))
        assert coverage['context_completeness'] == 'complete'
        assert any(c['outcome']=='CACHE_REUSED' for c in coverage['calls'])
        assert len(json.loads((initial/'extraction_chunks.json').read_text())['chunks']) > 1
        assert len(gateway.requests) > 4
        table = next(shape.table for shape in Presentation(deck).slides[1].shapes if shape.has_table)
        assert table.cell(1,1).text == '001' and table.cell(1,2).text == '1.2300'
        for directory in [initial]:
            quality = json.loads((directory/'quality_report.json').read_text())
            assert quality['rendered']['status']=='EVALUATED_PASS' and quality['rendered']['details']['backend']=='powerpoint-com'
        before = store.verify(initial)['artifacts']
        editor = NativeClient()
        edited = edit_slide(options.output_dir, ROOT/'config/ppt_default_layouts.json', 2,
            'Keep the source table and clarify its heading', base_revision=initial.name, caller=editor)
        current = store.resolve()
        assert store.verify(initial)['artifacts'] == before
        assert len(json.loads((current/'context_calls.json').read_text())['calls']) > len(json.loads((initial/'context_calls.json').read_text())['calls'])
        attempted = NativeClient()
        try:
            edit_slide(options.output_dir, ROOT/'config/ppt_default_layouts.json', 2, 'Preserve all required facts. '*100000,
                base_revision=current.name, caller=attempted)
            raise AssertionError('Oversized mandatory edit should fail')
        except ContextBudgetError as error:
            assert error.code == 'REQUIRED_CONTEXT_TOO_LARGE'
        assert not attempted.requests and store.active_id() == current.name
        resumed = NativeClient()
        result = generate_deck_from_sources(sources, parse_root, options, caller=resumed)
        assert result.parent == current and not resumed.requests
        manifest = json.loads((current/'source_manifest.json').read_text())
        for member in manifest['members']:
            bundle = current/member['bundle_directory']
            assert file_hash(bundle/'source_snapshot.bin') == member['document']['source_sha256']
            for name,digest in member['artifacts'].items():
                assert file_hash(bundle/name)==digest
        quality = json.loads((current/'quality_report.json').read_text())
        outcomes.append({'backend':backend,'revision_path':str(current),'pptx_sha256':file_hash(edited),
            'context_coverage_sha256':file_hash(current/'context_coverage.json'),
            'extraction_chunks_sha256':file_hash(current/'extraction_chunks.json'),
            'previews':{Path(p).name:file_hash(p) for p in quality['rendered']['details']['images']},
            'retained_source_hashes':[m['document']['source_sha256'] for m in manifest['members']],
            'resumed_extraction_checkpoint':True,'zero_generation_calls_on_resume':True,
            'mandatory_over_budget_edit_preserved_active_revision':True,'native_renderer':'powerpoint-com'})
    atomic_json(root/'native_acceptance.json', {'model_answers':'typed source-grounded fixtures', 'outcomes':outcomes})
    print(json.dumps({'evidence':str(root/'native_acceptance.json'), 'outcomes':outcomes}, indent=2))


if __name__ == '__main__':
    main()
