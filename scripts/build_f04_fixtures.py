"""Synthetic F04 sources: contents are fixtures, never production fallback facts."""
import sys
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from infrastructure.utils.artifacts import atomic_json, file_hash


def build(directory=None):
    from PIL import Image, ImageDraw
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from docx import Document
    from docx.oxml import parse_xml
    root=Path(directory or ROOT/'tests/fixtures/assets');root.mkdir(parents=True,exist_ok=True)
    (root/'results.csv').write_text('Method,Accuracy (%)\nMethod A,80\nMethod B,90\n',encoding='utf-8')
    (root/'lexemes.csv').write_text('ID,Value,Note\n001,1.2300,"line one\nline two"\n002,0,\n003,-2,\n',encoding='utf-8')
    (root/'formula.md').write_text('# Method\nThe fraction describes the weighted mean. The variable x denotes observations and n denotes their count.\n\\[\\frac{\\sum_i x_i}{n}\\]\nThe budget is $12 and $18; these are currency.\n\\[\\unknown{x}\\]\n',encoding='utf-8')
    image=Image.new('RGB',(900,450),'white');draw=ImageDraw.Draw(image)
    draw.rectangle((100,200,280,400),fill='#3366CC');draw.rectangle((500,80,680,400),fill='#3366CC')
    draw.text((120,410),'Method A',fill='black');draw.text((520,410),'Method B',fill='black')
    image.save(root/'chart-picture.png')
    for orientation in range(1,9):
        exif=Image.Exif();exif[274]=orientation;image.save(root/f'oriented-{orientation}.jpg',exif=exif)
    book=Workbook();sheet=book.active;sheet.title='Results'
    for row in [('Method','Score'),('Method A',80),('Method B',90)]:sheet.append(row)
    sheet['A199']='Merged boundary';sheet.merge_cells('A199:A202');sheet['B205']='End'
    chart=BarChart();chart.add_data(Reference(sheet,min_col=2,min_row=1,max_row=3),titles_from_data=True)
    chart.set_categories(Reference(sheet,min_col=1,min_row=2,max_row=3));sheet.add_chart(chart,'D2')
    hidden=book.create_sheet('Hidden');hidden.sheet_state='hidden';hidden['A1']='001';hidden['B1']='=SUM(1,2)'
    book.save(root/'charts.xlsx')
    document=Document();document.add_heading('Experimental comparison',0)
    document.add_paragraph('The figure contrasts Method A and Method B on the same experiment. Its ordering is qualitative evidence, with no exact values printed.')
    document.add_picture(str(root/'chart-picture.png'))
    document.add_paragraph('Figure 1. Comparison of methods.')
    table=document.add_table(rows=4,cols=2)
    for i,row in enumerate([['Method','Accuracy (%)'],['Method A','80'],['Method B','90'],['Note','Same conditions']]):
        for j,text in enumerate(row):table.cell(i,j).text=text
    table.cell(3,0).merge(table.cell(3,1))
    document.add_paragraph('The numerator x is an observation and n is the number of observations. Their fraction defines the mean term.')
    document.paragraphs[-1]._p.append(parse_xml('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:f><m:num><m:r><m:t>x</m:t></m:r></m:num><m:den><m:r><m:t>n</m:t></m:r></m:den></m:f></m:oMath>'))
    document.save(root/'assets.docx')
    # A genuine Word chart relationship with its original embedded workbook.
    from zipfile import ZipFile,ZIP_DEFLATED
    from lxml import etree
    with ZipFile(root/'charts.xlsx') as xlsx:
        chart_xml=etree.fromstring(xlsx.read('xl/charts/chart1.xml'))
    chart_ns='http://schemas.openxmlformats.org/drawingml/2006/chart'
    rel_ns='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    external=etree.SubElement(chart_xml,'{'+chart_ns+'}externalData')
    external.set('{'+rel_ns+'}id','rIdWorkbook')
    with ZipFile(root/'assets.docx') as source, ZipFile(root/'embedded-chart.docx','w',ZIP_DEFLATED) as target:
        body=etree.fromstring(source.read('word/document.xml'))
        w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
        paragraph=etree.Element('{'+w+'}p')
        run=etree.SubElement(paragraph,'{'+w+'}r')
        run.append(etree.fromstring('''<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
            xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
            xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"
            xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
            <wp:inline><wp:extent cx="5486400" cy="3200400"/><wp:docPr id="99" name="Source chart"/>
            <a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:chart r:id="rIdFixtureChart"/></a:graphicData></a:graphic></wp:inline></w:drawing>'''))
        body.find('{'+w+'}body').insert(1,paragraph)
        rels=etree.fromstring(source.read('word/_rels/document.xml.rels'))
        pn='http://schemas.openxmlformats.org/package/2006/relationships'
        etree.SubElement(rels,'{'+pn+'}Relationship',Id='rIdFixtureChart',Type=rel_ns+'/chart',Target='charts/chart1.xml')
        types=etree.fromstring(source.read('[Content_Types].xml'))
        ct='http://schemas.openxmlformats.org/package/2006/content-types'
        etree.SubElement(types,'{'+ct+'}Override',PartName='/word/charts/chart1.xml',ContentType='application/vnd.openxmlformats-officedocument.drawingml.chart+xml')
        etree.SubElement(types,'{'+ct+'}Override',PartName='/word/embeddings/source.xlsx',ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        for item in source.infolist():
            payload=etree.tostring(body) if item.filename=='word/document.xml' else etree.tostring(rels) if item.filename=='word/_rels/document.xml.rels' else etree.tostring(types) if item.filename=='[Content_Types].xml' else source.read(item.filename)
            target.writestr(item,payload)
        target.writestr('word/charts/chart1.xml',etree.tostring(chart_xml))
        target.writestr('word/charts/_rels/chart1.xml.rels',f'<Relationships xmlns="{pn}"><Relationship Id="rIdWorkbook" Type="{rel_ns}/package" Target="../embeddings/source.xlsx"/></Relationships>')
        target.writestr('word/embeddings/source.xlsx',(root/'charts.xlsx').read_bytes())
    files={p.name:{'sha256':file_hash(p),'fixture':True} for p in sorted(root.iterdir()) if p.name!='manifest.json'}
    atomic_json(root/'manifest.json',{'schema_version':1,'synthetic':True,'files':files,
        'expected':{'results':{'difference':'10','unit':'percentage points'},'merge':'Results!A199:A202',
                    'chart_categories':['Method A','Method B'],'chart_values':['80','90'],
                    'formula':r'\frac{\sum_i x_i}{n}','currency_formula_count':0}})
    return root


if __name__=='__main__':
    build()
