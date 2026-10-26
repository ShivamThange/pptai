"""Reproducible R06 actual PowerPoint acceptance; never calls a live LLM.

Run from the repository root. Evidence is written outside the checkout by default.
"""
from __future__ import annotations
import argparse
import json
import subprocess
import sys
import tempfile
from dataclasses import asdict, replace
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from PIL import Image, ImageDraw
from pptx import Presentation
from core.artifacts import atomic_json, file_hash, sha256_value
from core.asset_registry import AssetRegistry
from core.builder import PresentationBuilder
from renderers.base import legacy_ir2_to_scene
from renderers.factory import get_renderer
from renderers.artifact import verify_native_coverage
from pipeline.slide_renderer import render_pptx_slides
from pipeline.quality_controller import run_deterministic_quality_checks
from pipeline.orchestrator import runtime_dependencies, run_multi_agent_pipeline, MultiAgentOptions
from pipeline.slide_editor import edit_multi_agent_slide
from core.revisions import RevisionStore
from core.chart_content import native_series_name
from schemas.raw_extraction_schema import SourceAsset, SourceLocation


def catalog():
    types = ["heading", "text_group", "bullets", "numbered_list", "table", "chart", "image", "formula", "diagram", "stat"]
    return {"slide_dimensions": {"width_inches": 13.333, "height_inches": 7.5}, "layouts": {"fixture": {"regions": {
        "title": {"left": .06, "top": .05, "width": .88, "height": .1, "allowed_types": ["heading"], "default_style": {"font_size": 28}},
        "body": {"left": .08, "top": .25, "width": .84, "height": .6, "allowed_types": types, "default_style": {"font_size": 20}}}}}}


def source_assets(root):
    assets = root / "assets"; assets.mkdir(exist_ok=True)
    image = assets / "measured_source.png"
    pixels = Image.new("RGB", (600, 300), "#EFF6FF"); draw = ImageDraw.Draw(pixels)
    draw.rectangle((50, 50, 180, 250), fill="#2563EB"); draw.ellipse((350, 50, 550, 250), fill="#0F172A")
    draw.line((180, 150, 350, 150), fill="#64748B", width=5); pixels.save(image)
    formula = assets / "source_equation.png"
    from matplotlib.mathtext import math_to_image
    math_to_image(r"$E=mc^2$", str(formula), dpi=200, format="png")
    values = []
    for kind, path in (("image", image), ("formula", formula)):
        ref = "r06-source:" + kind
        values.append(SourceAsset(asset_id=kind, kind=kind, section_path=["Fixture source"], file_path=str(path.resolve()),
            document_id="r06-source", source_refs=[ref], locations=[SourceLocation(document_id="r06-source", source_id=ref, page=1)],
            content_sha256=file_hash(path), availability="AVAILABLE"))
    atomic_json(root / "source_registry.json", [a.model_dump(mode="json") for a in values])
    return AssetRegistry(values, run_root=root, source_item_ids=["r06-source:image", "r06-source:formula"])


def slides(contents):
    return {"slides": [{"layout": "fixture", "notes": "Native R06 fixture notes", "content": {
        "title": {"type": "heading", "text": f"Renderer fixture {i}"}, "body": c}} for i, c in enumerate(contents, 1)]}


def render_checked(path, output):
    result = render_pptx_slides(path, output)
    atomic_json(output / "acceptance_render.json", asdict(result) | {"image_paths": [str(p) for p in result.image_paths]})
    if result.status != "EVALUATED_PASS" or result.backend != "powerpoint-com" or result.text_bounds_status != "EVALUATED_PASS":
        raise ValueError(f"Native acceptance failed: {path}: {result}")
    deterministic = run_deterministic_quality_checks(path, output, raise_on_failure=False)
    if deterministic.readiness == "BLOCKED":
        raise ValueError(f"Deterministic acceptance failed: {path}: {deterministic.hard_failures}")
    return {"pptx_sha256": file_hash(path), "render": json.loads((output/"acceptance_render.json").read_text()),
            "deterministic": deterministic.status, "slides": len(Presentation(path).slides)}


def native_edit_roundtrip(path, output, scene):
    import pythoncom, win32com.client
    pythoncom.CoInitialize(); app = None; deck = None
    # Office save APIs have a shorter path limit than normal filesystem APIs.
    with tempfile.TemporaryDirectory(prefix="r06-edit-") as directory:
        target = Path(directory) / "edited.pptx"
        try:
            app = win32com.client.DispatchEx("PowerPoint.Application")
            deck = app.Presentations.Open(str(path.resolve()), ReadOnly=False, Untitled=False, WithWindow=False)
            original_text = deck.Slides(1).Shapes(1).TextFrame.TextRange.Text
            deck.Slides(1).Shapes(1).TextFrame.TextRange.Text = "Native edit survived"
            # Fixture slide 2 is a native table, slide 3 a native chart.
            deck.Slides(2).Shapes(2).Table.Cell(2,2).Shape.TextFrame.TextRange.Text = "99"
            chart = deck.Slides(3).Shapes(2).Chart
            series = chart.SeriesCollection(1)
            series.Name = "Edited native series"
            deck.SaveAs(str(target), 24)
            import shutil
            shutil.copyfile(target, output)
        finally:
            if deck is not None: deck.Close()
            if app is not None: app.Quit()
            deck = app = None
            pythoncom.CoUninitialize()
    reopened = Presentation(output)
    assert reopened.slides[0].shapes[0].text == "Native edit survived"
    assert reopened.slides[1].shapes[1].table.cell(1,1).text == "99"
    assert native_series_name(reopened.slides[2].shapes[1].chart.series[0]) == "Edited native series"
    assert len(reopened.slides) == len(scene.slides)
    edited_scene = scene.model_copy(deep=True)
    edited_scene.slides[0].objects[0].content["text"] = "Native edit survived"
    edited_scene.slides[1].objects[1].content["rows"][0][1] = "99"
    edited_scene.slides[2].objects[1].content["series"][0]["name"] = "Edited native series"
    verify_native_coverage(output, edited_scene)
    return {"pptx_sha256": file_hash(output), "native_text_edit": True, "native_table_edit": True, "native_chart_edit": True,
            "original_text": original_text}


def pipeline_acceptance(root):
    # The generation answers are typed fixtures; all builders/validators/gates/
    # revision writes and native PowerPoint rendering run without success mocks.
    sys.path.insert(0, str(ROOT / "tests"))
    from test_cache_resume_revision_repair import FixtureCaller, fixture_source
    root.mkdir(exist_ok=True)
    source = fixture_source(root); caller = FixtureCaller(count=4)
    options = MultiAgentOptions(layouts_path=str(ROOT/"config/ppt_default_layouts.json"), output_dir=str(root/"run"),
                                min_slides=4, max_slides=4, max_repairs_per_slide=0, max_concurrency=1, renderer="pptxgenjs")
    original = run_multi_agent_pipeline(source, options, caller=caller)
    edited = edit_multi_agent_slide(output_dir=options.output_dir, layouts_path=options.layouts_path, slide_index=2,
                instruction="Use learned weights", base_revision=original.parent.name, caller=caller, renderer="pptxgenjs")
    calls_before = len(caller.calls)
    resumed = run_multi_agent_pipeline(source, replace(options, resume=True), caller=caller)
    assert resumed == edited and len(caller.calls) == calls_before
    store = RevisionStore(options.output_dir); record = store.verify(edited.parent)
    assert record["renderer"]["backend"] == "pptxgenjs" and record["renderer"]["artifact_hash"] == file_hash(edited)
    return {"initial_pptx_sha256": file_hash(original), "edited_pptx_sha256": file_hash(edited), "revision_id": edited.parent.name,
            "revision_record_sha256": file_hash(edited.parent/"revision.json"), "resume_generation_calls": 0,
            "quality": json.loads((edited.parent/"quality_report.json").read_text()), "renderer": record["renderer"]}


def main():
    parser = argparse.ArgumentParser(); parser.add_argument("--output-dir", default=str(ROOT.parent/"audit-artifacts/r06"))
    args = parser.parse_args(); root = Path(args.output_dir).resolve(); root.mkdir(parents=True, exist_ok=True)
    layouts = catalog(); registry = source_assets(root)
    table_deck = slides([{"type":"table","headers":["Measure","Result","Unit"],
        "rows":[["Revenue",str(i*10),"USD"],["Change",str(-i),"USD"],["Count","0","items"]]} for i in range(1,10)])
    parity_deck = slides([
        {"type":"text_group","lines":["Identical source content", "Both adapters retain these lines."]},
        {"type":"table","headers":["Measure","Result"],"rows":[["Revenue","12"],["Change","-3"]]},
        {"type":"chart","chart_type":"bar","categories":["A","B","C"],"series":[{"name":"Measured","values":[-3,0,5]}]},
        {"type":"bullets","items":["Source-backed content", {"text":"Nested point","level":1}]},
        {"type":"numbered_list","items":["Inspect source","Validate facts","Render native objects"]},
        {"type":"stat","style":{"value_font_size":42},"items":[{"value":"0","label":"Items"},{"value":"-3","label":"Change"}]},
        {"type":"image","source_asset_id":"image","caption":"Fixture source pixels"},
        {"type":"formula","source_asset_id":"formula","caption":"Original equation pixels"},
        {"type":"diagram","nodes":[{"id":"source","label":"Source"},{"id":"validate","label":"Validate"},{"id":"render","label":"Render"}],
         "edges":[{"from_id":"source","to_id":"validate"},{"from_id":"validate","to_id":"render"}]}])
    atomic_json(root/"layouts.json", layouts); atomic_json(root/"table_ir2.json", table_deck); atomic_json(root/"parity_ir2.json", parity_deck)
    scene = legacy_ir2_to_scene(parity_deck, layouts, registry); atomic_json(root/"parity_scene.json", scene.model_dump(mode="json"))
    evidence = {"base_commit": subprocess.check_output(["git","rev-parse","HEAD"],cwd=ROOT,text=True).strip(),
                "upstream_commit": "e534ff1c6b9d6e6846b1508a39f0bd2add5e60df", "scene_hash": sha256_value(scene.model_dump(mode="json")), "backends": {}}
    for backend in ("python", "pptxgenjs"):
        directory = root/backend; directory.mkdir(exist_ok=True)
        table = directory/"nine_tables.pptx"; builder = PresentationBuilder(layouts, backend=backend)
        builder.build(table_deck, table)
        parity = directory/"parity.pptx"; artifact = get_renderer(backend).render(scene, registry, {}, parity)
        verify_native_coverage(parity, scene)
        table_result = render_checked(table, directory/"table_render")
        parity_result = render_checked(parity, directory/"parity_render")
        roundtrip = directory/"native_edited.pptx"
        edits = native_edit_roundtrip(parity, roundtrip, scene)
        edits["render"] = render_checked(roundtrip, directory/"edited_render")
        evidence["backends"][backend] = {"renderer": artifact.metadata(), "tables": table_result, "parity": parity_result, "native_edits": edits}
        atomic_json(root/"native_acceptance.json", evidence)
        print(f"{backend}: nine-table native opening/render, shared-scene parity and editable roundtrip passed", flush=True)
    evidence["pipeline"] = pipeline_acceptance(root/"pipeline")
    evidence["source_files"] = runtime_dependencies()
    evidence["source_bundle_sha256"] = sha256_value(evidence["source_files"])
    atomic_json(root/"native_acceptance.json", evidence)
    print("JS generation/edit/resume: actual native gates and immutable revision publication passed", flush=True)


if __name__ == "__main__": main()
