"""Real multi-source intake, parsing, native table generation, edit and resume.

Only model answers are fixtures. Both renderers and PowerPoint COM gates run.
"""
import argparse
import json
import sys
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT)); sys.path.insert(0,str(ROOT/'tests'))
from fastapi.testclient import TestClient
from pptx import Presentation
from api.app import create_app
from config.service_settings import ServiceSettings
from config.settings import get_multi_agent_settings
from infrastructure.utils.artifacts import atomic_json, file_hash
from infrastructure.persistence.revisions import RevisionStore
from workers.local import LocalWorker
from test_cache_resume_revision_repair import FixtureCaller
from contracts.ir2_content_schema import SlideSpec
from contracts.storyboard_schema import SlidePlan, Storyboard


class SourceFixtureCaller(FixtureCaller):
    cache_version='f02-native-v1'
    def generate_structured(self, *, stage, prompt, **kwargs):
        if stage=='knowledge_extraction':
            return super().generate_structured(stage=stage,prompt=prompt,**kwargs)
        self.calls.append(stage)
        chunks=json.loads(prompt)['SOURCE_DATA']['chunks']
        return Storyboard(document_title='Source collection',slides=[SlidePlan(slide_id=i,
            title='Attention combines values',objective='Explain source evidence',narrative_role='opening' if i==1 else 'evidence',
            chunk_refs=[c['chunk_id'] for c in chunks],source_unit_ids=[u for c in chunks for u in c['unit_ids']],
            preferred_layouts=['TABLE_FULL_WIDTH'] if i==2 else ['TITLE_CENTERED']) for i in range(1,5)])

    async def generate_structured_async(self, *, prompt, **kwargs):
        payload=json.loads(prompt); context=payload['SOURCE_DATA']['context']; index=context['plan']['slide_id']
        if index!=2:
            return await super().generate_structured_async(prompt=prompt,**kwargs)
        self.calls.append(index)
        table=next(a for a in context['assets'] if a['kind']=='table' and any(':csv:' in r for r in a['source_refs']))
        edited='user_edit_instruction' in payload.get('requested_changes',{})
        return SlideSpec(layout='TABLE_FULL_WIDTH',content={'header':{'type':'heading','text':'Source values retained' if edited else 'Source values'},
            'table':{'type':'table','source_asset_id':table['asset_id'],'headers':[],'rows':[]}})


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--output',default=str(ROOT.parent/'audit-artifacts/f02/native'))
    args=parser.parse_args(); root=Path(args.output).resolve(); root.mkdir(parents=True,exist_ok=True)
    fixtures=ROOT/'tests/fixtures/ingestion'
    settings=ServiceSettings(root=root/'service',poll_interval=.05)
    callers=[]
    def factory(context):
        caller=SourceFixtureCaller(); caller.settings=get_multi_agent_settings({}); callers.append(caller); return caller
    results=[]
    with TestClient(create_app(settings)) as client:
        store=client.app.state.store
        worker=LocalWorker(settings,store=store,caller_factory=factory,provider_settings=get_multi_agent_settings({}))
        sources=[]
        for name in ['source.pdf','records.csv','source.docx']:
            uploaded=client.post('/api/documents',files={'file':(name,(fixtures/name).read_bytes())})
            assert uploaded.status_code==201,uploaded.text
            source={'document_id':uploaded.json()['id']}
            if name.endswith('.pdf'):
                source['parse_options']={'ocr':'off'}
            sources.append(source)
        with worker.lifetime():
            for backend in ['python','pptxgenjs']:
                request={'sources':sources,'renderer':backend,'min_slides':4,'max_slides':4,'max_concurrency':1,
                         'max_repairs_per_slide':0,'extraction_mode':'source-passages','resume':True}
                submitted=client.post('/api/decks',json=request)
                assert submitted.status_code==202,submitted.text
                job=submitted.json(); assert worker.process_next()
                completed=store.get(job['job_id']); assert completed.status=='completed',completed.public()
                deck_root=store.deck(job['deck_id'])['output_dir']; initial=RevisionStore(deck_root).resolve()
                assert client.get(completed.result['pptx_url']).content==(initial/'op_ppt.pptx').read_bytes()
                raw=json.loads((initial/'raw_extraction.json').read_text(encoding='utf-8'))
                assert len(raw['documents'])==3 and raw['document_id'] is None
                manifest=json.loads((initial/'source_manifest.json').read_text(encoding='utf-8'))
                assert len(manifest['members'])==3
                for member in manifest['members']:
                    bundle=initial/member['bundle_directory']
                    assert file_hash(bundle/'source_snapshot.bin')==member['document']['source_sha256']
                    for name,digest in member['artifacts'].items():
                        assert file_hash(bundle/name)==digest
                native=Presentation(initial/'op_ppt.pptx')
                table=next(s.table for s in native.slides[1].shapes if s.has_table)
                assert table.cell(1,1).text=='001' and table.cell(1,2).text=='1.2300'
                quality=json.loads((initial/'quality_report.json').read_text(encoding='utf-8'))
                assert quality['rendered']['status']=='EVALUATED_PASS' and quality['rendered']['details']['backend']=='powerpoint-com'
                edit=client.post(f"/api/decks/{job['deck_id']}/slides/2/edit",json={'instruction':'Keep the source table and clarify its heading','base_revision':initial.name})
                assert edit.status_code==202 and worker.process_next()
                edited=store.get(edit.json()['job_id']); assert edited.status=='completed',edited.public()
                current=RevisionStore(deck_root).resolve()
                assert (current/'source_manifest.json').is_file() and initial.is_dir()
                resume=client.post(f"/api/decks/{job['deck_id']}/generate",json={k:v for k,v in request.items() if k!='sources'})
                calls=sum(len(c.calls) for c in callers)
                assert resume.status_code==202 and worker.process_next()
                resumed=store.get(resume.json()['job_id']); assert resumed.status=='completed',resumed.public()
                assert resumed.result['revision_id']==current.name and sum(len(c.calls) for c in callers)==calls
                results.append({'backend':backend,'initial':completed.public(),'edited':edited.public(),'resumed':resumed.public(),
                    'source_manifest_sha256':file_hash(current/'source_manifest.json'),'collection_sha256':file_hash(current/'structural_ir_snapshot.json'),
                    'pptx_sha256':file_hash(current/'op_ppt.pptx'),'zero_generation_calls_on_resume':True,'revision_path':str(current),
                    'retained_source_hashes':[m['document']['source_sha256'] for m in manifest['members']]})
    atomic_json(root/'native_acceptance.json',{'provider':'Typed source-grounded fixture; no live provider calls',
        'native_renderer':'powerpoint-com','outcomes':results})
    print(json.dumps({'evidence':str(root/'native_acceptance.json'),'pptx_sha256':[r['pptx_sha256'] for r in results]},indent=2))


if __name__=='__main__':
    main()
