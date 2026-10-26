"""Real F04 parse/enrich/Office plus shared API generation/edit/download/resume.

Narrative/interpreter answers are explicitly fixtures unless --live is used.
"""
import argparse
import json
import sys
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/'tests'));sys.path.insert(0,str(ROOT/'scripts'))
from infrastructure.utils.artifacts import atomic_json,file_hash
from domain.extraction.asset_registry import AssetRegistry
from application.jobs.builder import PresentationBuilder
from infrastructure.ingestion.collection import assemble,retain_sources
from infrastructure.ingestion.asset_adapters import enrich_assets
from infrastructure.ingestion.source_adapter import build_raw_extraction
from infrastructure.utils.artifacts import parse_document_cached
from contracts.asset_schema import AssetOptions
from infrastructure.ingestion.asset_interpreter import asset_insights_to_units
from infrastructure.ingestion.knowledge_extractor import generate_validated_ir1
from f04_helpers import FIXTURES,AssetCaller
from verify_r06_native import catalog,slides,render_checked,native_edit_roundtrip
from test_cache_resume_revision_repair import FixtureCaller
from contracts.storyboard_schema import SlidePlan,Storyboard
from contracts.ir2_content_schema import SlideSpec
from pptx import Presentation


class PipelineCaller(FixtureCaller,AssetCaller):
    cache_version='f04-native-v1'
    def generate_structured(self, *, stage,prompt,**kwargs):
        if stage=='asset_interpretation':
            return AssetCaller.generate_structured(self,prompt=prompt,**kwargs)
        self.calls.append(stage)
        chunks=json.loads(prompt)['SOURCE_DATA']['chunks']
        return Storyboard(document_title='Asset findings',slides=[SlidePlan(slide_id=i,title='Source evidence',
            objective='Explain a source-backed finding',narrative_role='evidence',chunk_refs=[c['chunk_id'] for c in chunks],
            source_unit_ids=[u for c in chunks for u in c['unit_ids']],preferred_layouts=['TITLE_CENTERED']) for i in range(1,3)])
    async def generate_structured_async(self, *, prompt,**kwargs):
        data=json.loads(prompt);context=data['SOURCE_DATA']['context'];self.calls.append('composition')
        findings=[u for u in context['knowledge_units'] if u.get('origin_insight_ids')]
        text=findings[0]['summary'] if findings else 'Source evidence retained'
        return SlideSpec(layout='TITLE_CENTERED',content={'title':{'type':'heading','text':'Measured comparison'},
            'subtitle':{'type':'text_group','text':text}},source_unit_ids=context['plan']['source_unit_ids'])


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--output',required=True);parser.add_argument('--live',action='store_true')
    args=parser.parse_args();root=Path(args.output).resolve();root.mkdir(parents=True,exist_ok=True)
    caller=AssetCaller();members=[]
    for name in ['results.csv','lexemes.csv','charts.xlsx','formula.md','assets.docx','embedded-chart.docx']:
        structural=parse_document_cached(FIXTURES/name,root/'parses'/name,resume=True,
            parse_options={})
        result=json.loads((structural.parent/'parse_result.json').read_text())
        members.append({'path':str(structural),'result':result})
    structural=assemble(members,root/'collection');work=root/'enriched';work.mkdir(exist_ok=True)
    retain_sources(structural,work);raw=build_raw_extraction(structural,work)
    raw=enrich_assets(raw,structural,work,AssetOptions(mode='structured'),caller=caller)
    atomic_json(work/'raw_extraction.json',raw.model_dump(mode='json'))
    registry=AssetRegistry(raw.assets,run_root=work,source_item_ids=raw.source_item_ids,asset_source_items=raw.asset_source_items)
    table=next(a for a in raw.assets if a.kind=='table' and a.content.interpretation and any(i.calculation for i in a.content.interpretation.insights))
    finding=asset_insights_to_units(raw)[0].summary
    chart=next(a for a in raw.assets if a.kind=='chart' and a.content.reconstruction_status=='EDITABLE')
    merged=next(a for a in raw.assets if a.asset_id.endswith(':1:full-table'))
    figure=next(a for a in raw.assets if a.kind=='image' and a.availability=='AVAILABLE')
    formula=next(a for a in raw.assets if a.kind=='formula' and a.content and a.content.render_status=='RENDERED')
    contents=[{'type':'text_group','text':finding}, {'type':'table','source_asset_id':table.asset_id},
        {'type':'chart',**registry.resolve_chart(chart.asset_id)},
        {'type':'table','source_asset_id':merged.asset_id,'selection':{'asset_id':merged.asset_id,'rows':[198,199,200,201],'columns':[0]}},
        {'type':'image','source_asset_id':figure.asset_id,'caption':'Original source figure'},
        {'type':'formula','source_asset_id':formula.asset_id,'caption':'Source-backed Mathtext raster'}]
    visual_unit=next(u for u in asset_insights_to_units(raw) if u.structured_data['asset_eligibility']=='QUALIFIED_TEXT')
    contents.append({'type':'text_group','text':visual_unit.summary})
    word_chart=next(a for a in raw.assets if a.kind=='chart' and a.content.data_state=='VERIFIED_SOURCE_DATA' and any(
        item.embedded_package_part for item in raw.asset_source_items if item.source_id in a.source_refs))
    contents.append({'type':'chart',**registry.resolve_chart(word_chart.asset_id)})
    from contracts.ir1_schema import ValidatedIR1
    from contracts.storyboard_schema import TopicChunk
    from infrastructure.ingestion.slide_composer import retrieve_slide_context,validate_slide_grounding
    units=asset_insights_to_units(raw)
    ir1=ValidatedIR1(schema_version='4',document_title=raw.document_name,units=units,assets=[a.model_dump(mode='json') for a in raw.assets],
        source_item_ids=raw.source_item_ids,source_blocks=raw.blocks,asset_source_items=raw.asset_source_items)
    chunk=TopicChunk(chunk_id='asset',label='Finding',unit_ids=[visual_unit.unit_id],digest=visual_unit.summary,token_estimate=100)
    plan=SlidePlan(slide_id=1,title='Figure finding',objective='Explain visual evidence',narrative_role='evidence',chunk_refs=['asset'],source_unit_ids=[visual_unit.unit_id])
    visual_context=retrieve_slide_context(plan,ir1,[chunk],Storyboard(document_title='Source findings',slides=[plan]))
    visual_slide=SlideSpec(layout='fixture',content={'body':{'type':'text_group','text':visual_unit.summary}},source_unit_ids=[visual_unit.unit_id])
    validate_slide_grounding(visual_slide,visual_context)
    evidence={'fixture_answers':True,'native':[],'semantic_findings':finding,'live':'not_requested',
              'visual_finding':visual_unit.summary,'visual_object_usage':visual_slide.asset_usage}
    for backend in ['python','pptxgenjs']:
        deck=root/(backend+'.pptx');builder=PresentationBuilder(catalog(),asset_registry=registry,backend=backend)
        builder.build(slides(contents),deck)
        checked=render_checked(deck,root/(backend+'-render'))
        edited=root/(backend+'-edited.pptx')
        # R06 helper edits row 2. Its fixture has a separate header row;
        # rich grids include original headers in rows, so adjust its expectation.
        edit_scene=builder.last_scene.model_copy(deep=True)
        original_table=edit_scene.slides[1].objects[1].content
        original_table['headers']=original_table['rows'].pop(0)
        edit=native_edit_roundtrip(deck,edited,edit_scene)
        checked_edit=render_checked(edited,root/(backend+'-edited-render'))
        evidence['native'].append({'backend':backend,'initial':checked,'native_edit':edit,'edited':checked_edit})
    # Shared queue and R05 published assets, real source parsing, COM and editing.
    from api.app import create_app
    from config.service_settings import ServiceSettings
    from config.settings import get_multi_agent_settings
    from fastapi.testclient import TestClient
    from workers.local import LocalWorker
    from infrastructure.persistence.revisions import RevisionStore
    settings=ServiceSettings(root=root/'service',layouts_path=ROOT/'config/ppt_default_layouts.json')
    callers=[]
    def factory(context):
        value=PipelineCaller();value.settings=get_multi_agent_settings({});callers.append(value);return value
    evidence['api']=[]
    with TestClient(create_app(settings)) as client:
        uploaded=client.post('/api/documents',files={'file':('results.csv',(FIXTURES/'results.csv').read_bytes())})
        upload2=client.post('/api/documents',files={'file':('lexemes.csv',(FIXTURES/'lexemes.csv').read_bytes())})
        worker=LocalWorker(settings,store=client.app.state.store,caller_factory=factory,provider_settings=get_multi_agent_settings({}))
        with worker.lifetime():
            for backend in ['python','pptxgenjs']:
                submitted=client.post('/api/decks',json={'sources':[{'document_id':uploaded.json()['id']},{'document_id':upload2.json()['id']}],
                    'renderer':backend,'min_slides':2,'max_slides':2,'max_concurrency':1,'max_repairs_per_slide':0,
                    'extraction_mode':'source-passages','asset_options':{'mode':'structured'}})
                assert submitted.status_code==202,submitted.text
                job=submitted.json();worker.process_next();completed=client.app.state.store.get(job['job_id'])
                assert completed.status=='completed',completed.public()
                initial=RevisionStore(client.app.state.store.deck(job['deck_id'])['output_dir']).resolve()
                assert client.get(completed.result['pptx_url']).content==(initial/'op_ppt.pptx').read_bytes()
                url=f"/api/decks/{job['deck_id']}/revisions/{initial.name}/assets"
                asset_id=client.get(url).json()['items'][0]['asset_id']
                assert client.get(url+'/detail',params={'asset_id':asset_id}).status_code==200
                assert client.get(url+'/file',params={'asset_id':asset_id,'representation_id':'original'}).status_code==200
                assert client.get(url+'/detail',params={'asset_id':'foreign'}).status_code==404
                assert client.get(url,params={'limit':101}).status_code==422
                edit=client.post(f"/api/decks/{job['deck_id']}/slides/1/edit",json={'instruction':'Keep the supported finding','base_revision':initial.name,'max_repairs':0})
                assert edit.status_code==202,edit.text
                worker.process_next();edited=client.app.state.store.get(edit.json()['job_id']);assert edited.status=='completed',edited.public()
                revision=RevisionStore(client.app.state.store.deck(job['deck_id'])['output_dir']).resolve()
                assert '10 percentage points' in next(s.text for s in Presentation(revision/'op_ppt.pptx').slides[0].shapes if s.has_text_frame and 'percentage' in s.text)
                regenerated=client.post(f"/api/decks/{job['deck_id']}/generate",json={'resume':True,'renderer':backend,
                    'min_slides':2,'max_slides':2,'max_concurrency':1,'max_repairs_per_slide':0,
                    'extraction_mode':'source-passages','asset_options':{'mode':'structured'}})
                assert regenerated.status_code==202,regenerated.text
                worker.process_next();resumed=client.app.state.store.get(regenerated.json()['job_id']);assert resumed.status=='completed',resumed.public()
                assert not callers[-1].calls,'Unchanged resume made new calls'
                evidence['api'].append({'backend':backend,'revision':revision.name,'pptx_sha256':file_hash(revision/'op_ppt.pptx'),
                    'resume_calls':len(callers[-1].calls),'asset_download':'passed','edit':'passed'})
    if args.live:
        from config.settings import resolve_environment
        from infrastructure.llm.llm_client import LLMClient,LLMRequestError
        env=resolve_environment(ROOT.parent/'BTPs5/.env');settings=get_multi_agent_settings(env)
        live=LLMClient(settings=settings,api_key=env.get('OPENROUTER_API_KEY' if settings.provider=='openrouter' else 'GEMINI_API_KEY'))
        try:
            retain_sources(structural,root/'live')
            live_raw=enrich_assets(build_raw_extraction(structural,root/'live'),structural,root/'live',AssetOptions(mode='structured'),caller=live)
            atomic_json(root/'live-result.json',live_raw.model_dump(mode='json'));evidence['live']='completed'
        except LLMRequestError as error:
            evidence['live']={'status':'unavailable','error_class':error.error_class}
    atomic_json(root/'native_acceptance.json',evidence)
    print(json.dumps({'native_backends':len(evidence['native']),'api_backends':len(evidence['api']),'live':evidence['live']}))


if __name__=='__main__':main()
