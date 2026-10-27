"""
Native Python primitives retained from the migration backend.
"""

from pathlib import Path
from typing import Dict, Any, List, Optional

from domain.extraction.asset_registry import AssetBindingError

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn

# --------------------------------------------------------------------------
# Helper Functions & Mappings
# --------------------------------------------------------------------------

ALIGN_MAP = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY,
}

CHART_TYPE_MAP = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
    "pie": XL_CHART_TYPE.PIE,
    "area": XL_CHART_TYPE.AREA,
}

DEFAULT_ASSET_SEARCH_DIRS = [
    Path("output/assets/images"),
    Path("output/assets/formulas"),
    Path("output/assets/tables"),
    Path("output/assets"),
    Path("."),
]


def hex_to_rgb(hex_str: Optional[str]) -> RGBColor:
    """Converts hex color string (e.g. '#0F172A') to RGBColor."""
    if not hex_str:
        return RGBColor(0x00, 0x00, 0x00)
    h = str(hex_str).lstrip("#")
    if len(h) != 6:
        return RGBColor(0x00, 0x00, 0x00)
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def merged_style(region_def: Dict[str, Any], content_obj: Dict[str, Any]) -> Dict[str, Any]:
    style = dict(region_def.get("default_style", {}))
    style.update(content_obj.get("style", {}))
    return style


def resolve_asset_path(
    asset_ref: Optional[str], path_ref: Optional[str] = None, search_dirs: Optional[List[Path]] = None
) -> Optional[Path]:
    """Resolves image or asset path across standard asset directories."""
    candidates = []
    if path_ref:
        candidates.append(Path(path_ref))
    if asset_ref:
        candidates.append(Path(asset_ref))
        for sdir in search_dirs or DEFAULT_ASSET_SEARCH_DIRS:
            candidates.append(sdir / asset_ref)

    for cand in candidates:
        if cand.is_file():
            return cand.resolve()
    return None


class RegionBox:
    """Computes absolute EMU geometry for a region from fractional coordinates."""

    def __init__(self, region_def: Dict[str, Any], slide_w_in: float, slide_h_in: float):
        self.left = Inches(region_def["left"] * slide_w_in)
        self.top = Inches(region_def["top"] * slide_h_in)
        self.width = Inches(region_def["width"] * slide_w_in)
        self.height = Inches(region_def["height"] * slide_h_in)


# --------------------------------------------------------------------------
# Presentation Builder Class
# --------------------------------------------------------------------------

class PythonPrimitives:
    """Native primitives for validated, resolved scenes from PythonRenderer."""

    def __init__(self, slide_w_in: float, slide_h_in: float):
        self.slide_w_in = slide_w_in
        self.slide_h_in = slide_h_in

    # --------------------------------------------------------------------------
    # Region Dispatch
    # --------------------------------------------------------------------------

    def _render_region_content(
        self, slide: Any, region_def: Dict[str, Any], content_obj: Dict[str, Any], region_name: str
    ) -> None:
        ctype = content_obj.get("type")
        box = RegionBox(region_def, self.slide_w_in, self.slide_h_in)

        if ctype == "heading":
            self._render_heading(slide, box, content_obj, region_def)
        elif ctype == "text_group":
            self._render_text_group(slide, box, content_obj, region_def)
        elif ctype in ("bullets", "numbered_list"):
            self._render_bullets(slide, box, content_obj, region_def, numbered=(ctype == "numbered_list"))
        elif ctype == "chart":
            self._render_chart(slide, box, content_obj, region_def)
        elif ctype == "table":
            self._render_table(slide, box, content_obj, region_def)
        else:
            raise ValueError(f"Unsupported native primitive: {ctype}")

    def _set_run_style(self, run: Any, style: Dict[str, Any], default_size: int = 18) -> None:
        run.font.size = Pt(style.get("font_size", default_size))
        run.font.bold = bool(style.get("bold", False))
        run.font.italic = bool(style.get("italic", False))
        run.font.color.rgb = hex_to_rgb(style.get("color", "#0F172A"))
        if style.get("font_name"):
            run.font.name = style["font_name"]

    def _apply_bullet_formatting(
        self, paragraph: Any, numbered: bool = False, level: int = 0, bullet_char: str = "\u2022"
    ) -> None:
        pPr = paragraph._p.get_or_add_pPr()
        for tag in ("a:buChar", "a:buAutoNum", "a:buNone"):
            el = pPr.find(qn(tag))
            if el is not None:
                pPr.remove(el)

        indent_step = 228600  # 0.25 inches
        marL = indent_step * (level + 1)
        pPr.set("marL", str(marL))
        pPr.set("indent", str(-indent_step))

        if numbered:
            bu = pPr.makeelement(qn("a:buAutoNum"), {"type": "arabicPeriod"})
        else:
            bu = pPr.makeelement(qn("a:buChar"), {"char": bullet_char})
        pPr.append(bu)

    # --------------------------------------------------------------------------
    # Element Renderers
    # --------------------------------------------------------------------------

    def _render_heading(self, slide: Any, box: RegionBox, content_obj: Dict[str, Any], region_def: Dict[str, Any]) -> None:
        style = merged_style(region_def, content_obj)
        tb = slide.shapes.add_textbox(box.left, box.top, box.width, box.height)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = ALIGN_MAP.get(style.get("align", "left"), PP_ALIGN.LEFT)
        run = p.add_run()
        text_val = content_obj.get("text") or content_obj.get("title") or content_obj.get("heading") or ""
        run.text = str(text_val)
        self._set_run_style(run, style, default_size=32)

    def _render_text_group(self, slide: Any, box: RegionBox, content_obj: Dict[str, Any], region_def: Dict[str, Any]) -> None:
        style = merged_style(region_def, content_obj)
        tb = slide.shapes.add_textbox(box.left, box.top, box.width, box.height)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE

        lines = content_obj.get("lines") or content_obj.get("items") or content_obj.get("bullets")
        if lines is None:
            text = content_obj.get("text", "")
            lines = text.split("\n") if isinstance(text, str) and text else ([str(text)] if text else [""])
        elif isinstance(lines, str):
            lines = [lines]

        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = ALIGN_MAP.get(style.get("align", "left"), PP_ALIGN.LEFT)
            run = p.add_run()
            line_str = line.get("text", "") if isinstance(line, dict) else str(line)
            run.text = line_str
            self._set_run_style(run, style, default_size=18)

    def _render_bullets(
        self, slide: Any, box: RegionBox, content_obj: Dict[str, Any], region_def: Dict[str, Any], numbered: bool = False
    ) -> None:
        style = merged_style(region_def, content_obj)
        
        # Robust fallback extraction for bullet items
        items = (
            content_obj.get("items")
            or content_obj.get("bullets")
            or content_obj.get("points")
            or content_obj.get("lines")
        )
        if isinstance(items, str):
            items = [items]
        elif items is None:
            text_val = content_obj.get("text")
            if text_val:
                items = text_val.split("\n") if isinstance(text_val, str) else [str(text_val)]
            else:
                items = []

        if not items:
            return

        base_size = style.get("font_size", 16)

        tb = slide.shapes.add_textbox(box.left, box.top, box.width, box.height)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE

        for i, item in enumerate(items):
            if isinstance(item, dict):
                text = item.get("text", "")
                level = int(item.get("level", 0))
            else:
                text = str(item)
                level = 0

            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = ALIGN_MAP.get(style.get("align", "left"), PP_ALIGN.LEFT)
            p.space_after = Pt(4)
            run = p.add_run()
            run.text = text
            
            item_style = dict(style)
            item_style["font_size"] = base_size
            self._set_run_style(run, item_style, default_size=base_size)
            self._apply_bullet_formatting(p, numbered=numbered, level=level)

    def _render_chart(self, slide: Any, box: RegionBox, content_obj: Dict[str, Any], region_def: Dict[str, Any]) -> None:
        style = merged_style(region_def, content_obj)
        chart_type = CHART_TYPE_MAP.get(content_obj.get("chart_type", "column"), XL_CHART_TYPE.COLUMN_CLUSTERED)
        categories = content_obj.get("categories", [])
        series_list = content_obj.get("series", [])

        chart_data = CategoryChartData()
        chart_data.categories = categories

        for index, series in enumerate(series_list, 1):
            chart_data.add_series(series.get("name", f"Series {index}"), series["values"])

        gframe = slide.shapes.add_chart(chart_type, box.left, box.top, box.width, box.height, chart_data)
        chart = gframe.chart

        # Theme-aware chart typography and series colors (F13 dynamic theming).
        font_name = style.get("font_name")
        try:
            from pptx.util import Pt as _Pt
            chart.font.size = _Pt(12)
            if font_name:
                chart.font.name = font_name
            if style.get("color"):
                chart.font.color.rgb = hex_to_rgb(style.get("color"))
        except Exception:
            pass  # Chart-level font is cosmetic; never fail a build over it.
        chart_colors = style.get("chart_colors") or []
        for index, series_plot in enumerate(chart.plots):
            for s_index, series in enumerate(series_plot.series):
                color = chart_colors[(s_index) % len(chart_colors)] if chart_colors else None
                if not color:
                    continue
                try:
                    series.format.fill.solid()
                    series.format.fill.fore_color.rgb = hex_to_rgb(color)
                except Exception:
                    continue  # Line/pie series may reject solid fills; keep defaults.

        if content_obj.get("title"):
            chart.has_title = True
            chart.chart_title.text_frame.text = content_obj["title"]
        else:
            chart.has_title = False

        chart.has_legend = len(series_list) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False

    def _render_table(self, slide: Any, box: RegionBox, content_obj: Dict[str, Any], region_def: Dict[str, Any]) -> None:
        """Renders dynamic tables cleanly without XML corruption or PowerPoint repair errors."""
        def clean_txt(val):
            return str(val)

        style = merged_style(region_def, content_obj)
        headers = content_obj.get("headers", [])
        rows = content_obj.get("rows", [])

        if not rows:
            raise AssetBindingError("Table has no usable rows")

        # Determine strict column count across headers and rows
        n_cols = len(headers) if headers else len(rows[0])

        n_rows = len(rows) + (1 if headers else 0)

        font_size = style.get("font_size", 13)

        gtable = slide.shapes.add_table(n_rows, n_cols, box.left, box.top, box.width, box.height)
        table = gtable.table

        # Explicit measured column widths (from the fitting pass) instead of
        # PowerPoint's even distribution, which wraps text unpredictably.
        col_widths_in = content_obj.get("col_widths_in") or []
        if col_widths_in and len(col_widths_in) == n_cols:
            total = sum(col_widths_in) or 1.0
            available_in = box.width / 914400.0
            for index, width_in in enumerate(col_widths_in):
                table.columns[index].width = Emu(int(914400.0 * width_in * available_in / total))

        # Neutral slate borders/banding instead of the blue default table style.
        tbl = table._tbl
        tblPr = tbl.tblPr
        tblPr.set("firstRow", "1" if headers else "0")
        tblPr.set("bandRow", "1")
        for child in list(tblPr):
            if child.tag == qn("a:tableStyleId"):
                tblPr.remove(child)
        style_id = tblPr.makeelement(qn("a:tableStyleId"), {})
        style_id.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
        tblPr.append(style_id)

        # Apply cell padding
        for row in table.rows:
            for cell in row.cells:
                cell.margin_left = Inches(0.06)
                cell.margin_right = Inches(0.06)
                cell.margin_top = Inches(0.04)
                cell.margin_bottom = Inches(0.04)

        # 1. Render Header Row safely
        r_offset = 0
        if headers:
            header_text_color = style.get("header_text", "#FFFFFF")
            for c in range(n_cols):
                htext = clean_txt(headers[c])
                cell = table.cell(0, c)
                cell.text = htext
                for p in cell.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.LEFT
                    for run in p.runs:
                        run.font.bold = True
                        run.font.size = Pt(font_size)
                        run.font.color.rgb = hex_to_rgb(header_text_color)
                        if style.get("font_name"):
                            run.font.name = style["font_name"]
                cell.fill.solid()
                cell.fill.fore_color.rgb = hex_to_rgb(style.get("header_bg", "#0F172A"))
            r_offset = 1
        # 2. Render Data Rows strictly bound to (n_rows, n_cols) grid
        band_bg = style.get("band_bg", "#F8FAFC")
        # Unbanded rows explicitly take the slide background (stashed per slide
        # by the adapter), matching the JS backend instead of implicit white -
        # dark themes otherwise show light stripes inside dark tables.
        page_bg = style.get("row_bg") or getattr(self, "page_background", None) or "#FFFFFF"
        for r_idx, row in enumerate(rows):
            row_num = r_idx + r_offset
            for c in range(n_cols):
                val = clean_txt(row[c])
                cell = table.cell(row_num, c)
                cell.text = val
                for p in cell.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.LEFT
                    for run in p.runs:
                        run.font.size = Pt(font_size)
                        run.font.color.rgb = hex_to_rgb(style.get("color", "#334155"))
                        if style.get("font_name"):
                            run.font.name = style["font_name"]
                if r_idx % 2 == 1:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = hex_to_rgb(band_bg)
                else:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = hex_to_rgb(page_bg)

        for r, c, rs, cs in content_obj.get('merges', []):
            anchor = table.cell(r, c)
            original_text = anchor.text
            anchor.merge(table.cell(r+rs-1, c+cs-1))
            anchor.text = original_text


