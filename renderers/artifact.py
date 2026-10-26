"""Shared package coverage and role metadata, before atomic output publication."""
import hashlib
import json
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
from pptx import Presentation
from core.quality_metadata import PREFIX
from core.chart_content import native_series_name

NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def seal_native_metadata(path, scene):
    """Patch only native metadata/connections; never reserialize the whole deck.

    PptxGenJS 3.12 numbers table frames independently of other shapes. Normalize
    IDs per slide before binding connectors so mixed slides also open natively.
    """
    updates = {}
    with ZipFile(path) as source:
        for index, slide in enumerate(scene.slides, 1):
            name = f"ppt/slides/slide{index}.xml"
            tree = etree.fromstring(source.read(name))
            nodes = tree.xpath("./p:cSld/p:spTree/*[p:nvSpPr or p:nvPicPr or p:nvGraphicFramePr or p:nvCxnSpPr]", namespaces=NS)
            mapping = {}
            for shape_id, node in enumerate(nodes, 2):
                props = node.xpath(".//p:cNvPr", namespaces=NS)[0]
                props.set("id", str(shape_id))
                mapping[props.get("name")] = (node, props)
            if len(mapping) != len(nodes) or set(mapping) != {o.object_id for o in slide.objects}:
                raise ValueError("Native artifact object coverage differs from scene")
            for obj in slide.objects:
                node, props = mapping[obj.object_id]
                props.set("descr", PREFIX + json.dumps({"role": obj.role, "required": obj.required,
                    "object_id": obj.object_id, "source_unit_ids": obj.source_unit_ids}, ensure_ascii=False))
                if obj.kind == "connector":
                    # Native lines become actual attached p:cxnSp primitives.
                    node.tag = "{" + NS["p"] + "}cxnSp"
                    nv = node[0]
                    nv.tag = "{" + NS["p"] + "}nvCxnSpPr"
                    cnv = nv[1]
                    cnv.tag = "{" + NS["p"] + "}cNvCxnSpPr"
                    for child in list(cnv):
                        cnv.remove(child)
                    for tag, endpoint, site in (("stCxn", "from_id", "3"), ("endCxn", "to_id", "1")):
                        etree.SubElement(cnv, "{" + NS["a"] + "}" + tag,
                                         id=mapping[obj.content[endpoint]][1].get("id"), idx=site)
            updates[name] = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
        temp = path.with_suffix(".sealed.pptx")
        with ZipFile(temp, "w", ZIP_DEFLATED) as target:
            for item in source.infolist():
                target.writestr(item, updates.get(item.filename, source.read(item.filename)))
    temp.replace(path)


def verify_native_coverage(path, scene):
    deck = Presentation(path)
    if len(deck.slides) != len(scene.slides):
        raise ValueError("Renderer changed slide count")
    for slide, expected in zip(deck.slides, scene.slides):
        shapes = {s.name: s for s in slide.shapes}
        if len(shapes) != len(slide.shapes) or set(shapes) != {o.object_id for o in expected.objects}:
            raise ValueError("Renderer omitted or duplicated scene objects")
        for obj in expected.objects:
            shape = shapes[obj.object_id]
            if obj.kind != "connector":
                actual_box = [v / 914400 for v in (shape.left, shape.top, shape.width, shape.height)]
                wanted_box = [obj.box.x, obj.box.y, obj.box.w, obj.box.h]
                if any(abs(a-b) > .005 for a, b in zip(actual_box, wanted_box)):
                    raise ValueError("Native object changed resolved geometry; reduce content or solve a larger box")
            if obj.kind == "table":
                if not shape.has_table:
                    raise ValueError("Required native table was rasterized")
                rows = ([obj.content["headers"]] if obj.content.get("headers") else []) + obj.content["rows"]
                actual = [[c.text for c in row.cells] for row in shape.table.rows]
                if actual != [[str(v) for v in row] for row in rows]:
                    raise ValueError("Native table changed values or column lengths")
            elif obj.kind == "chart":
                if not shape.has_chart:
                    raise ValueError("Required native chart was rasterized")
                chart = shape.chart
                if [cat.label for cat in chart.plots[0].categories] != [str(v) for v in obj.content["categories"]]:
                    raise ValueError("Native chart changed categories")
                if [list(s.values) for s in chart.series] != [s["values"] for s in obj.content["series"]]:
                    raise ValueError("Native chart changed series values")
                if [native_series_name(s) for s in chart.series] != [s["name"] for s in obj.content["series"]]:
                    raise ValueError("Native chart changed series names")
            elif obj.kind in {"image", "formula"}:
                if not hasattr(shape, "image") or hashlib.sha256(shape.image.blob).hexdigest() != obj.content["content_sha256"]:
                    raise ValueError("Required real image/formula pixels changed or are absent")
            elif obj.kind in {"heading", "text_group", "shape"}:
                text = "\n".join(obj.content["lines"]) if obj.content.get("lines") else obj.content.get("text", "")
                if not shape.has_text_frame or shape.text.replace("\v", "\n") != text:
                    raise ValueError("Native text changed or is absent")
            elif obj.kind in {"bullets", "numbered_list"}:
                items = [i.get("text", "") if isinstance(i, dict) else str(i) for i in obj.content["items"]]
                if not shape.has_text_frame or shape.text.replace("\v", "\n").split("\n") != items:
                    raise ValueError("Native list text changed or is absent")
        if expected.notes and expected.notes not in slide.notes_slide.notes_text_frame.text:
            raise ValueError("Renderer omitted speaker notes")
    return {"native_text": True, "native_tables": True, "native_charts": True, "native_diagrams": True,
            "images": "source-or-derived-raster", "formulas": "source-or-mathtext-raster", "required_objects": "CHECKED", "rich-assets-v1": True}
