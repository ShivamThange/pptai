import json
import subprocess
import sys
from pathlib import Path

from pptx import Presentation

from evalute.metrics.faithfulness import _claims

ROOT = Path(__file__).parents[2]


def test_claims_are_numbered_sequentially():
    slides = [
        {"slide_no": 1, "text": "This sentence has more than seven words for evaluation."},
        {"slide_no": 2, "text": "The 2026 result is included here."},
    ]

    assert _claims(slides) == [
        {"claim_id": "CL001", "slide_no": 1, "claim": "This sentence has more than seven words for evaluation."},
        {"claim_id": "CL002", "slide_no": 2, "claim": "The 2026 result is included here."},
    ]


def test_cli(tmp_path):
    ir = {"document_name": "source.pdf", "blocks": [{"block_id": "block_0001", "section": "Intro", "text": "Attention uses weighted combinations of values."}]}
    ir1 = {"document_title": "Attention", "executive_summary": "Attention is central.", "concepts": [{"title": "Attention", "summary": "Weighted combinations of values", "bullets": [], "source_block_ids": ["block_0001"]}]}
    (tmp_path / "ir.json").write_text(json.dumps(ir)); (tmp_path / "ir1.json").write_text(json.dumps(ir1))
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[1]); slide.shapes.title.text = "Attention"; slide.placeholders[1].text = "Attention uses weighted combinations of values."
    pptx = tmp_path / "deck.pptx"; deck.save(pptx); reference = tmp_path / "reference.json"; result = tmp_path / "result"
    base = [sys.executable, "-m", "evalute.cli"]
    subprocess.run(base + ["prepare-reference", "--ir", str(tmp_path / "ir.json"), "--ir1", str(tmp_path / "ir1.json"), "--out", str(reference)], cwd=ROOT, check=True)
    subprocess.run(base + ["evaluate", "--reference", str(reference), "--pptx", str(pptx), "--tool-name", "test", "--run-id", "test", "--out", str(result), "--no-ai-judging"], cwd=ROOT, check=True)
    assert (result / "report.html").exists()
